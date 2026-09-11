# -*- coding: utf-8 -*-
"""
Australia TV Price Tracking Dashboard Generator
Compiles 'data/au_price_dashboard_template.html' into:
- data/au_price_dashboard.html
- public_au/index.html
- History_AU/{YYYY MMDD}/au_price_dashboard.html
"""

import openpyxl
import json
import os
import re
import sys
import glob
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DATA_DIR = os.path.join(ROOT_DIR, "data")
PUBLIC_AU_DIR = os.path.join(ROOT_DIR, "public_au")
HISTORY_DIR = os.path.join(ROOT_DIR, "History_AU")
os.makedirs(PUBLIC_AU_DIR, exist_ok=True)

def find_latest_au_workbook():
    today_mmdd = datetime.now().strftime("%m%d")
    p = os.path.join(DATA_DIR, f"price tracker_AU_2026 {today_mmdd}_v1.xlsx")
    if os.path.exists(p):
        return p
    candidates = sorted(glob.glob(os.path.join(DATA_DIR, "price tracker_AU_*.xlsx")))
    return candidates[-1] if candidates else None

def generate_au_dashboard():
    wb_path = find_latest_au_workbook()
    if not wb_path or not os.path.exists(wb_path):
        print("[ERROR] AU Workbook not found!")
        return
        
    print(f"[PARSING EXCEL] Loading latest report: {wb_path}")
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    
    now_dt = datetime.now()
    week_no = now_dt.isocalendar()[1]
    today_date_str = now_dt.strftime("%Y-%m-%d")
    week_key = f"{now_dt.year}-W{week_no}"
    
    # 1. Parse all models from retailer sheets
    ret_sheets = [
        ("JBHIFI_Samsung_Full", "msh", "JB Hi-Fi", "SAMSUNG"),
        ("JBHIFI_LG_Full", "msh", "JB Hi-Fi", "LG"),
        ("TheGoodGuys_Samsung_Full", "id", "The Good Guys", "SAMSUNG"),
        ("TheGoodGuys_LG_Full", "id", "The Good Guys", "LG"),
        ("HarveyNorman_Samsung_Full", "digi", "Harvey Norman", "SAMSUNG"),
        ("HarveyNorman_LG_Full", "digi", "Harvey Norman", "LG")
    ]
    
    all_models = []
    for s_name, ret_key, ret_name, brand in ret_sheets:
        if s_name not in wb.sheetnames:
            continue
        ws = wb[s_name]
        for r in range(2, ws.max_row + 1):
            yr = ws.cell(row=r, column=2).value
            cat = ws.cell(row=r, column=3).value
            size = ws.cell(row=r, column=4).value
            mc = ws.cell(row=r, column=5).value
            price = ws.cell(row=r, column=6).value
            orig_p = ws.cell(row=r, column=7).value
            promo = ws.cell(row=r, column=10).value
            title = ws.cell(row=r, column=11).value
            url = ws.cell(row=r, column=12).value
            
            if mc and price:
                try:
                    p_val = float(price)
                    op_val = float(orig_p) if orig_p else p_val
                    all_models.append({
                        "brand": brand,
                        "year": int(yr) if yr else 2026,
                        "category": str(cat or "UHD"),
                        "size": int(size) if size else 55,
                        "model_code": str(mc),
                        "price": p_val,
                        "orig_price": op_val,
                        "retailer_key": ret_key,
                        "retailer": ret_name,
                        "promo": str(promo or "None"),
                        "title": str(title or ""),
                        "url": str(url or "")
                    })
                except ValueError:
                    pass
                    
    print(f"  ➔ Extracted {len(all_models)} total active models across 3 AU retailers.")
    
    # Pairs configuration across all segments
    pairs_cfg = {
        2026: {
            "OLED": [
                {"lg": "G6", "sam": "S95H", "sizes": [83, 77, 65, 55]},
                {"lg": "C6", "sam": "S90H", "sizes": [83, 77, 65, 55, 48, 42]},
                {"lg": "B6", "sam": "S85H", "sizes": [83, 77, 65, 55, 48]}
            ],
            "MRGB": [
                {"lg": "MRGB87B", "sam": "R85H", "sizes": [86, 75, 65, 55, 50]}
            ],
            "QNED/QLED": [
                {"lg": "QNED86B", "sam": "QN80H", "sizes": [86, 75, 65, 55, 50, 43]},
                {"lg": "QNED80B", "sam": "QN85H", "sizes": [86, 75, 65, 55, 50, 43]},
                {"lg": "QNED70B", "sam": "M70H", "sizes": [86, 75, 65, 55, 50, 43]}
            ],
            "UHD 4K": [
                {"lg": "UT80", "sam": "U8000H", "sizes": [85, 75, 65, 55, 50, 43]},
                {"lg": "UA77", "sam": "U8000H", "sizes": [85, 75, 65, 55, 50, 43]}
            ]
        },
        2025: {
            "OLED": [
                {"lg": "G5", "sam": "S95F", "sizes": [83, 77, 65, 55]},
                {"lg": "C5", "sam": "S90F", "sizes": [83, 77, 65, 55, 48, 42]},
                {"lg": "B5", "sam": "S85F", "sizes": [83, 77, 65, 55, 48]}
            ],
            "QNED/QLED": [
                {"lg": "QNED86A", "sam": "QN80F", "sizes": [86, 75, 65, 55, 50, 43]},
                {"lg": "QNED80A", "sam": "Q7F", "sizes": [86, 75, 65, 55, 50, 43]}
            ],
            "UHD 4K": [
                {"lg": "UA75", "sam": "U8000F", "sizes": [85, 75, 65, 55, 50, 43]}
            ]
        }
    }
    
    def find_best_model(model_list, ret_k, brand_name, sz, series_name):
        # 1. Exact match with size & series code
        candidates = [
            m for m in model_list 
            if m["retailer_key"] == ret_k and m["brand"] == brand_name and m["size"] == sz and series_name.upper() in m["model_code"].upper()
        ]
        # 2. Match with title
        if not candidates:
            candidates = [
                m for m in model_list 
                if m["retailer_key"] == ret_k and m["brand"] == brand_name and m["size"] == sz and series_name.upper() in m["title"].upper()
            ]
        # 3. Micro RGB fallback (R85 / MRGB)
        if not candidates and series_name in ["R85H", "MRGB87B"]:
            candidates = [
                m for m in model_list 
                if m["retailer_key"] == ret_k and m["brand"] == brand_name and m["size"] == sz and any(k in m["model_code"].upper() for k in ["R85", "MRGB", "MICRO"])
            ]
        if candidates:
            candidates.sort(key=lambda x: x["price"])
            return candidates[0]
        return None

    pairs_out = {
        "2026": {},
        "2025": {}
    }
    
    history_out = {
        "2026": {},
        "2025": {}
    }

    for yr in [2026, 2025]:
        yr_str = str(yr)
        yr_models = [m for m in all_models if m["year"] == yr]
        
        for cat, series_list in pairs_cfg[yr].items():
            pairs_out[yr_str][cat] = []
            for item in series_list:
                lg_fam = item["lg"]
                sam_fam = item["sam"]
                for sz in item["sizes"]:
                    lg_series_key = f"{sz}{lg_fam}"
                    sam_series_key = f"{sz}{sam_fam}"
                    
                    row = {
                        "lg_series": lg_series_key,
                        "sam_series": sam_series_key,
                        "size": str(sz),
                        # JB Hi-Fi (msh)
                        "lg_price_msh": 0, "sam_price_msh": 0,
                        "lg_net_msh": 0, "sam_net_msh": 0,
                        "lg_model_msh": "", "sam_model_msh": "",
                        "lg_promo_msh": "", "sam_promo_msh": "",
                        # The Good Guys (id)
                        "lg_price_id": 0, "sam_price_id": 0,
                        "lg_net_id": 0, "sam_net_id": 0,
                        "lg_model_id": "", "sam_model_id": "",
                        "lg_promo_id": "", "sam_promo_id": "",
                        # Harvey Norman (digi)
                        "lg_price_digi": 0, "sam_price_digi": 0,
                        "lg_net_digi": 0, "sam_net_digi": 0,
                        "lg_model_digi": "", "sam_model_digi": "",
                        "lg_promo_digi": "", "sam_promo_digi": ""
                    }
                    
                    for ret_k in ["msh", "id", "digi"]:
                        l_m = find_best_model(yr_models, ret_k, "LG", sz, lg_fam)
                        s_m = find_best_model(yr_models, ret_k, "SAMSUNG", sz, sam_fam)
                        
                        if l_m:
                            row[f"lg_price_{ret_k}"] = int(round(l_m["price"]))
                            row[f"lg_net_{ret_k}"] = int(round(l_m["price"]))
                            row[f"lg_model_{ret_k}"] = l_m["model_code"]
                            row[f"lg_promo_{ret_k}"] = l_m["promo"] if l_m["promo"] != "None" else ""
                        if s_m:
                            row[f"sam_price_{ret_k}"] = int(round(s_m["price"]))
                            row[f"sam_net_{ret_k}"] = int(round(s_m["price"]))
                            row[f"sam_model_{ret_k}"] = s_m["model_code"]
                            row[f"sam_promo_{ret_k}"] = s_m["promo"] if s_m["promo"] != "None" else ""
                            
                    pairs_out[yr_str][cat].append(row)
                    
                    # Register in history_out
                    for s_key, b_name, fam, p_prefix in [(lg_series_key, "LG", lg_fam, "lg_"), (sam_series_key, "SAMSUNG", sam_fam, "sam_")]:
                        if any(row[f"{p_prefix}price_{rk}"] > 0 for rk in ["msh", "id", "digi"]):
                            history_out[yr_str][s_key] = {
                                "brand": b_name,
                                "size": sz,
                                "series": fam,
                                "history": [
                                    {
                                        "date": today_date_str,
                                        "msh_price": row[f"{p_prefix}price_msh"],
                                        "id_price": row[f"{p_prefix}price_id"],
                                        "digi_price": row[f"{p_prefix}price_digi"],
                                        "msh_net": row[f"{p_prefix}net_msh"],
                                        "id_net": row[f"{p_prefix}net_id"],
                                        "digi_net": row[f"{p_prefix}net_digi"]
                                    }
                                ]
                            }
                            
    # Build weekly_changes structure
    weekly_changes_out = {
        week_key: {
            "date": today_date_str,
            "changes": []
        }
    }
    
    price_data = {
        "2026": pairs_out["2026"],
        "2025": pairs_out["2025"],
        "history": history_out,
        "weekly_changes": weekly_changes_out,
        "currency": "AUD",
        "currency_symbol": "A$",
        "survey_date": now_dt.strftime("%Y.%m.%d"),
        "retailers": ["JB Hi-Fi", "The Good Guys", "Harvey Norman"],
        "models": all_models
    }
    
    # Compile Dashboard HTML
    template_path = os.path.join(DATA_DIR, "au_price_dashboard_template.html")
    if not os.path.exists(template_path):
        print(f"[ERROR] Template not found: {template_path}")
        return
        
    with open(template_path, 'r', encoding='utf-8') as f:
        tmpl = f.read()
        
    json_str = json.dumps(price_data, ensure_ascii=False, separators=(',', ':'))
    injection = f"const priceData = {json_str};"
    
    survey_date_kr = now_dt.strftime(f"%Y년 %m월 %d일 (W{week_no})")
    survey_date_dot = now_dt.strftime(f"%Y.%m.%d(W{week_no})")
    
    compiled = tmpl.replace("// {{INSERT_PRICE_DATA}}", injection)
    compiled = compiled.replace("{{SURVEY_DATE}}", survey_date_kr)
    compiled = compiled.replace("{{SURVEY_DATE_DOT}}", survey_date_dot)
    
    # Save outputs
    out_workspace = os.path.join(DATA_DIR, "au_price_dashboard.html")
    out_public = os.path.join(PUBLIC_AU_DIR, "index.html")
    
    with open(out_workspace, 'w', encoding='utf-8') as f:
        f.write(compiled)
    with open(out_public, 'w', encoding='utf-8') as f:
        f.write(compiled)
        
    # Mirror to History_AU
    mmdd = now_dt.strftime("%m%d")
    hist_dir = os.path.join(HISTORY_DIR, f"{now_dt.strftime('%Y')} {mmdd}")
    os.makedirs(hist_dir, exist_ok=True)
    out_hist = os.path.join(hist_dir, "au_price_dashboard.html")
    with open(out_hist, 'w', encoding='utf-8') as f:
        f.write(compiled)
        
    print(f"✅ [COMPILED HTML] {out_workspace}")
    print(f"✅ [FIREBASE PUBLIC] {out_public}")
    print(f"✅ [HISTORY_AU BACKUP] {out_hist}")

if __name__ == '__main__':
    generate_au_dashboard()
