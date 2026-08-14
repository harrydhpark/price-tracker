# -*- coding: utf-8 -*-
import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('data/price tracker_EU_2026 0814_v1.xlsx', data_only=True)
print('=== LG & Samsung MRGB Models Across European Retailers ===')

for brand in ['LG', 'SAMSUNG']:
    print(f"\n==================== {brand} MRGB ====================")
    for sname in wb.sheetnames:
        if f"_{brand}" in sname.upper():
            ws = wb[sname]
            retailer = sname
            matches = []
            for r in range(2, ws.max_row+1):
                code = str(ws.cell(r, 5).value or '')
                name = str(ws.cell(r, 7).value or '')
                price = ws.cell(r, 6).value
                size = ws.cell(r, 4).value
                if any(k in code.upper() for k in ['MRGB', 'MR9', 'MR8', 'MRE', 'TMR', 'R85', 'R86', 'R95', 'R96']) or any(k in name.upper() for k in ['MICRO RGB', 'MICRO-RGB']):
                    matches.append((code, size, price, name))
            print(f"[{retailer}] ({len(matches)} models):")
            if not matches:
                print("  - None listed")
            for code, size, price, name in matches:
                print(f"  - {size}\" | Code: {code} | Price: {price} | Name: {name[:60]}")
