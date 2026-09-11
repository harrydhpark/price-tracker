import os
import sys
import re
import json
import glob
import time
from datetime import datetime
from bs4 import BeautifulSoup
from scrapling.fetchers import StealthySession
import openpyxl

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding='utf-8')

DATA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
HISTORY_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "History_EU"))

def parse_price(price_str):
    if not price_str:
        return 0.0
    cleaned = str(price_str).replace('€', '').replace(' ', '').replace('\xa0', '').strip()
    if not cleaned:
        return 0.0
    if ',' in cleaned and '.' in cleaned:
        cleaned = cleaned.replace('.', '').replace(',', '.')
    elif ',' in cleaned:
        cleaned = cleaned.replace(',', '.')
    try:
        val = float(cleaned)
        return val if val >= 50.0 else 0.0
    except ValueError:
        return 0.0

def extract_model_code_and_info(title_upper, brand, href):
    year_val = None
    series_code = "Unknown"
    model_code = "Unknown"
    size_val = 0

    m_size = re.search(r'(\d{2,3})[\"\”\s]*(?:4K|Smart|OLED|QLED|LED|QNED|Nano|Mini|Micro|\b)', title_upper)
    if m_size:
        try:
            size_val = int(m_size.group(1))
        except ValueError:
            size_val = 0
            
    if size_val < 22:
        return None  # Exclude small screens, monitors and accessories

    if brand.lower() == "samsung":
        patterns_2026 = [
            (r'S99H', 'S99H', 2026), (r'S95H', 'S95H', 2026), (r'S90H', 'S90H', 2026), (r'S85H', 'S85H', 2026),
            (r'QN95H', 'QN95H', 2026), (r'QN90H', 'QN90H', 2026), (r'QN85H', 'QN85H', 2026),
            (r'QN80H', 'QN80H', 2026), (r'QN70H', 'QN70H', 2026), (r'M80H', 'M80H', 2026),
            (r'M70H', 'M70H', 2026), (r'R95H', 'R95H', 2026), (r'R86H', 'R86H', 2026), (r'R85H', 'R85H', 2026),
            (r'U8072H', 'U8000H', 2026), (r'U8070H', 'U8000H', 2026), (r'U8092H', 'U8000H', 2026),
            (r'U8072', 'U8000H', 2026), (r'U8092', 'U8000H', 2026), (r'U8000H', 'U8000H', 2026),
            (r'LS03HW', 'The Frame Pro', 2026), (r'LS03H', 'The Frame', 2026), (r'F6000H', 'F6000H', 2026)
        ]
        patterns_2025 = [
            (r'S95F', 'S95F', 2025), (r'S90F', 'S90F', 2025), (r'S85F', 'S85F', 2025),
            (r'QN95F', 'QN95F', 2025), (r'QN90F', 'QN90F', 2025), (r'QN85F', 'QN85F', 2025),
            (r'QN80F', 'QN80F', 2025), (r'QN70F', 'QN70F', 2025), (r'Q80F', 'Q8F', 2025),
            (r'Q8F', 'Q8F', 2025), (r'Q70F', 'Q7F', 2025), (r'Q7F', 'Q7F', 2025),
            (r'Q60F', 'Q6F', 2025), (r'Q6F', 'Q6F', 2025), (r'U8072F', 'U8000F', 2025), (r'U8000F', 'U8000F', 2025),
            (r'LS03F', 'The Frame', 2025)
        ]

        m_code = re.search(r'([A-Z]{2}\d{2}[A-Z0-9]+|\b\d{2}[A-Z]{1,3}\d{2,3}[A-Z0-9]*\b)', title_upper)
        if m_code:
            model_code = m_code.group(1)
        else:
            slug_m = re.search(r'/([a-z0-9\-]+)/?\d*$', href.lower())
            if slug_m:
                model_code = slug_m.group(1).upper()

        for pat, ser, yr in patterns_2026:
            if re.search(pat, title_upper):
                series_code = ser
                year_val = yr
                break

        if not year_val:
            for pat, ser, yr in patterns_2025:
                if re.search(pat, title_upper):
                    series_code = ser
                    year_val = yr
                    break

        if not year_val:
            if "2026" in title_upper or ("AI TV" in title_upper and any(x in title_upper for x in ["H ", "H\b", "95H", "90H", "85H", "80H", "70H"])):
                year_val = 2026
            elif "2025" in title_upper:
                year_val = 2025

    elif brand.lower() == "lg":
        lg_2026 = [
            # QNED series (must come before loose OLED B6 pattern)
            (r'QNED93B', 'QNED93B', 2026), (r'QNED87B', 'QNED87B', 2026), (r'QNED86B', 'QNED86B', 2026),
            (r'QNED85B', 'QNED85B', 2026), (r'QNED83B', 'QNED81B', 2026), (r'QNED82B', 'QNED82B', 2026),
            (r'QNED81B', 'QNED81B', 2026), (r'QNED80B', 'QNED80B', 2026),
            (r'QNED72B', 'QNED72B', 2026), (r'QNED71B', 'QNED70B', 2026), (r'QNED70B', 'QNED70B', 2026),
            (r'QNED7EB', 'QNED7EB', 2026),
            # Micro RGB series
            (r'MRGB96B', 'MRGB96B', 2026), (r'MRGB88B', 'MRGB88B', 2026),
            (r'MRGB87B', 'MRGB87B', 2026), (r'MRGB86B', 'MRGB86B', 2026), (r'MRGB85B', 'MRGB85B', 2026),
            # Lifestyle
            (r'LX7B', 'LX7B', 2026), (r'LX6', 'LX6', 2026), (r'27LX6', '27LX6', 2026), (r'STANBYME', 'STANBYME', 2026),
            # OLED series (digit-aware patterns to match G64, C64, B66 etc.)
            (r'OLED\d{2}G6', 'OLED G6', 2026), (r'G6[0-9]L', 'OLED G6', 2026), (r'G6[0-9]', 'OLED G6', 2026),
            (r'OLED\d{2}C6', 'OLED C6', 2026), (r'C6[0-9]L', 'OLED C6', 2026), (r'C6[0-9]', 'OLED C6', 2026),
            (r'OLED\d{2}B6', 'OLED B6', 2026), (r'B6[0-9]L', 'OLED B6', 2026), (r'B6[0-9]', 'OLED B6', 2026),
            (r'OLED\d{2}W6', 'OLED W6', 2026), (r'W6[0-9]L', 'OLED W6', 2026),
            # UHD 4K
            (r'NU90', 'NU90', 2026), (r'NU85', 'NU85', 2026), (r'NU8E', 'NU85', 2026),
            (r'UA77', 'UA77', 2026), (r'LB650B', 'LB650B', 2026)
        ]
        lg_2025 = [
            # QNED series (including sub-model aliases)
            (r'QNED93A', 'QNED93A', 2025),
            (r'QNED87A', 'QNED86A', 2025), (r'QNED86A', 'QNED86A', 2025), (r'QNED85A', 'QNED86A', 2025),
            (r'QNED84A', 'QNED80A', 2025), (r'QNED83A', 'QNED80A', 2025), (r'QNED82A', 'QNED80A', 2025),
            (r'QNED80A', 'QNED80A', 2025), (r'QNED8EA', 'QNED8EA', 2025),
            (r'QNED70A', 'QNED70A', 2025),
            # NanoCell / Nano series
            (r'NANO90A', 'NANO90A', 2025), (r'NANO81A', 'NANO81A', 2025), (r'NANO80A', 'NANO80A', 2025),
            # OLED series (digit-aware patterns to match G56, C55, B56 etc.)
            (r'OLED\d{2}G5', 'OLED G5', 2025), (r'G5[0-9]L', 'OLED G5', 2025), (r'G5[0-9]', 'OLED G5', 2025),
            (r'OLED\d{2}C5', 'OLED C5', 2025), (r'C5[0-9]L', 'OLED C5', 2025), (r'C5[0-9]', 'OLED C5', 2025),
            (r'OLED\d{2}B5', 'OLED B5', 2025), (r'B5[0-9]L', 'OLED B5', 2025), (r'B5[0-9]', 'OLED B5', 2025),
            # UHD 4K
            (r'UA75', 'UA75', 2025), (r'UA74', 'UA75', 2025), (r'UA73', 'UA75', 2025)
        ]

        m_code = re.search(r'([A-Z0-9]{2,4}\d{2}[A-Z0-9]+|\b\d{2}[A-Z0-9]{5,10}\b|\bOLED\d{2}[A-Z0-9]+\b)', title_upper)
        if m_code:
            model_code = m_code.group(1)
        else:
            m_code2 = re.search(r'\((OLED[A-Z0-9]+|[A-Z0-9]{8,15})\)', title_upper)
            if m_code2:
                model_code = m_code2.group(1)

        for pat, ser, yr in lg_2026:
            if re.search(pat, title_upper):
                series_code = ser
                year_val = yr
                break

        if not year_val:
            for pat, ser, yr in lg_2025:
                if re.search(pat, title_upper):
                    series_code = ser
                    year_val = yr
                    break

        if not year_val:
            if "2026" in title_upper:
                year_val = 2026
            elif "2025" in title_upper:
                year_val = 2025

    if not year_val:
        return None  # Purge older legacy models

    disp_type = "UHD 4K"
    if "OLED" in series_code.upper() or "OLED" in title_upper:
        disp_type = "OLED"
    elif "MRGB" in series_code.upper() or "MICRO RGB" in title_upper:
        disp_type = "MRGB"
    elif "QNED" in series_code.upper() or "QNED" in title_upper:
        disp_type = "QNED/QLED"
    elif any(x in series_code.upper() or x in title_upper for x in ["QLED", "NEO QLED", "MINI LED", "Q8", "Q7", "QN"]):
        disp_type = "QNED/QLED"

    return {
        "brand": brand.upper(),
        "title": title_upper,
        "series": series_code,
        "model_code": model_code,
        "size": size_val,
        "year": year_val,
        "display_type": disp_type
    }

def scrape_public_gr(brand):
    if brand.lower() == "samsung":
        queries = [
            # Core brand + product type queries
            "https://www.public.gr/search?q=samsung%20tv",
            "https://www.public.gr/search?q=samsung%20oled",
            "https://www.public.gr/search?q=samsung%20qled",
            "https://www.public.gr/search?q=samsung%20neo%20qled",
            "https://www.public.gr/search?q=samsung%204k",
            # Year-specific queries
            "https://www.public.gr/search?q=samsung%20tv%202026",
            "https://www.public.gr/search?q=samsung%20tv%202025",
            # Technology-specific keyword reorder queries
            "https://www.public.gr/search?q=samsung%20tv%20oled",
            "https://www.public.gr/search?q=samsung%20tv%20qled",
            "https://www.public.gr/search?q=samsung%20tv%204k",
            # Size-specific queries (high yield: 21-28 new each)
            "https://www.public.gr/search?q=samsung%20tv%2055",
            "https://www.public.gr/search?q=samsung%20tv%2065",
            "https://www.public.gr/search?q=samsung%20tv%2075",
            "https://www.public.gr/search?q=samsung%20tv%2077",
            "https://www.public.gr/search?q=samsung%20tv%2083",
            "https://www.public.gr/search?q=samsung%20tv%2085",
            # Product line queries
            "https://www.public.gr/search?q=samsung%20mini%20led",
            "https://www.public.gr/search?q=samsung%20micro",
            "https://www.public.gr/search?q=samsung%20the%20frame",
            # Greek-language & alternative queries
            "https://www.public.gr/search?q=tileoraseis%20samsung",
            "https://www.public.gr/search?q=samsung%20smart%20tv",
            # Category browsing fallback
            "https://www.public.gr/cat/tileoraseis/tileoraseis/"
        ]
    else:
        queries = [
            # Core brand + product type queries
            "https://www.public.gr/search?q=lg%20tv",
            "https://www.public.gr/search?q=lg%20oled",
            "https://www.public.gr/search?q=lg%20qned",
            "https://www.public.gr/search?q=lg%204k",
            # OLED series-specific queries
            "https://www.public.gr/search?q=lg%20oled%20c6",
            "https://www.public.gr/search?q=lg%20oled%20g6",
            "https://www.public.gr/search?q=lg%20oled%20b6",
            "https://www.public.gr/search?q=lg%20oled%20c5",
            "https://www.public.gr/search?q=lg%20c5",
            "https://www.public.gr/search?q=lg%20g5",
            "https://www.public.gr/search?q=lg%20b5",
            "https://www.public.gr/search?q=lg%20b6",
            "https://www.public.gr/search?q=lg%20c6",
            "https://www.public.gr/search?q=lg%20g6",
            # QNED series-specific queries
            "https://www.public.gr/search?q=lg%20qned%2093",
            "https://www.public.gr/search?q=lg%20qned%2087",
            "https://www.public.gr/search?q=lg%20qned%2086",
            "https://www.public.gr/search?q=lg%20qned%2080",
            "https://www.public.gr/search?q=lg%20qned%2072",
            # Product line queries
            "https://www.public.gr/search?q=lg%20nano",
            "https://www.public.gr/search?q=lg%20nu85",
            "https://www.public.gr/search?q=lg%20mrgb",
            "https://www.public.gr/search?q=lg%20micro",
            "https://www.public.gr/search?q=lg%20ua",
            # Year-specific queries (high yield: 36 new each)
            "https://www.public.gr/search?q=lg%20tv%202026",
            "https://www.public.gr/search?q=lg%20tv%202025",
            # Technology reorder queries
            "https://www.public.gr/search?q=lg%20tv%20oled",
            "https://www.public.gr/search?q=lg%20tv%20qned",
            "https://www.public.gr/search?q=lg%20tv%204k",
            # Size-specific queries (high yield: 18-24 new each)
            "https://www.public.gr/search?q=lg%20tv%2055",
            "https://www.public.gr/search?q=lg%20tv%2065",
            "https://www.public.gr/search?q=lg%20tv%2075",
            "https://www.public.gr/search?q=lg%20tv%2077",
            "https://www.public.gr/search?q=lg%20tv%2083",
            # Greek-language & alternative queries
            "https://www.public.gr/search?q=tileoraseis%20lg",
            "https://www.public.gr/search?q=lg%20smart%20tv",
            # Category browsing fallback
            "https://www.public.gr/cat/tileoraseis/tileoraseis/"
        ]

    print(f"\n[INFO] Starting Multi-Query Targeted Public Greece collection for {brand} ({len(queries)} query configurations)...")
    items = []
    seen_hrefs = set()

    with StealthySession(headless=True) as session:
        for idx, url in enumerate(queries, 1):
            print(f"[FETCH {idx}/{len(queries)}] Requesting query URL: {url}")
            
            try:
                res = session.fetch(url, wait=3000)
                if res.status != 200:
                    print(f"[WARN] Non-200 status {res.status} on query {idx}")
                    continue
                    
                body = res.body.decode('utf-8', errors='ignore') if isinstance(res.body, bytes) else res.body
                soup = BeautifulSoup(body, 'html.parser')
                
                product_anchors = soup.find_all('a', href=re.compile(r'/product/tileoraseis/tileoraseis/'))
                print(f" -> Found {len(product_anchors)} raw product anchors")
                
                page_extracted = 0
                for a in product_anchors:
                    href = a['href']
                    if href in seen_hrefs:
                        continue
                    title = a.get_text(strip=True)
                    if not title or brand.lower() not in title.lower():
                        continue
                        
                    info = extract_model_code_and_info(title, brand, href)
                    if not info:
                        continue
                    size_val = info["size"]
                    seen_hrefs.add(href)
                    
                    # Climb up parent to card container
                    card = a
                    for _ in range(10):
                        if card.parent:
                            p_text = card.parent.get_text(' ', strip=True)
                            if '€' in p_text and len(p_text) < 3000:
                                card = card.parent
                                break
                            card = card.parent
                            
                    card_text = card.get_text(' ', strip=True)
                    
                    # Extract prices
                    prices = re.findall(r'(\d[\d\.\s]*[,\.]\d{2})\s*€', card_text)
                    selling_price = 0.0
                    orig_price = 0.0
                    
                    if prices:
                        parsed_prices = []
                        for p in prices:
                            val = parse_price(p)
                            # Screen size threshold guard
                            if size_val >= 70 and val < 700.0:
                                continue
                            if size_val >= 55 and val < 350.0:
                                continue
                            if size_val >= 48 and val < 250.0:
                                continue
                            if val >= 100.0:
                                parsed_prices.append(val)
                                
                        if parsed_prices:
                            selling_price = min(parsed_prices)
                            if len(parsed_prices) > 1:
                                orig_price = max(parsed_prices)
                            else:
                                orig_price = selling_price
                                
                    if selling_price < 100.0:
                        continue  # Skip items with no valid price
                        
                    # Extract promotion
                    promo_text = "None"
                    if "Δώρο" in card_text or "Soundbar" in card_text or "Projector" in card_text:
                        m_promo = re.search(r'(Δώρο[^\n\.\,]+|Soundbar[^\n\.\,]+)', card_text)
                        if m_promo:
                            promo_text = m_promo.group(1).strip()
                            
                    full_url = f"https://www.public.gr{href}" if href.startswith('/') else href
                    
                    item = {
                        "brand": info["brand"],
                        "title": title,
                        "series": info["series"],
                        "model_code": info["model_code"],
                        "size": info["size"],
                        "year": info["year"],
                        "display_type": info["display_type"],
                        "original_price": orig_price,
                        "price": selling_price,
                        "cashback": 0.0,
                        "net_price": selling_price,
                        "promotion": promo_text,
                        "url": full_url,
                        "image_url": ""
                    }
                    items.append(item)
                    page_extracted += 1
                    
                print(f" -> Extracted {page_extracted} new unique items (Total accumulated: {len(items)})")
                
            except Exception as e:
                print(f"[ERROR] Exception on query {idx}: {e}")
                continue

    # Save raw JSON
    json_path = os.path.join(DATA_DIR, f"raw_public_gr_{brand.lower()}.json")
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(items, f, ensure_ascii=False, indent=2)
    print(f"[SAVE] Raw JSON saved: {json_path} ({len(items)} items)")
    return items

def update_excel_workbook(samsung_items, lg_items):
    today_mmdd = datetime.now().strftime("%m%d")
    excel_path = os.path.join(DATA_DIR, f"price tracker_EU_2026 {today_mmdd}_v1.xlsx")
    if not os.path.exists(excel_path):
        eu_files = glob.glob(os.path.join(DATA_DIR, "price tracker_EU_2026 *.xlsx"))
        eu_files = [f for f in eu_files if not os.path.basename(f).startswith("~$")]
        if eu_files:
            eu_files.sort()
            import shutil
            shutil.copy2(eu_files[-1], excel_path)
            print(f"[CLONE] Created today's workbook from {eu_files[-1]} -> {excel_path}")
        else:
            print(f"[ERROR] Target Excel path not found: {excel_path}")
            return

    print(f"[EXCEL] Updating EU workbook: {excel_path}")
    wb = openpyxl.load_workbook(excel_path)
    
    sheets_config = [
        ("Public_GR_Samsung", samsung_items),
        ("Public_GR_LG", lg_items)
    ]
    
    for sheet_name, items in sheets_config:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(title=sheet_name)
            
        headers = [
            "Brand", "Model Year", "Display Type", "Size (Inch)", "Model Code",
            "Price (EUR)", "Original Price (EUR)", "Series", "Cashback", "General Promotions",
            "URL", "Title"
        ]
        ws.append(headers)
        
        for item in items:
            ws.append([
                item["brand"], item["year"], item["display_type"], item["size"], item["model_code"],
                item["price"], item["original_price"], item["series"], item["cashback"], item["promotion"],
                item["url"], item["title"]
            ])
            
    wb.save(excel_path)
    print(f"[EXCEL] Successfully updated workbook: {excel_path} with Public_GR sheets!")
    
    try:
        import shutil
        today_folder = datetime.now().strftime("%Y %m%d")
        hist_folder_path = os.path.join(HISTORY_DIR, today_folder)
        os.makedirs(hist_folder_path, exist_ok=True)
        hist_wb_path = os.path.join(hist_folder_path, os.path.basename(excel_path))
        shutil.copy2(excel_path, hist_wb_path)
        print(f"[MIRROR] Saved backup to History_EU: {hist_wb_path}")
    except Exception as e:
        print(f"[WARN] Failed to mirror history workbook: {e}")

if __name__ == "__main__":
    import time
    start_time = time.time()
    print("=" * 60)
    print(" GREECE PUBLIC (public.gr) MULTI-QUERY EXPANDED TV SURVEY ")
    print(" Target Brands: LG & Samsung | Target Model Years: 2025 & 2026")
    print("=" * 60)
    
    samsung_data = scrape_public_gr("Samsung")
    lg_data = scrape_public_gr("LG")
    
    update_excel_workbook(samsung_data, lg_data)
    print("\n[COMPLETE] Multi-Query Public Greece price collection finished cleanly!")
    
    try:
        from scrape_logger import emit_summary
        elapsed = round(time.time() - start_time, 2)
        emit_summary(country="GR", retailer="Public", brand="SAMSUNG", total_extracted=len(samsung_data), final_deduplicated=len(samsung_data), execution_time_sec=elapsed)
        emit_summary(country="GR", retailer="Public", brand="LG", total_extracted=len(lg_data), final_deduplicated=len(lg_data), execution_time_sec=elapsed)
    except Exception:
        pass
