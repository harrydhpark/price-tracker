# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb_old = openpyxl.load_workbook(r'History_EU\2026 0903\price tracker_EU_2026 0903_v1.xlsx', data_only=True)
wb_new = openpyxl.load_workbook(r'data\price tracker_EU_2026 0907_v1.xlsx', data_only=True)

def get_sheet_dict(wb, sname):
    if sname not in wb.sheetnames:
        return {}
    ws = wb[sname]
    res = {}
    for r in range(2, ws.max_row + 1):
        brand = ws.cell(row=r, column=1).value
        year = ws.cell(row=r, column=2).value
        size = ws.cell(row=r, column=4).value
        code = ws.cell(row=r, column=5).value
        price = ws.cell(row=r, column=6).value
        promo = ws.cell(row=r, column=10).value
        if code and price:
            try:
                p_val = float(str(price).replace(',', '').replace(' ', ''))
                res[str(code).strip().upper()] = {
                    'brand': brand,
                    'year': year,
                    'size': size,
                    'code': str(code).strip(),
                    'price': p_val,
                    'promo': str(promo or 'None')
                }
            except:
                pass
    return res

common_sheets = [s for s in wb_new.sheetnames if s in wb_old.sheetnames]
print(f"Comparing {len(common_sheets)} sheets between W36 (0903) and W37 (0907)...\n")

for s in common_sheets:
    old_map = get_sheet_dict(wb_old, s)
    new_map = get_sheet_dict(wb_new, s)
    
    drops = []
    hikes = []
    p_changes = []
    new_items = []
    
    for c, it in new_map.items():
        if c in old_map:
            old_it = old_map[c]
            diff = round(it['price'] - old_it['price'], 2)
            if diff < -0.5:
                drops.append((it, old_it['price'], diff))
            elif diff > 0.5:
                hikes.append((it, old_it['price'], diff))
            elif it['promo'] != old_it['promo']:
                p_changes.append((it, old_it['promo']))
        else:
            new_items.append(it)
            
    if drops or hikes or p_changes or new_items:
        print(f"=== [{s}] ===")
        print(f"  - 인하(Drops): {len(drops)}건 | 인상(Hikes): {len(hikes)}건 | 프로모션변경: {len(p_changes)}건 | 신규: {len(new_items)}건")
        for it, old_p, diff in drops[:5]:
            pct = round(diff / old_p * 100, 1)
            print(f"    * [인하] {it['code']} ({it['size']}\"): {old_p:,.0f} ➔ {it['price']:,.0f} ({diff:+,.0f}, {pct}%)")
        for it, old_p, diff in hikes[:5]:
            pct = round(diff / old_p * 100, 1)
            print(f"    * [인상] {it['code']} ({it['size']}\"): {old_p:,.0f} ➔ {it['price']:,.0f} ({diff:+,.0f}, {pct}%)")
        print()
