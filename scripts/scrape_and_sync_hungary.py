# -*- coding: utf-8 -*-
"""
Full Deep Scraper & Multi-Sheet Excel Sync for Hungary MediaMarkt (mediamarkt.hu)
Matches Czech (Alza) & Greece (Public) standard formatting with 2026 / 2025 year separation.
"""

import sys
import os
import json
import re
import glob
from datetime import datetime
from bs4 import BeautifulSoup
from scrapling.fetchers import StealthySession
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding='utf-8')

# Exchange Rate: 1 EUR = 398 HUF
HUF_TO_EUR_RATE = 398.0

# 1. Advanced Inch Extractor
def extract_size(title_upper, model_code=""):
    m_mc = re.search(r'(?:OLED|QE|UE|GQ|QNED|NANO|NU|UA|DU|CU)?(\d{2,3})(?:C\d|G\d|B\d|M\d|QN|S\d|U\d|Q\d|QNED|NANO|UA|NU|DU|CU|H\d)', model_code.upper())
    if m_mc:
        val = int(m_mc.group(1))
        if 24 <= val <= 115:
            return val

    m_inch = re.search(r'(\d{2,3})\s*(?:\"|\'\'|”|COL|INCH)', title_upper)
    if m_inch:
        val = int(m_inch.group(1))
        if 24 <= val <= 115:
            return val
            
    cm_map = {
        80: 32, 108: 43, 109: 43, 126: 50, 127: 50, 139: 55, 140: 55,
        164: 65, 165: 65, 189: 75, 190: 75, 195: 77, 196: 77, 210: 83,
        215: 85, 216: 85, 218: 86, 248: 98, 249: 98
    }
    m_cm = re.search(r'(\d{2,3})\s*CM\b', title_upper)
    if m_cm:
        cm_val = int(m_cm.group(1))
        if cm_val in cm_map:
            return cm_map[cm_val]

def extract_size_from_title(title_upper):
    # Check explicit inch pattern like 86", 55", 85"
    m_inch = re.search(r'(\d{2,3})\s*["”]', title_upper)
    if m_inch:
        v = int(m_inch.group(1))
        if v in [24, 27, 32, 40, 42, 43, 48, 50, 55, 60, 65, 70, 75, 77, 83, 85, 86, 98, 100, 115]:
            return v
            
    # Check model code start inch e.g., 86MRGB87B3B, 55MRGB87B3B, OLED65C6, QE55S90, MRE85R95
    m_code_size = re.search(r'(?:OLED|QE|UE|GQ|MRE|QNED)?(\d{2,3})(?:MRGB|QNED|NANO|UA|NU|UT|UR|UQ|LQ|LM|G|C|B|M|W|S|QN|Q|U|R)\d', title_upper)
    if m_code_size:
        v = int(m_code_size.group(1))
        if v in [24, 27, 32, 40, 42, 43, 48, 50, 55, 60, 65, 70, 75, 77, 83, 85, 86, 98, 100, 115]:
            return v

    # Fallback to cm mapping
    m_cm = re.search(r'(\d{2,3})\s*CM', title_upper)
    if m_cm:
        cm_val = int(m_cm.group(1))
        cm_map = {
            106: 42, 108: 43, 109: 43, 121: 48, 126: 50, 127: 50,
            139: 55, 140: 55, 164: 65, 165: 65, 189: 75, 190: 75, 195: 77, 196: 77,
            210: 83, 211: 83, 215: 85, 216: 85, 217: 86, 218: 86, 248: 98, 249: 98,
            254: 100, 290: 115, 292: 115
        }
        for k, v in cm_map.items():
            if abs(cm_val - k) <= 2:
                return v
                
    for token in title_upper.replace('"', ' ').replace("'", ' ').split():
        clean = re.sub(r'[^\d]', '', token)
        if clean.isdigit():
            v = int(clean)
            if v in [24, 27, 32, 40, 42, 43, 48, 50, 55, 60, 65, 70, 75, 77, 83, 85, 86, 98, 100, 115]:
                return v
    return 0

# 2. Comprehensive Model & Series Classifier with Global Reference DB Rules
def extract_model_and_series(brand, title_upper):
    brand_upper = brand.upper()
    
    if brand_upper == "LG":
        m_mrgb = re.search(r'(\d{2,3}MRGB[A-Z0-9]+)', title_upper)
        if m_mrgb:
            mc = m_mrgb.group(1)
            sm = re.search(r'MRGB(\d{2}[A-Z]?)', mc)
            series = f"MRGB {sm.group(1)}" if sm else "MRGB"
            return mc, series, "MRGB"

        m_oled = re.search(r'(OLED\d{2,3}[A-Z0-9]+)', title_upper)
        if m_oled:
            mc = m_oled.group(1)
            sm = re.search(r'OLED\d{2,3}([A-Z]\d)', mc)
            series = f"OLED {sm.group(1)}" if sm else "OLED"
            return mc, series, "OLED"
            
        m_qned = re.search(r'(\d{2,3}QNED[A-Z0-9]+)', title_upper)
        if m_qned:
            mc = m_qned.group(1)
            sm = re.search(r'QNED(\d{2,3}[A-Z0-9]*)', mc)
            series = f"QNED {sm.group(1)[:4]}" if sm else "QNED"
            return mc, series, "QNED"

        m_nano = re.search(r'(\d{2,3}NANO[A-Z0-9]+)', title_upper)
        if m_nano:
            mc = m_nano.group(1)
            sm = re.search(r'NANO(\d{2,3}[A-Z0-9]*)', mc)
            series = f"NANO {sm.group(1)[:3]}" if sm else "NanoCell"
            return mc, series, "NanoCell"
            
        m_uhd = re.search(r'(\d{2,3}(?:UA|NU|UT|UR|UQ|LQ|LM|QNED|NANO)[A-Z0-9]+)', title_upper)
        if m_uhd:
            mc = m_uhd.group(1)
            if "LQ" in mc or "LM" in mc:
                return mc, "FHD / HD", "FHD"
            prefix = mc[2:4] if len(mc) >= 4 else "UHD"
            return mc, f"UHD {prefix}", "UHD"
            
        return "Unknown LG", "LG TV", "UHD"
        
    elif brand_upper == "SAMSUNG":
        m_mre = re.search(r'(MRE\d{2,3}[A-Z0-9]+)', title_upper)
        if m_mre:
            mc = m_mre.group(1)
            sm = re.search(r'(R\d{2,3}[A-Z]?)', mc)
            series = f"Micro RGB {sm.group(1)}" if sm else "Micro RGB"
            return mc, series, "MRGB"

        m_full = re.search(r'([QUG][EAN]\d{2,3}[A-Z0-9]+)', title_upper)
        if m_full:
            mc = m_full.group(1)
            
            if any(x in mc for x in ["S95", "S90", "S85", "S99"]):
                sm = re.search(r'(S\d{2}[A-Z]?)', mc)
                series = sm.group(1) if sm else "OLED"
                return mc, f"OLED {series}", "OLED"
                
            if "QN" in mc:
                sm = re.search(r'(QN\d{2,3}[A-Z]?)', mc)
                series = sm.group(1) if sm else "Neo QLED"
                return mc, f"Neo QLED {series}", "Neo QLED"
                
            if "LS03" in mc:
                return mc, "The Frame", "Lifestyle"
                
            if re.search(r'Q\d[A-Z0-9]', mc):
                sm = re.search(r'(Q\d{1,2}[A-Z]?)', mc)
                series = sm.group(1) if sm else "QLED"
                return mc, f"QLED {series}", "QLED"
                
            if any(x in mc for x in ["U8", "U7", "DU", "CU", "AU"]):
                sm = re.search(r'([A-Z]*U\d{4}[A-Z]*)', mc)
                series = sm.group(1) if sm else "Crystal UHD"
                return mc, f"UHD {series}", "UHD"
                
            return mc, "Samsung TV", "UHD"
            
        return "Unknown Samsung", "Samsung TV", "UHD"

    return "Unknown", "Other", "Other"

# 3. Model Year Classifier based on Pan-European Rules
def classify_year(brand, model_code, title_upper):
    brand_upper = brand.upper()
    
    if "2026" in title_upper:
        return 2026
    if "2025" in title_upper:
        return 2025
    if "2024" in title_upper:
        return 2024
        
    if brand_upper == "LG":
        if any(x in model_code for x in [
            "C6", "G6", "B6", "M6", "QNED86B", "QNED80B", "QNED87B", "QNED81B", "QNED7EB", "QNED72B", "QNED71B", "QNED70B",
            "NU8E", "MRGB87B", "MRGB96B", "UA77", "LX7B", "LX6", "27LX6", "STANBYME", "QLED7EB"
        ]):
            return 2026
            
        if any(x in model_code for x in [
            "C5", "G5", "B5", "M5", "QNED86A", "QNED80A", "QNED87A", "QNED7EA", "QNED72A", "QNED70A",
            "NANO82A", "NANO81A", "NANO80A", "UA75", "UA73", "MRGB87A", "LX7A", "LX5", "QNED93A"
        ]):
            return 2025
            
        if any(x in model_code for x in ["C4", "G4", "B4", "M4", "QNED85", "UA70", "UR", "UQ", "UT"]):
            return 2024
            
    elif brand_upper == "SAMSUNG":
        mc_clean = re.sub(r'(?:XXH|XXC|XXN|XXU|XXE|XAU|XEN)$', '', model_code.upper())
        
        # 2026 Series check
        if any(x in mc_clean for x in [
            "S95H", "S90H", "S85H", "S99H",
            "QN900H", "QN800H", "QN95H", "QN90H", "QN85H", "QN80H", "QN70H", "QN72H", "QN82H",
            "M80H", "M70H", "M82H", "M72H", "R95H", "R85H", "R86H",
            "MRE85", "MRE75", "MRE65", "MRE55",
            "LS03H", "LS03HW", "LS03HE", "LS03HA",
            "U8000H", "U8072H", "U8070H", "U8090H", "U8005H", "U8500H",
            "Q70H", "Q80H", "Q60H"
        ]):
            return 2026
            
        # 2025 Series check
        if any(x in mc_clean for x in [
            "S95F", "S90F", "S85F", "S91F", "S92F",
            "QN990F", "QN900F", "QN800F", "QN95F", "QN90F", "QN85F", "QN80F", "QN70F", "QN72F", "QN74F", "QN82F",
            "Q8F", "Q7F", "Q6F", "Q8FA", "Q7FA", "Q6FA", "Q80F", "Q70F", "Q60F",
            "LS03F", "LS03FW", "LS03FA",
            "U8000F", "U8072F", "U8070F", "U8090F", "U8005F", "U7022F", "U7020F", "U7025F", "U7000F",
            "H5002F"
        ]):
            return 2025
            
        # 2024 Series check
        if any(x in mc_clean for x in [
            "S95D", "S90D", "S85D", "QN900D", "QN90D", "QN85D", "QN80D", "QN70D", "Q60D", "LS03D", "DU8000", "DU7000", "CU8000"
        ]):
            return 2024
            
        # General pattern match (excluding country suffix)
        m_yr = re.search(r'(?:S\d{2}|QN\d{2,3}|Q\d{1,2}|M\d{2}|R\d{2}|LS\d{2}|U\d{4}|H\d{4})([HFD])', mc_clean)
        if m_yr:
            letter = m_yr.group(1)
            if letter == 'H': return 2026
            if letter == 'F': return 2025
            if letter == 'D': return 2024
            
        if any(x in title_upper for x in ["S90H", "S95H", "S85H", "S99H", "QN900H", "QN800H", "QN95H", "QN90H", "QN85H", "QN80H", "M80H", "M70H", "R95H", "R85H", "LS03H", "U8000H", "U8090H", "2026"]):
            return 2026
        if any(x in title_upper for x in ["S90F", "S95F", "S85F", "QN900F", "QN800F", "QN95F", "QN90F", "QN85F", "QN80F", "LS03F", "U8000F", "U8072F", "Q7F", "Q8F", "Q6F", "2025"]):
            return 2025
            
    return 2025

# 4. Apollo State Parser
def parse_apollo_state(html):
    idx = html.find('window.__PRELOADED_STATE__')
    if idx == -1:
        return []
    end_idx = html.find('</script>', idx)
    raw = html[idx:end_idx].split('window.__PRELOADED_STATE__ =', 1)[1].strip()
    if raw.endswith(';'):
        raw = raw[:-1]
    raw = re.sub(r':\s*undefined\b', ': null', raw)
    try:
        data = json.loads(raw)
    except Exception as e:
        return []
        
    apollo = data.get("apolloState", {})
    
    # Map price features
    prices = {}
    for k, v in apollo.items():
        if k.startswith("CofrPriceFeature:"):
            media_id = v.get("id", "")
            raw_id = media_id.split(":")[-1] if ":" in media_id else media_id
            p_obj = v.get("price", {})
            amount = p_obj.get("amount") if isinstance(p_obj, dict) else None
            promo_p = v.get("promoPrice", {})
            promo_amount = promo_p.get("amount") if isinstance(promo_p, dict) else None
            
            strike_obj = v.get("strikePrice")
            strike_price = None
            if isinstance(strike_obj, dict):
                strike_price = strike_obj.get("amount")
            elif isinstance(strike_obj, (int, float)):
                strike_price = strike_obj
                
            is_mp = v.get("isProductOfTypeMarketplace", False)
            seller = v.get("marketplaceSeller")
            
            prices[raw_id] = {
                "amount": amount,
                "promo_amount": promo_amount,
                "strike_price": strike_price,
                "currency": v.get("currency", "HUF"),
                "is_marketplace": is_mp,
                "seller": seller
            }
            
    products = []
    for k, v in apollo.items():
        if k.startswith("GraphqlProduct:"):
            prod_id = str(v.get("id") or v.get("productId") or "")
            title = v.get("title") or v.get("name") or v.get("description") or ""
            brand = v.get("manufacturer") or v.get("brand") or ""
            pdp_url = v.get("pdpUrl") or v.get("url") or ""
            
            price_info = prices.get(prod_id, {})
            
            products.append({
                "product_id": prod_id,
                "title": title,
                "brand": brand,
                "pdp_url": f"https://www.mediamarkt.hu{pdp_url}" if pdp_url.startswith("/") else pdp_url,
                "price_huf": price_info.get("amount"),
                "promo_price_huf": price_info.get("promo_amount"),
                "strike_price_huf": price_info.get("strike_price"),
                "is_marketplace": price_info.get("is_marketplace", False),
                "seller": price_info.get("seller") or "MediaMarkt"
            })
            
    return products

# 5. Deep Multi-Query & Category Scraping Engine
def run_deep_scrape_for_brand(session, brand):
    brand_upper = brand.upper()
    print(f"\n=======================================================")
    print(f"🚀 [DEEP SCRAPING] Starting {brand} Deep TV Collection")
    print(f"=======================================================")
    
    if brand_upper == "LG":
        search_targets = [
            ("LG TV Base", "https://www.mediamarkt.hu/hu/search.html?query=LG+TV", 12),
            ("LG OLED", "https://www.mediamarkt.hu/hu/search.html?query=LG+OLED", 8),
            ("LG QNED", "https://www.mediamarkt.hu/hu/search.html?query=LG+QNED", 8),
            ("LG NanoCell", "https://www.mediamarkt.hu/hu/search.html?query=LG+NanoCell", 5),
            ("LG UHD", "https://www.mediamarkt.hu/hu/search.html?query=LG+UHD", 8),
            ("LG 2026", "https://www.mediamarkt.hu/hu/search.html?query=LG+2026", 5),
            ("LG C5 C6 G5 G6", "https://www.mediamarkt.hu/hu/search.html?query=LG+C5+C6+G5+G6", 5),
            ("OLED Category", "https://www.mediamarkt.hu/hu/category/oled-tv-679555.html", 5),
            ("QNED Category", "https://www.mediamarkt.hu/hu/category/qled-8k-4k-tv-649494.html", 5),
            ("4K UHD Category", "https://www.mediamarkt.hu/hu/category/4k-uhd-tv-679556.html", 6),
            ("NanoCell Category", "https://www.mediamarkt.hu/hu/category/nanocell-tv-765211.html", 4),
            ("MiniLED Category", "https://www.mediamarkt.hu/hu/category/miniled-tv-858585.html", 4)
        ]
    else:  # SAMSUNG
        search_targets = [
            ("Samsung TV Base", "https://www.mediamarkt.hu/hu/search.html?query=Samsung+TV", 12),
            ("Samsung OLED", "https://www.mediamarkt.hu/hu/search.html?query=Samsung+OLED", 8),
            ("Samsung Neo QLED", "https://www.mediamarkt.hu/hu/search.html?query=Samsung+Neo+QLED", 8),
            ("Samsung QLED", "https://www.mediamarkt.hu/hu/search.html?query=Samsung+QLED", 8),
            ("Samsung The Frame", "https://www.mediamarkt.hu/hu/search.html?query=Samsung+The+Frame", 5),
            ("Samsung Crystal UHD", "https://www.mediamarkt.hu/hu/search.html?query=Samsung+UHD", 8),
            ("Samsung 2026", "https://www.mediamarkt.hu/hu/search.html?query=Samsung+2026", 5),
            ("Samsung S90 S95 QN90", "https://www.mediamarkt.hu/hu/search.html?query=Samsung+S90+S95+QN90", 5),
            ("OLED Category", "https://www.mediamarkt.hu/hu/category/oled-tv-679555.html", 5),
            ("QLED Category", "https://www.mediamarkt.hu/hu/category/qled-8k-4k-tv-649494.html", 5),
            ("4K UHD Category", "https://www.mediamarkt.hu/hu/category/4k-uhd-tv-679556.html", 6),
            ("MiniLED Category", "https://www.mediamarkt.hu/hu/category/miniled-tv-858585.html", 4)
        ]
        
    collected_dict = {}
    
    for label, base_url, max_p in search_targets:
        print(f"\n  ➔ [TARGET] {label} (Max {max_p} pages): {base_url}")
        for page in range(1, max_p + 1):
            sep = "&" if "?" in base_url else "?"
            url = f"{base_url}{sep}page={page}"
            
            try:
                res = session.fetch(url, solve_cloudflare=True, wait=2500, timeout=35000)
                if res.status != 200:
                    break
                    
                items = parse_apollo_state(str(res.html_content))
                if not items:
                    break
                    
                added = 0
                for it in items:
                    item_b = it["brand"].upper()
                    item_t = it["title"].upper()
                    if brand_upper not in item_b and brand_upper not in item_t:
                        continue
                        
                    pid = it["product_id"]
                    if pid not in collected_dict:
                        collected_dict[pid] = it
                        added += 1
                        
                print(f"    Page {page}: Extracted {len(items)} raw | Added +{added} unique for {brand} (Subtotal: {len(collected_dict)})")
                
                if len(items) == 0:
                    break
                    
            except Exception as e:
                print(f"    ⚠️ Fetch error on {url}: {e}")
                break
                
    print(f"\n✅ Finished Deep Scraping for {brand}. Total Unique Models: {len(collected_dict)}")
    return list(collected_dict.values())

# 6. Data Cleaning & Standard Formatting
def clean_and_standardize(raw_items, brand):
    cleaned = []
    
    for item in raw_items:
        title = item["title"]
        title_upper = title.upper()
        
        if any(x in title_upper for x in ["MONITOR", "ODYSSEY", "ULTRAGEAR", "MYVIEW", "SOUNDBAR", "HIFI", "FALI TARTÓ", "TÁVIRÁNYÍTÓ", "KIEGÉSZÍTŐ", "BRACKET", "HANGFAL"]):
            continue
            
        p_raw = item["price_huf"] or item["promo_price_huf"]
        if not p_raw or not isinstance(p_raw, (int, float)) or p_raw <= 0:
            continue
        price_huf = float(p_raw)
            
        model_code, series, category = extract_model_and_series(brand, title_upper)
        inch = extract_size_from_title(title_upper) or extract_size(title_upper, model_code) or 0
        year = classify_year(brand, model_code, title_upper)
        
        # Strict Model Year filter: Only 2025 and 2026 models are allowed
        if year not in [2025, 2026]:
            continue
        
        price_eur = round(price_huf / HUF_TO_EUR_RATE, 2)
        
        w_raw = item.get("strike_price_huf")
        was_price = 0.0
        if isinstance(w_raw, (int, float)):
            was_price = float(w_raw)
        elif isinstance(w_raw, dict) and "amount" in w_raw:
            was_price = float(w_raw["amount"])
            
        discount_pct = round((was_price - price_huf) / was_price * 100, 1) if was_price > price_huf else 0.0
        
        promo_text = "Special Price" if item.get("promo_price_huf") and float(item["promo_price_huf"]) < price_huf else "Standard"
        if item.get("is_marketplace"):
            promo_text += f" (Marketplace: {item.get('seller', 'Partner')})"
            
        record = {
            "Brand": brand.capitalize(),
            "Model Year": year,
            "Series": series,
            "Size (inch)": int(inch),
            "Model Code": model_code,
            "Selling Price (HUF)": int(price_huf),
            "Selling Price (EUR)": price_eur,
            "Was Price (HUF)": int(was_price) if was_price > 0 else 0,
            "Discount (%)": discount_pct,
            "Promotion": promo_text,
            "Cashback (HUF)": 0,
            "Product Link": item["pdp_url"],
            "Title": title,
            "Category": category,
            "Is Marketplace": item.get("is_marketplace", False),
            "Seller": item.get("seller", "MediaMarkt")
        }
        
        cleaned.append(record)
        
    cleaned.sort(key=lambda x: (-(x["Model Year"] or 0), -(x["Size (inch)"] or 0), -(x["Selling Price (HUF)"] or 0)))
    return cleaned

# 7. Excel Formatter & Sheet Builder
def write_standard_sheet(ws, products, sheet_title):
    headers = [
        "Brand", "Model Year", "Series", "Size (inch)", "Model Code",
        "Selling Price (HUF)", "Selling Price (EUR)", "Was Price (HUF)",
        "Discount (%)", "Promotion", "Cashback (HUF)", "Product Link", "Title"
    ]
    
    header_fill = PatternFill(start_color="051C2C", end_color="051C2C", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    
    font_2026 = Font(name="Segoe UI", size=9.5, bold=True, color="00538C")
    font_2025 = Font(name="Segoe UI", size=9.5, color="1E293B")
    font_normal = Font(name="Segoe UI", size=9.5, color="334155")
    
    fill_2026 = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    ws.append(headers)
    ws.row_dimensions[1].height = 24
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    for row_idx, p in enumerate(products, 2):
        row_data = [
            p["Brand"],
            p["Model Year"],
            p["Series"],
            p["Size (inch)"],
            p["Model Code"],
            p["Selling Price (HUF)"],
            p["Selling Price (EUR)"],
            p["Was Price (HUF)"],
            p["Discount (%)"],
            p["Promotion"],
            p["Cashback (HUF)"],
            p["Product Link"],
            p["Title"]
        ]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 19
        
        is_new_2026 = (p["Model Year"] == 2026)
        
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border_thin
            cell.font = font_2026 if is_new_2026 else (font_2025 if p["Model Year"] == 2025 else font_normal)
            if is_new_2026:
                cell.fill = fill_2026
                
            if col_idx in [1, 2, 3, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [5, 10]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 6:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0 "Ft"'
            elif col_idx == 7:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '€#,##0.00'
            elif col_idx == 8:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0 "Ft"'
            elif col_idx == 9:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '0.0"%"'
            elif col_idx == 11:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0'
            elif col_idx in [12, 13]:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 50)
        
    ws.auto_filter.ref = ws.dimensions

def main():
    data_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
    os.makedirs(data_dir, exist_ok=True)
    raw_cache_file = os.path.join(data_dir, "raw_mediamarkt_hu_deep.json")
    
    print("🌐 [LIVE SCRAPE] Starting live scrape for MediaMarkt Hungary...")
    with StealthySession(headless=True) as session:
        raw_sec = run_deep_scrape_for_brand(session, "Samsung")
        raw_lg = run_deep_scrape_for_brand(session, "LG")
        
    with open(raw_cache_file, "w", encoding="utf-8") as f:
        json.dump({"samsung": raw_sec, "lg": raw_lg}, f, ensure_ascii=False, indent=2)
            
    clean_sec = clean_and_standardize(raw_sec, "Samsung")
    clean_lg = clean_and_standardize(raw_lg, "LG")
    
    print(f"\n=======================================================")
    print(f"📊 [PROCESSED] Cleaned Data Summary:")
    print(f"  • Samsung Total Unique Models: {len(clean_sec)}")
    print(f"    - 2026 Models: {sum(1 for x in clean_sec if x['Model Year'] == 2026)}")
    print(f"    - 2025 Models: {sum(1 for x in clean_sec if x['Model Year'] == 2025)}")
    print(f"    - 2024 Models: {sum(1 for x in clean_sec if x['Model Year'] == 2024)}")
    print(f"  • LG Total Unique Models: {len(clean_lg)}")
    print(f"    - 2026 Models: {sum(1 for x in clean_lg if x['Model Year'] == 2026)}")
    print(f"    - 2025 Models: {sum(1 for x in clean_lg if x['Model Year'] == 2025)}")
    print(f"    - 2024 Models: {sum(1 for x in clean_lg if x['Model Year'] == 2024)}")
    print(f"  • Grand Total: {len(clean_sec) + len(clean_lg)} Models")
    print(f"=======================================================")

    # 1. Save standalone Hungary Excel
    date_str = datetime.now().strftime("%Y %m%d")
    history_eu_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "History_EU", date_str))
    os.makedirs(history_eu_dir, exist_ok=True)
    
    hu_excel_path = os.path.join(history_eu_dir, f"price tracker_HU_{date_str}.xlsx")
    wb_hu = openpyxl.Workbook()
    ws_sec = wb_hu.active
    ws_sec.title = "MediaMarkt_HU_Samsung"
    write_standard_sheet(ws_sec, clean_sec, "MediaMarkt_HU_Samsung")
    
    ws_lg = wb_hu.create_sheet(title="MediaMarkt_HU_LG")
    write_standard_sheet(ws_lg, clean_lg, "MediaMarkt_HU_LG")
    
    wb_hu.save(hu_excel_path)
    print(f"💾 [SAVED] Standalone Hungary Excel saved to: {hu_excel_path}")

    # 2. Merge into Pan-European Excel (price tracker_EU_2026 0814_v1.xlsx)
    eu_excels = glob.glob(os.path.join(history_eu_dir, "price tracker_EU_*.xlsx"))
    if eu_excels:
        target_eu_excel = sorted(eu_excels)[-1]
        print(f"🔄 [MERGING] Merging Hungary sheets into Pan-European Excel: {target_eu_excel}")
        try:
            wb_eu = openpyxl.load_workbook(target_eu_excel)
            
            for sname in ["MediaMarkt_HU_Samsung", "MediaMarkt_HU_LG"]:
                if sname in wb_eu.sheetnames:
                    del wb_eu[sname]
                    
            ws_eu_sec = wb_eu.create_sheet(title="MediaMarkt_HU_Samsung")
            write_standard_sheet(ws_eu_sec, clean_sec, "MediaMarkt_HU_Samsung")
            
            ws_eu_lg = wb_eu.create_sheet(title="MediaMarkt_HU_LG")
            write_standard_sheet(ws_eu_lg, clean_lg, "MediaMarkt_HU_LG")
            
            wb_eu.save(target_eu_excel)
            print(f"✅ [MERGED] Successfully updated Pan-European Excel with Hungary sheets!")
        except Exception as e:
            print(f"⚠️ Error merging into Pan-European Excel: {e}")

    # 3. Save JSON DB
    all_json = clean_sec + clean_lg
    with open(os.path.join(data_dir, "mediamarkt_hu_full.json"), "w", encoding="utf-8") as f:
        json.dump(all_json, f, ensure_ascii=False, indent=2)
    print(f"💾 [SAVED] Full JSON DB saved ({len(all_json)} total records).")

if __name__ == "__main__":
    main()
