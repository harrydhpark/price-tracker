import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('data/price tracker_EU_2026 0901_v1.xlsx', data_only=True)

for sheet in ['Currys_UK_SAMSUNG_Full', 'Currys_UK_LG_Full']:
    ws = wb[sheet]
    print(f'=== {sheet} ===')
    for r in range(2, ws.max_row+1):
        year = ws.cell(r, 2).value
        disp = ws.cell(r, 3).value
        size = ws.cell(r, 4).value
        mc = ws.cell(r, 5).value
        price = ws.cell(r, 6).value
        promo = ws.cell(r, 10).value
        if mc and price:
            try:
                pval = float(price)
                print(f'{str(year):<5} | {str(disp):<10} | {str(size):<3} | {str(mc):<20} | £{pval:>8.2f} | Promo: {promo}')
            except Exception:
                pass
