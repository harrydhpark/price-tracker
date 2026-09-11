# -*- coding: utf-8 -*-
"""
Czech Alza (alza.cz) LG & Samsung TV Price Collector (Full 100% Coverage)
Scrapes 2025 & 2026 TV model listings sold by Alza Czech Republic,
filters for Condition = NEW (Nové), parses model codes, screen sizes,
selling prices, original prices, and promotion details.
Outputs raw JSON files and appends sheets to price tracker_EU_2026 {MMDD}_v1.xlsx.
"""

import sys
import os
import re
import json
import glob
import time
from datetime import datetime
from bs4 import BeautifulSoup
from scrapling.fetchers import StealthySession
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding='utf-8')

DATA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
HISTORY_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "History_EU"))

def extract_model_code_and_info(title_upper, brand, size_val):
    year_val = None
    series_code = "Unknown"
    model_code = "Unknown"

    if brand.lower() == "samsung":
        patterns_2026 = [
            (r'S99H', 'S99H', 2026), (r'S95H', 'S95H', 2026), (r'S90H', 'S90H', 2026), (r'S85H', 'S85H', 2026),
            (r'QN95H', 'QN95H', 2026), (r'QN90H', 'QN90H', 2026), (r'QN85H', 'QN85H', 2026),
            (r'QN80H', 'QN80H', 2026), (r'QN70H', 'QN70H', 2026), (r'M80H', 'M80H', 2026),
            (r'M70H', 'M70H', 2026), (r'R95H', 'R95H', 2026), (r'R85H', 'R85H', 2026), (r'R86H', 'R86H', 2026),
            (r'U8072H', 'U8000H', 2026), (r'U8070H', 'U8000H', 2026), (r'U8092H', 'U8000H', 2026),
            (r'U8072', 'U8000H', 2026), (r'U8092', 'U8000H', 2026),
            (r'U8000H', 'U8000H', 2026), (r'F6000H', 'F6000H', 2026)
        ]
        patterns_2025 = [
            (r'S95F', 'S95F', 2025), (r'S90F', 'S90F', 2025), (r'S85F', 'S85F', 2025),
            (r'QN95F', 'QN95F', 2025), (r'QN90F', 'QN90F', 2025), (r'QN85F', 'QN85F', 2025),
            (r'QN80F', 'QN80F', 2025), (r'QN70F', 'QN70F', 2025), (r'Q80F', 'Q8F', 2025),
            (r'Q8F', 'Q8F', 2025), (r'Q70F', 'Q7F', 2025), (r'Q7F', 'Q7F', 2025),
            (r'Q60F', 'Q6F', 2025), (r'Q6F', 'Q6F', 2025), (r'U8000F', 'U8000F', 2025),
            (r'LS03F', 'The Frame', 2025)
        ]

        m_code = re.search(r'([A-Z]{2}\d{2}[A-Z0-9]+)', title_upper)
        if m_code:
            model_code = m_code.group(1)

        for pat, ser, yr in patterns_2026:
            if re.search(pat, title_upper):
                series_code = ser
                year_val = yr
                break

        if not year_val:
            for pat, ser, yr in patterns_2025:
                if re.search(pat, title_upper):
                    series_code = ser
                    year_val = yr
                    break

        if not year_val:
            if "2026" in title_upper:
                year_val = 2026
            elif "2025" in title_upper:
                year_val = 2025

    elif brand.lower() == "lg":
        lg_2026 = [
            (r'QNED93B', 'QNED93B', 2026),
            (r'QNED87B', 'QNED87B', 2026),
            (r'QNED86B', 'QNED86B', 2026),
            (r'QNED85B', 'QNED85B', 2026),
            (r'QNED82B', 'QNED82B', 2026),
            (r'QNED81B', 'QNED81B', 2026),
            (r'QNED80B', 'QNED80B', 2026),
            (r'QNED72B', 'QNED72B', 2026),
            (r'QNED71B', 'QNED70B', 2026),
            (r'QNED70B', 'QNED70B', 2026),
            (r'QNED7EB', 'QNED7EB', 2026),
            (r'MRGB96B', 'MRGB96B', 2026),
            (r'MRGB87B', 'MRGB87B', 2026),
            (r'MRGB86B', 'MRGB85B', 2026),
            (r'LX7B', 'LX7B', 2026), (r'27LX6', '27LX6', 2026), (r'STANBYME', 'STANBYME', 2026),
            (r'OLED\d{2}G6', 'OLED G6', 2026), (r'OLED.*G6', 'OLED G6', 2026),
            (r'OLED\d{2}C6', 'OLED C6', 2026), (r'OLED.*C6', 'OLED C6', 2026),
            (r'OLED\d{2}B6', 'OLED B6', 2026), (r'OLED.*B6', 'OLED B6', 2026),
            (r'\bG6\b', 'OLED G6', 2026), (r'\bC6\b', 'OLED C6', 2026), (r'\bB6\b', 'OLED B6', 2026),
            (r'NU85', 'NU85', 2026), (r'NU8E', 'NU85', 2026), (r'UA77', 'UA77', 2026)
        ]
        lg_2025 = [
            (r'QNED93A', 'QNED93A', 2025),
            (r'QNED87A', 'QNED86A', 2025),
            (r'QNED86A', 'QNED86A', 2025),
            (r'QNED85A', 'QNED86A', 2025),
            (r'QNED80A', 'QNED80A', 2025),
            (r'QNED70A', 'QNED70A', 2025),
            (r'OLED\d{2}G5', 'OLED G5', 2025), (r'OLED.*G5', 'OLED G5', 2025),
            (r'OLED\d{2}C5', 'OLED C5', 2025), (r'OLED.*C5', 'OLED C5', 2025),
            (r'OLED\d{2}B5', 'OLED B5', 2025), (r'OLED.*B5', 'OLED B5', 2025),
            (r'\bG5\b', 'OLED G5', 2025), (r'\bC5\b', 'OLED C5', 2025), (r'\bB5\b', 'OLED B5', 2025),
            (r'UA75', 'UA75', 2025), (r'UA73', 'UA75', 2025)
        ]

        m_code = re.search(r'(OLED\d{2}[A-Z0-9]+|\d{2}QNED[A-Z0-9]+|\d{2}UA[A-Z0-9]+|\d{2}NU[A-Z0-9]+)', title_upper)
        if m_code:
            model_code = m_code.group(1)
        else:
            m_code2 = re.search(r'LG\s+([A-Z0-9]{5,15})', title_upper)
            if m_code2:
                model_code = m_code2.group(1)

        for pat, ser, yr in lg_2026:
            if re.search(pat, title_upper):
                series_code = ser
                year_val = yr
                break

        if not year_val:
            for pat, ser, yr in lg_2025:
                if re.search(pat, title_upper):
                    series_code = ser
                    year_val = yr
                    break

        if not year_val:
            if "2026" in title_upper:
                year_val = 2026
            elif "2025" in title_upper:
                year_val = 2025

    if model_code == "Unknown":
        if series_code != "Unknown":
            if brand.lower() == "lg":
                model_code = f"{size_val}{series_code.replace('OLED ', '')}"
            else:
                model_code = f"QE{size_val}{series_code}"
        else:
            m_title = re.search(r'(\d{2,3})"\s+(?:LG|Samsung)\s+([A-Za-z0-9-]+)', title_upper)
            if m_title:
                model_code = m_title.group(2)

    return model_code, series_code, year_val

def parse_alza_page(html_body, brand):
    soup = BeautifulSoup(html_body, 'html.parser')
    products = []

    # Select all potential product card containers
    items = soup.select('.browsingitem, .js-box, [data-id], .box')
    
    for item in items:
        item_text = item.get_text(separator=' ', strip=True)
        item_upper = item_text.upper()

        if brand.lower() not in item_text.lower():
            continue
            
        # Condition check: exclude Rozbaleno / Zánovní / Použité (keep Nové)
        if any(unboxed in item_text for unboxed in ["Rozbaleno", "Zánovní", "použité", "Použité"]):
            continue

        link_el = item.select_one('a[href*="-d"]') or item.select_one('.name') or item.select_one('a')
        if not link_el:
            continue
            
        href = link_el.get('href', '')
        if not href or href == '#' or 'javascript:' in href:
            continue
        if not href.startswith('http'):
            href = "https://www.alza.cz" + href if href.startswith('/') else "https://www.alza.cz/" + href

        title = link_el.get_text(strip=True) or link_el.get('title', '')
        
        if len(title) < 5 or title.lower() in ["zobrazit", "detail", "koupit", "doporučujeme"]:
            title_m = re.search(r'(\d{2,3}"\s+(?:LG|Samsung)\s+[A-Za-z0-9\s-]+)', item_text)
            if title_m:
                title = title_m.group(1)
            else:
                title = item_text[:60]

        title_upper = title.upper()

        size_m = re.search(r'(\d{2,3})\s*(?:[″"]|inch|cm|\b)', title)
        size_val = 55
        if size_m:
            try:
                size_candidate = int(size_m.group(1))
                if 22 <= size_candidate <= 100:
                    size_val = size_candidate
            except ValueError:
                pass

        if size_val < 22:
            continue
        if any(m in title_upper for m in ["MONITOR", "ODYSSEY", "ULTRAGEAR", "MYVIEW", "STANDBYME", "SOUNDBAR", "REPRODUKTOR", "SLUCHÁTKA"]):
            continue

        # 1. Selling Price from specific DOM price elements only
        price_el = item.select_one('.js-price-box__primary-price__value, .ads-pb__price-value, .price-box__primary-price, .c2')
        selling_price = 0.0
        if price_el:
            p_text = price_el.get_text(strip=True)
            num_match = re.search(r'(\d[\d\s\xa0]*)', p_text)
            if num_match:
                num_str = num_match.group(1).replace('\xa0', '').replace(' ', '').replace('.', '').replace(',', '')
                if num_str.isdigit():
                    selling_price = float(num_str)

        # Fallback if DOM element selector is missing
        if selling_price < 500.0:
            prices_raw = re.findall(r'(\d[\d\s\xa0]*)(?:,-|Kč)', item_text)
            clean_prices = []
            for pr in prices_raw:
                num_str = pr.replace('\xa0', '').replace(' ', '').replace('.', '').replace(',', '')
                if num_str.isdigit():
                    val = float(num_str)
                    if val >= 500.0 and val != 100000.0: # Purge 100,000 promo text false positive
                        clean_prices.append(val)
            if clean_prices:
                selling_price = clean_prices[0]

        if selling_price < 500.0:
            continue

        # 2. Was Price (Original Price)
        was_price_el = item.select_one('.ads-pb__price-value--orig, .price-box__compare-price')
        was_price = 0.0
        if was_price_el:
            w_text = was_price_el.get_text(strip=True)
            w_match = re.search(r'(\d[\d\s\xa0]*)', w_text)
            if w_match:
                w_str = w_match.group(1).replace('\xa0', '').replace(' ', '').replace('.', '').replace(',', '')
                if w_str.isdigit():
                    val = float(w_str)
                    if val > selling_price:
                        was_price = val

        discount_pct = 0
        if was_price > selling_price:
            discount_pct = int(round((was_price - selling_price) / was_price * 100))

        model_code, series_code, year_val = extract_model_code_and_info(title_upper, brand, size_val)

        if year_val and year_val < 2025:
            continue

        promo_text = "None"
        if "Alza dny" in item_text:
            promo_text = "Alza Days Discount Campaign"
        elif "Doručení zdarma" in item_text:
            promo_text = "Free Delivery"
        elif "Slevový kód" in item_text:
            promo_text = "Voucher Code Available"

        # 3. Cashback Net Price & Cashback Amount Calculation
        cashback_val = 0
        cb_price_el = item.select_one('.coupon-block--cashback .coupon-block__price')
        if cb_price_el:
            cb_p_text = cb_price_el.get_text(strip=True)
            cb_num_match = re.search(r'(\d[\d\s\xa0]*)', cb_p_text)
            if cb_num_match:
                cb_num_str = cb_num_match.group(1).replace('\xa0', '').replace(' ', '').replace('.', '').replace(',', '')
                if cb_num_str.isdigit():
                    cb_net_price = float(cb_num_str)
                    if selling_price > cb_net_price:
                        cashback_val = int(selling_price - cb_net_price)
                        if promo_text == "None":
                            promo_text = f"Cashback CZK {cashback_val:,}"
        else:
            cb_m = re.search(r'Cashback\s*(?:Kč)?\s*(\d[\d\s\xa0]*)', item_text, re.IGNORECASE)
            if cb_m:
                cb_str = cb_m.group(1).replace('\xa0', '').replace(' ', '')
                if cb_str.isdigit() and int(cb_str) != 100000:
                    cashback_val = int(cb_str)

        products.append({
            "brand": brand.capitalize(),
            "year": year_val or 2025,
            "series": series_code,
            "size": size_val,
            "model_code": model_code,
            "price": selling_price,
            "was_price": was_price,
            "discount_pct": discount_pct,
            "promo": promo_text,
            "cashback": cashback_val,
            "title": title,
            "url": href
        })

    return products

def collect_alza_brand(brand):
    print(f"\n[INFO] Starting Alza Czech Republic collection for {brand} (Expanded 12-page coverage)...")
    brand_id = "18862345" if brand.lower() == "lg" else "18862344"
    all_products = []
    seen_urls = set()

    # Target URLs: Search query + Category pages 1 to 12
    target_urls = [
        f"https://www.alza.cz/search.htm?exps={brand.lower()}+tv",
        f"https://www.alza.cz/televize-{brand.lower()}/{brand_id}.htm"
    ] + [f"https://www.alza.cz/televize-{brand.lower()}/{brand_id}-p{p}.htm" for p in range(1, 13)]

    with StealthySession(headless=True) as session:
        for idx, page_url in enumerate(target_urls):
            print(f"[FETCH {idx+1}/{len(target_urls)}] Requesting: {page_url}")
            try:
                res = session.fetch(page_url, wait=2500)
                if res.status != 200:
                    print(f"[WARN] Non-200 status {res.status} on {page_url}")
                    continue
                    
                body = res.body.decode('utf-8', errors='ignore') if isinstance(res.body, bytes) else res.body
                items = parse_alza_page(body, brand)
                
                new_count = 0
                for item in items:
                    if item["url"] not in seen_urls:
                        seen_urls.add(item["url"])
                        all_products.append(item)
                        new_count += 1
                        
                print(f" -> Extracted {len(items)} items ({new_count} unique new). Total: {len(all_products)}")
                    
            except Exception as e:
                print(f"[ERROR] Failed to fetch {page_url}: {e}")

    return all_products

def save_raw_json(data, brand):
    os.makedirs(DATA_DIR, exist_ok=True)
    filename = f"raw_alza_{brand.lower()}.json"
    filepath = os.path.join(DATA_DIR, filename)
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"[SAVE] Raw JSON saved: {filepath} ({len(data)} items)")

def update_eu_excel_tracker(lg_items, samsung_items):
    today_mmdd = datetime.now().strftime("%m%d")
    target_wb_path = os.path.join(DATA_DIR, f"price tracker_EU_2026 {today_mmdd}_v1.xlsx")
    
    if not os.path.exists(target_wb_path):
        existing = glob.glob(os.path.join(DATA_DIR, "price tracker_EU_2026 *.xlsx"))
        existing = [f for f in existing if not os.path.basename(f).startswith("~$")]
        if existing:
            existing.sort()
            import shutil
            shutil.copy2(existing[-1], target_wb_path)
            print(f"[CLONE] Created today's workbook from {existing[-1]} -> {target_wb_path}")
        else:
            print("[ERROR] No EU price tracker template found!")
            return

    print(f"[EXCEL] Target workbook: {target_wb_path}")

    wb = None
    for attempt in range(1, 10):
        try:
            wb = openpyxl.load_workbook(target_wb_path)
            break
        except PermissionError:
            print(f"[WARN] Excel file locked. Generating new version suffix _v{attempt+1}...")
            base, ext = os.path.splitext(target_wb_path)
            target_wb_path = f"{base.rsplit('_v', 1)[0]}_v{attempt+1}{ext}"

    if not wb:
        print("[ERROR] Could not open Excel workbook after multiple attempts.")
        return

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    data_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = [
        "Brand", "Model Year", "Series", "Size (inch)", "Model Code",
        "Selling Price (CZK)", "Was Price (CZK)", "Discount (%)", "Promotion",
        "Cashback (CZK)", "Product Link"
    ]

    sheets_data = [
        ("Alza_CZ_Samsung", samsung_items),
        ("Alza_CZ_LG", lg_items)
    ]

    for sheet_name, items in sheets_data:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True
        
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for p in items:
            ws.append([
                p.get("brand", ""),
                p.get("year", 2025),
                p.get("series", ""),
                p.get("size", 55),
                p.get("model_code", ""),
                p.get("price", 0.0),
                p.get("was_price", 0.0),
                p.get("discount_pct", 0),
                p.get("promo", "None"),
                p.get("cashback", 0),
                p.get("url", "")
            ])

        for r in range(2, ws.max_row + 1):
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = data_font
                cell.border = thin_border
                if c in [2, 4, 8, 10]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c in [6, 7]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0"

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        print(f"[EXCEL] Written {len(items)} rows into sheet '{sheet_name}'")

    wb.save(target_wb_path)
    print(f"[SUCCESS] Excel successfully saved to: {target_wb_path}")

    try:
        today_folder = datetime.now().strftime("%Y %m%d")
        hist_folder_path = os.path.join(HISTORY_DIR, today_folder)
        os.makedirs(hist_folder_path, exist_ok=True)
        hist_wb_path = os.path.join(hist_folder_path, os.path.basename(target_wb_path))
        wb.save(hist_wb_path)
        print(f"[MIRROR] Saved backup to History_EU: {hist_wb_path}")
    except Exception as e:
        print(f"[WARN] Failed to mirror history workbook: {e}")

def main():
    import time
    start_time = time.time()
    print("=" * 60)
    print(" CZECH REPUBLIC ALZA (alza.cz) TV PRICE SURVEY COLLECTOR ")
    print(" Target Brands: LG & Samsung | Target Model Years: 2025 & 2026")
    print("=" * 60)

    lg_products = collect_alza_brand("LG")
    save_raw_json(lg_products, "LG")

    samsung_products = collect_alza_brand("Samsung")
    save_raw_json(samsung_products, "Samsung")

    update_eu_excel_tracker(lg_products, samsung_products)

    print("\n[COMPLETE] Alza Czech Republic price collection finished cleanly!")

    try:
        from scrape_logger import emit_summary
        elapsed = round(time.time() - start_time, 2)
        emit_summary(country="CZ", retailer="Alza", brand="LG", total_extracted=len(lg_products), final_deduplicated=len(lg_products), execution_time_sec=elapsed)
        emit_summary(country="CZ", retailer="Alza", brand="SAMSUNG", total_extracted=len(samsung_products), final_deduplicated=len(samsung_products), execution_time_sec=elapsed)
    except Exception:
        pass

if __name__ == "__main__":
    main()
