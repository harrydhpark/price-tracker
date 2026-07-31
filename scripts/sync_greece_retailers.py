# -*- coding: utf-8 -*-
import sys
import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import re
import json
import glob
import shutil
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

GREECE_SOURCES = {
    'kotsovolos': {'retailer': 'Kotsovolos', 'cc': 'GR', 'currency': 'EUR'},
    'public': {'retailer': 'Public', 'cc': 'GR', 'currency': 'EUR'},
}

def save_workbook_with_retry(wb, filename):
    base, ext = os.path.splitext(filename)
    version = 1
    current_filename = filename
    saved = False
    
    while not saved:
        try:
            wb.save(current_filename)
            saved = True
            print(f"  ➔ [SAVED] Successfully saved to: {current_filename}")
        except PermissionError:
            v_match = re.search(r'(_v|v)(\d+)$', base)
            if v_match:
                prefix = v_match.group(1)
                num = int(v_match.group(2)) + 1
                base_clean = base[:v_match.start()]
                current_filename = f"{base_clean}{prefix}{num}{ext}"
                base = f"{base_clean}{prefix}{num}"
            else:
                current_filename = f"{base}_v{version}{ext}"
                version += 1
            print(f"  ➔ [LOCKED] Permission denied. Retrying with: {current_filename}...")
            
    return current_filename

def classify_year(brand, model_code, title_upper):
    year_val = None  # Default value is always None
    
    # Samsung rules
    if brand.upper() == "SAMSUNG":
        if model_code != "Unknown" and len(model_code) > 3:
            sub = model_code[2:]
            if "H" in sub:
                year_val = 2026
            elif "F" in sub:
                year_val = 2025
            elif "D" in sub or "E" in sub:
                year_val = 2024
        
        # Fallback to title keywords
        if year_val is None:
            if any(x in title_upper for x in ["S90H", "S95H", "S85H", "S99H", "QN900H", "QN800H", "QN95H", "QN90H", "QN85H", "QN80H", "QN70H", "LS03H", "U8000H", "U8090H", "R85H", "R95H", "2026"]):
                year_val = 2026
            elif any(x in title_upper for x in ["S90F", "S95F", "S85F", "S99F", "QN900F", "QN800F", "QN95F", "QN90F", "QN85F", "QN80F", "QN70F", "LS03F", "U8000F", "U8090F", "Q7F", "Q8F", "2025"]):
                year_val = 2025
            elif any(x in title_upper for x in ["S90D", "S95D", "S85D", "QN900D", "QN800D", "QN95D", "QN90D", "QN85D", "QN80D", "QN70D", "LS03D", "U8000D", "U8090D", "2024"]):
                year_val = 2024

    # LG rules
    elif brand.upper() == "LG":
        if model_code != "Unknown":
            if any(x in model_code for x in ["C6", "G6", "B6", "W6", "M6", "QNED86B", "QNED80B", "QNED87B", "QNED81B", "QNED72B", "QNED7EB", "UA77", "MRGB87B", "MRGB87", "LX7B", "LX6", "QLED7EB", "MRGB96B", "NU850", "LB700"]):
                year_val = 2026
            elif any(x in model_code for x in ["C5", "G5", "B5", "W5", "M5", "QNED86A", "QNED80A", "QNED87A", "QNED87", "QNED72A", "QNED7EA", "UA75", "MRGB87A", "LX7A", "LX5", "QNED70A", "NANO81A", "NANO81", "NANO80A", "QNED93A"]):
                year_val = 2025
            elif any(x in model_code for x in ["C4", "G4", "B4", "M4", "UA73", "LX4"]):
                year_val = 2024
        
        if year_val is None:
            if any(x in title_upper for x in ["C6", "G6", "B6", "W6", "M6", "QNED86B", "QNED80B", "QNED87B", "QNED72B", "QNED7EB", "UA77", "2026"]):
                year_val = 2026
            elif any(x in title_upper for x in ["C5", "G5", "B5", "W5", "M5", "QNED86A", "QNED80A", "QNED87A", "QNED72A", "QNED7EA", "UA75", "2025"]):
                year_val = 2025
            elif any(x in title_upper for x in ["C4", "G4", "B4", "M4", "2024"]):
                year_val = 2024
                
        if model_code != "Unknown" and "UA73" in model_code and "2025" in title_upper:
            year_val = 2025
            
    return year_val

def parse_product_name(title, brand):
    title_upper = title.upper()
    model_code = "Unknown"
    
    if brand.upper() == "LG":
        oled_match = re.search(r'\b(OLED\d{2,3}[A-Z]{1,3}\d{0,2}[A-Z]*)\b', title_upper)
        if oled_match:
            model_code = oled_match.group(1)
        else:
            lg_code_match = re.search(r'\b(\d{2,3}(?:QNED|NANO|MRGB|LX|NU|UA|UT|UR|UQ|LB|QLED)\w*)\b', title_upper)
            if lg_code_match:
                model_code = lg_code_match.group(1)
            else:
                words = title_upper.split()
                for w in words:
                    clean_w = re.sub(r'[^\w-]', '', w)
                    if any(junk in clean_w for junk in ["LEDLG", "OLEDLG", "TV", "CINEMA"]):
                        continue
                    if any(char.isdigit() for char in clean_w) and len(clean_w) >= 5:
                        model_code = clean_w
                        break
    else: # SAMSUNG
        words = title_upper.split()
        for w in words:
            clean_w = w.replace("(", "").replace(")", "").replace(",", "").replace("\"", "").strip()
            if any(char.isdigit() for char in clean_w) and len(clean_w) >= 6:
                if any(clean_w.startswith(pre) for pre in ["GQ", "TQ", "QE", "UE", "GU", "MRE"]):
                    model_code = clean_w
                    break
        if model_code == "Unknown":
            for w in words:
                clean_w = w.replace("(", "").replace(")", "").replace(",", "").replace("\"", "").strip()
                if any(char.isdigit() for char in clean_w) and len(clean_w) >= 6:
                    model_code = clean_w
                    break

    # Screen Size
    size_val = 55
    size_match = re.search(r'(\d+)\s*"', title)
    if not size_match:
        size_match = re.search(r'(\d+)\s*(?:inch|인치|″)', title, re.IGNORECASE)
    if not size_match:
        size_match = re.search(r'(\d{2,3})\s*cm', title, re.IGNORECASE)
        if size_match:
            cm_val = int(size_match.group(1))
            if cm_val > 100:
                size_val = int(round(cm_val / 2.54))
            elif cm_val >= 22:
                size_val = cm_val
            size_match = True
            
    if isinstance(size_match, re.Match):
        size_val = int(size_match.group(1))
        
    if model_code != "Unknown":
        code_nums = re.findall(r'\d+', model_code)
        if code_nums and len(code_nums[0]) in [2, 3]:
            parsed_code_size = int(code_nums[0])
            if parsed_code_size in [42, 43, 48, 50, 55, 65, 75, 77, 83, 85, 86, 98, 100]:
                size_val = parsed_code_size
            
    # Year
    year_val = classify_year(brand, model_code, title_upper)
    
    # Display Type
    display_type = "LED"
    code_upper = model_code.upper()
    if brand.upper() == "SAMSUNG":
        if "OLED" in code_upper or any(x in code_upper for x in ["S90", "S91", "S92", "S93", "S94", "S95", "S99", "S85", "S83", "S84", "S86", "S89"]) or bool(re.search(r'S(?:9[0-9]|8[0-9])', code_upper)):
            display_type = "OLED"
        elif "MRE" in code_upper or "MRGB" in code_upper:
            display_type = "Micro RGB"
        elif "QN" in code_upper:
            display_type = "Neo QLED"
        elif "Q" in code_upper or "LS" in code_upper:
            display_type = "QLED"
        elif any(x in code_upper for x in ["U8", "UA", "UT", "UR", "UQ", "UX"]):
            display_type = "UHD 4K"
    elif brand.upper() == "LG":
        if "MRGB" in code_upper or "MRGB" in title_upper:
            display_type = "Micro RGB"
        elif "QNED" in code_upper or "QNED" in title_upper:
            display_type = "QNED"
        elif "OLED" in code_upper or "OLED" in title_upper or re.search(r'OLED\d{2}', code_upper):
            display_type = "OLED evo" if ("EVO" in title_upper or any(x in code_upper for x in ["G5", "G6", "C5", "C6", "M5", "M6"])) else "OLED"
        elif "NANO" in code_upper or "NANO" in title_upper:
            display_type = "NanoCell"
        elif any(x in code_upper for x in ["NU900", "NU90", "NU85", "NU80", "NU800", "NU850", "LB650", "LB", "UA75", "UA73", "UA77", "UT", "UR", "UQ", "LQ"]):
            display_type = "UHD 4K"
        elif "LX" in code_upper or "STANBYME" in code_upper:
            display_type = "Lifestyle"

    return model_code, size_val, year_val, display_type

def clean_numeric_price(price_str):
    if price_str is None or price_str == "":
        return 0.0
    if isinstance(price_str, (int, float)):
        val = float(price_str)
        return val if val >= 50.0 else 0.0
        
    clean_p = str(price_str).replace("€", "").replace("EUR", "").strip()
    if re.match(r'^\d{1,3}\.\d{3}$', clean_p):
        clean_p = clean_p.replace(".", "")
    elif re.match(r'^\d{1,3},\d{3}$', clean_p):
        clean_p = clean_p.replace(",", "")
    elif "," in clean_p and "." in clean_p:
        if clean_p.find(".") < clean_p.find(","):
            clean_p = clean_p.replace(".", "").replace(",", ".")
        else:
            clean_p = clean_p.replace(",", "")
    elif "," in clean_p:
        parts = clean_p.split(",")
        if len(parts[-1]) == 2:
            clean_p = clean_p.replace(",", ".")
        else:
            clean_p = clean_p.replace(",", "")
            
    clean_p = re.sub(r'[^\d.]', '', clean_p)
    try:
        val = float(clean_p) if clean_p else 0.0
        return val if val >= 50.0 else 0.0
    except ValueError:
        return 0.0

def sync_greece_sheet(filename, sheetname, brand, products, currency):
    if not os.path.exists(filename):
        wb = openpyxl.Workbook()
        default_sheet = wb.active
        wb.remove(default_sheet)
    else:
        wb = openpyxl.load_workbook(filename)
        
    print(f"  ➔ Syncing to sheet: {sheetname} ({currency})...")
    
    if sheetname in wb.sheetnames:
        sheet = wb[sheetname]
    else:
        sheet = wb.create_sheet(sheetname)
        
    headers = [
        "Brand", "Model Year", "Display Type", "Size (Inch)", "Model Code",
        f"Price ({currency})", "Shipping", "Installment", "Cashback", "General Promotions"
    ]
    if sheet.max_row <= 1:
        sheet.append(headers)
        
    removed_count = 0
    for r in range(sheet.max_row, 1, -1):
        brand_val = sheet.cell(row=r, column=1).value
        if str(brand_val).upper() == str(brand).upper():
            sheet.delete_rows(r)
            removed_count += 1
    if removed_count > 0:
        print(f"    ➔ Removed {removed_count} obsolete {brand} rows.")
        
    added_count = 0
    for p in products:
        cashback_val = p.get("cashback", 0)
        promo_text = p.get("promo", "None") or "None"
        final_price = clean_numeric_price(p["price"])
        
        sheet.append([
            p["brand"],
            p["year"],
            p["display_type"],
            p["size"],
            p["model_code"],
            final_price,
            "Free",
            None,
            cashback_val,
            promo_text
        ])
        added_count += 1
        
    print(f"    ➔ Added {added_count} live {brand} rows to {sheetname}.")
    final_filename = save_workbook_with_retry(wb, filename)
    wb.close()
    return final_filename

def main():
    data_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
    if not os.path.exists(data_dir):
        os.makedirs(data_dir)
        
    today_mmdd = datetime.now().strftime("%m%d")
    excel_filename = os.path.join(data_dir, f"price tracker_greece_2026 {today_mmdd}_v1.xlsx")
    
    latest_files = glob.glob(os.path.join(data_dir, f"price tracker_greece_2026 {today_mmdd}*.xlsx"))
    latest_files = [f for f in latest_files if not os.path.basename(f).startswith("~$")]
    if latest_files:
        latest_files.sort()
        excel_filename = latest_files[-1]
        print(f"[SYNC] Found existing workbook: {excel_filename}")
    else:
        print(f"[SYNC] Creating new workbook: {excel_filename}")

    for src_key, cfg in GREECE_SOURCES.items():
        retailer_name = cfg['retailer']
        for brand in ["samsung", "lg"]:
            sheetname = f"{retailer_name}_{brand.capitalize()}"
            pattern = os.path.join(data_dir, f"raw_{src_key}_{brand}*.json")
            json_files = glob.glob(pattern)
            
            scraped_items = []
            for jf in json_files:
                try:
                    with open(jf, "r", encoding="utf-8") as f:
                        data = json.load(f)
                        if isinstance(data, list):
                            scraped_items.extend(data)
                except Exception as e:
                    print(f"[WARN] Failed to load JSON {jf}: {e}")
                    
            if not scraped_items:
                print(f"[SKIP] No scraped items for {sheetname}")
                continue
                
            processed_products = []
            seen_codes = {}
            
            for item in scraped_items:
                title = item.get("title") or item.get("name") or ""
                raw_price = item.get("price")
                promo = item.get("promo") or "None"
                
                if not title or len(title) < 8:
                    continue
                    
                title_upper = title.upper()
                if any(x in title_upper for x in ["MONITOR", "ODYSSEY", "ULTRAGEAR", "MYVIEW", "SOUNDBAR", "HIFI", "PROJECTOR"]):
                    continue
                if any(x in title_upper for x in ["REFURBISHED", "USED", "GEBRAUCHT"]):
                    continue
                    
                model_code, size, year, display_type = parse_product_name(title, brand)
                
                # Model year strict filtering: 2025 & 2026 only
                if year not in [2025, 2026]:
                    continue
                    
                # Size threshold: >= 22 inches
                if size < 22:
                    continue
                    
                price_val = clean_numeric_price(raw_price)
                if price_val < 50.0:
                    continue
                    
                prod_obj = {
                    "brand": brand.upper(),
                    "year": year,
                    "display_type": display_type,
                    "size": size,
                    "model_code": model_code,
                    "price": price_val,
                    "cashback": item.get("cashback", 0),
                    "promo": promo,
                    "url": item.get("url", "")
                }
                
                # Lowest price deduplication per model code
                code_key = model_code.upper() if model_code != "Unknown" else title_upper
                if code_key not in seen_codes:
                    seen_codes[code_key] = prod_obj
                else:
                    if price_val < seen_codes[code_key]["price"]:
                        seen_codes[code_key] = prod_obj
                        
            processed_products = list(seen_codes.values())
            processed_products.sort(key=lambda x: (x["year"], x["display_type"], -x["size"]))
            
            sync_greece_sheet(excel_filename, sheetname, brand.upper(), processed_products, cfg['currency'])

    print("[SUCCESS] Greece retailer sheets synchronization finished.")

if __name__ == '__main__':
    main()
