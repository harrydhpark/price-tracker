# -*- coding: utf-8 -*-
"""
Australia TV Retailers Excel Synchronization Engine
Syncs raw JSON data to 'data/price tracker_AU_2026 {MMDD}_v1.xlsx'
and mirrors to History_AU/{YYYY MMDD}/
"""

import os, sys, json, re, glob, shutil
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DATA_DIR = os.path.join(ROOT_DIR, "data")
HISTORY_DIR = os.path.join(ROOT_DIR, "History_AU")

def get_survey_date_str():
    return datetime.now().strftime("%m%d")

def check_preflight_gate(raw_files):
    print("=" * 65)
    print(" 🛡️ MANDATORY PRE-FLIGHT LIVE ASSERTION GATE (AUSTRALIA) ")
    print("=" * 65)
    today_str = datetime.now().strftime("%Y-%m-%d")
    for fpath in raw_files:
        if not os.path.exists(fpath):
            raise RuntimeError(f"❌ PRE-FLIGHT GATE FAILED: Missing dataset {fpath}")
        mtime = datetime.fromtimestamp(os.path.getmtime(fpath)).strftime("%Y-%m-%d")
        if mtime != today_str:
            raise RuntimeError(f"❌ PRE-FLIGHT GATE FAILED: {os.path.basename(fpath)} is stale ({mtime} != {today_str})")
        with open(fpath, 'r', encoding='utf-8') as f:
            items = json.load(f)
        if len(items) == 0:
            raise RuntimeError(f"❌ PRE-FLIGHT GATE FAILED: {os.path.basename(fpath)} has 0 items")
        print(f"✅ {os.path.basename(fpath)} verified fresh ({len(items)} items, modified: {mtime})")
    print("All Australian raw datasets verified fresh for today!\n")

def get_category(model_code, brand):
    mc = str(model_code).upper()
    if any(x in mc for x in ["MRGB", "MICRO RGB", "R85", "R95"]):
        return "MRGB"
    elif any(x in mc for x in ["OLED", "S90", "S95", "S85", "S99", "C6", "G6", "B6", "C5", "G5", "B5"]):
        return "OLED"
    elif any(x in mc for x in ["QNED"]):
        return "QNED"
    elif any(x in mc for x in ["QN", "QLED", "NEO QLED", "M70", "M80", "Q7", "Q8", "Q6", "Q5"]):
        return "QLED"
    elif any(x in mc for x in ["LS03", "THE FRAME", "LX"]):
        return "LIFESTYLE"
    else:
        return "UHD"

def sync_au():
    mmdd = get_survey_date_str()
    wb_filename = f"price tracker_AU_2026 {mmdd}_v1.xlsx"
    wb_path = os.path.join(DATA_DIR, wb_filename)
    
    raw_files = [
        os.path.join(DATA_DIR, "raw_jbhifi_samsung.json"),
        os.path.join(DATA_DIR, "raw_jbhifi_lg.json"),
        os.path.join(DATA_DIR, "raw_goodguys_samsung.json"),
        os.path.join(DATA_DIR, "raw_goodguys_lg.json"),
        os.path.join(DATA_DIR, "raw_harveynorman_samsung.json"),
        os.path.join(DATA_DIR, "raw_harveynorman_lg.json")
    ]
    
    check_preflight_gate(raw_files)
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    # Styles
    hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    headers = ["Brand", "Year", "Category", "Size (Inch)", "Model Code", "Price (AUD)", "Original Price", "Currency", "Retailer", "Promotion", "Product Name", "URL"]
    
    sheet_configs = [
        ("JBHIFI_Samsung_Full", "raw_jbhifi_samsung.json", "JB Hi-Fi", "SAMSUNG"),
        ("JBHIFI_LG_Full", "raw_jbhifi_lg.json", "JB Hi-Fi", "LG"),
        ("TheGoodGuys_Samsung_Full", "raw_goodguys_samsung.json", "The Good Guys", "SAMSUNG"),
        ("TheGoodGuys_LG_Full", "raw_goodguys_lg.json", "The Good Guys", "LG"),
        ("HarveyNorman_Samsung_Full", "raw_harveynorman_samsung.json", "Harvey Norman", "SAMSUNG"),
        ("HarveyNorman_LG_Full", "raw_harveynorman_lg.json", "Harvey Norman", "LG")
    ]
    
    all_models_by_year = {2026: [], 2025: []}
    
    for s_name, raw_file, ret_name, brand in sheet_configs:
        ws = wb.create_sheet(title=s_name)
        fpath = os.path.join(DATA_DIR, raw_file)
        with open(fpath, 'r', encoding='utf-8') as f:
            items = json.load(f)
            
        # Write headers
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = center_align
            
        # Write data rows
        row_idx = 2
        for it in sorted(items, key=lambda x: (-x.get('year', 2026), -x.get('size', 0), x.get('price', 0))):
            yr = it.get('year', 2026)
            size = it.get('size', 55)
            mc = it.get('model_code', '')
            price = it.get('price', 0)
            orig_p = it.get('original_price', price)
            cat = get_category(mc, brand)
            title = it.get('name') or it.get('title') or ''
            url = it.get('url', '')
            promo = it.get('promo', 'None')
            
            row = [brand, yr, cat, size, mc, price, orig_p, "AUD", ret_name, promo, title, url]
            ws.append(row)
            
            # Format row
            ws.cell(row=row_idx, column=1).alignment = center_align
            ws.cell(row=row_idx, column=2).alignment = center_align
            ws.cell(row=row_idx, column=3).alignment = center_align
            ws.cell(row=row_idx, column=4).alignment = center_align
            ws.cell(row=row_idx, column=5).alignment = left_align
            
            p_cell = ws.cell(row=row_idx, column=6)
            p_cell.number_format = '$#,##0'
            p_cell.alignment = right_align
            
            op_cell = ws.cell(row=row_idx, column=7)
            op_cell.number_format = '$#,##0'
            op_cell.alignment = right_align
            
            ws.cell(row=row_idx, column=8).alignment = center_align
            ws.cell(row=row_idx, column=9).alignment = center_align
            
            for c in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=c).border = thin_border
                
            if yr in all_models_by_year:
                all_models_by_year[yr].append({
                    'brand': brand, 'retailer': ret_name, 'size': size, 'category': cat,
                    'model_code': mc, 'price': price, 'orig_price': orig_p
                })
            row_idx += 1
            
        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        print(f"  ➔ Written {len(items)} rows into sheet '{s_name}'")
        
    # Create Summary Sheets
    for year_target in [2026, 2025]:
        s_ws = wb.create_sheet(title=f"Summary_{year_target}", index=0 if year_target == 2026 else 1)
        s_headers = ["Category", "Size (Inch)", "LG Model", "LG Price (AUD)", "Samsung Model", "Samsung Price (AUD)", "Gap (AUD)", "Gap (%)", "Retailer"]
        s_ws.append(s_headers)
        
        for c_idx in range(1, len(s_headers) + 1):
            cell = s_ws.cell(row=1, column=c_idx)
            cell.font = hdr_font
            cell.fill = PatternFill(start_color="0D5C75" if year_target == 2026 else "4A6572", end_color="0D5C75" if year_target == 2026 else "4A6572", fill_type="solid")
            cell.alignment = center_align
            
        s_models = all_models_by_year[year_target]
        lg_m = [m for m in s_models if m['brand'] == 'LG']
        sam_m = [m for m in s_models if m['brand'] == 'SAMSUNG']
        
        s_row_idx = 2
        for ret in ["JB Hi-Fi", "The Good Guys", "Harvey Norman"]:
            ret_lg = [m for m in lg_m if m['retailer'] == ret]
            ret_sam = [m for m in sam_m if m['retailer'] == ret]
            
            for l in ret_lg:
                # Find matching Samsung
                s_match = next((s for s in ret_sam if s['size'] == l['size'] and s['category'] == l['category']), None)
                if not s_match:
                    s_match = next((s for s in ret_sam if s['size'] == l['size']), None)
                    
                s_mc = s_match['model_code'] if s_match else "-"
                s_price = s_match['price'] if s_match else 0
                gap = (l['price'] - s_price) if s_price > 0 else 0
                gap_pct = (gap / s_price * 100) if s_price > 0 else 0
                
                s_row = [l['category'], l['size'], l['model_code'], l['price'], s_mc, s_price if s_price > 0 else "-", gap if s_price > 0 else "-", f"{gap_pct:+.1f}%" if s_price > 0 else "-", ret]
                s_ws.append(s_row)
                
                s_ws.cell(row=s_row_idx, column=1).alignment = center_align
                s_ws.cell(row=s_row_idx, column=2).alignment = center_align
                s_ws.cell(row=s_row_idx, column=3).alignment = left_align
                s_ws.cell(row=s_row_idx, column=4).number_format = '$#,##0'
                s_ws.cell(row=s_row_idx, column=5).alignment = left_align
                if s_price > 0:
                    s_ws.cell(row=s_row_idx, column=6).number_format = '$#,##0'
                    s_ws.cell(row=s_row_idx, column=7).number_format = '$#,##0'
                s_ws.cell(row=s_row_idx, column=9).alignment = center_align
                
                for c in range(1, len(s_headers) + 1):
                    s_ws.cell(row=s_row_idx, column=c).border = thin_border
                s_row_idx += 1
                
        for col in s_ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            s_ws.column_dimensions[col_letter].width = max(max_len + 3, 14)
            
    # Save workbook
    wb.save(wb_path)
    print(f"\n✅ [EXCEL SAVED] {wb_path}")
    
    # Mirror to History_AU
    hist_date_dir = os.path.join(HISTORY_DIR, f"{datetime.now().strftime('%Y')} {mmdd}")
    os.makedirs(hist_date_dir, exist_ok=True)
    hist_path = os.path.join(hist_date_dir, wb_filename)
    shutil.copy2(wb_path, hist_path)
    print(f"✅ [MIRRORED TO HISTORY_AU] {hist_path}")

if __name__ == '__main__':
    sync_au()
