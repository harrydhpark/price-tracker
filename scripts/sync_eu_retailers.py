# -*- coding: utf-8 -*-
import sys
import os
import openpyxl
from openpyxl.styles import Font, Alignment
import re
import json
import glob
import shutil

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding='utf-8')

# Source configurations
EU_SOURCES = {
    'mm-de': {'retailer': 'MediaMarkt', 'cc': 'DE', 'currency': 'EUR'},
    'mm-at': {'retailer': 'MediaMarkt', 'cc': 'AT', 'currency': 'EUR'},
    'mm-ch': {'retailer': 'MediaMarkt', 'cc': 'CH', 'currency': 'CHF'},
    'mm-es': {'retailer': 'MediaMarkt', 'cc': 'ES', 'currency': 'EUR'},
    'mm-nl': {'retailer': 'MediaMarkt', 'cc': 'NL', 'currency': 'EUR'},
    'currys': {'retailer': 'Currys', 'cc': 'UK', 'currency': 'GBP'},
    'fnac': {'retailer': 'Fnac', 'cc': 'FR', 'currency': 'EUR'},
    'mw-it': {'retailer': 'MediaWorld', 'cc': 'IT', 'currency': 'EUR'},
}

def save_workbook_with_retry(wb, filename):
    base, ext = os.path.splitext(filename)
    version = 1
    current_filename = filename
    saved = False
    
    while not saved:
        try:
            wb.save(current_filename)
            saved = True
            print(f"  ➔ [SAVED] Successfully saved to: {current_filename}")
        except PermissionError:
            v_match = re.search(r'(_v|v)(\d+)$', base)
            if v_match:
                prefix = v_match.group(1)
                num = int(v_match.group(2)) + 1
                base_clean = base[:v_match.start()]
                current_filename = f"{base_clean}{prefix}{num}{ext}"
                base = f"{base_clean}{prefix}{num}"
            else:
                current_filename = f"{base}_v{version}{ext}"
                version += 1
            print(f"  ➔ [LOCKED] Permission denied. Retrying with: {current_filename}...")
            
    return current_filename

def classify_year(brand, model_code, title_upper):
    year_val = None  # Default value is always None
    
    # Samsung rules
    if brand.upper() == "SAMSUNG":
        if model_code != "Unknown" and len(model_code) > 3:
            sub = model_code[2:]
            if "H" in sub:
                year_val = 2026
            elif "F" in sub:
                year_val = 2025
            elif "D" in sub or "E" in sub:
                year_val = 2024
        
        # Fallback to title keywords
        if year_val is None:
            if any(x in title_upper for x in ["S90H", "S92H", "S95H", "S85H", "S99H", "QN900H", "QN800H", "QN95H", "QN90H", "QN85H", "QN80H", "QN70H", "QN72H", "QN82H", "M70H", "M72H", "M80H", "M82H", "R95H", "R85H", "R86H", "LS03H", "U8000H", "U8070H", "U8079H", "U8075H", "U8090H", "2026"]):
                year_val = 2026
            elif any(x in title_upper for x in ["S90F", "S92F", "S95F", "S85F", "QN900F", "QN800F", "QN95F", "QN90F", "QN85F", "QN80F", "QN70F", "QN72F", "QN82F", "M70F", "M72F", "M80F", "M82F", "R95F", "LS03F", "U8000F", "U8070F", "U8079F", "U8075F", "U8090F", "2025"]):
                year_val = 2025
            elif any(x in title_upper for x in ["S90D", "S95D", "S85D", "QN900D", "QN800D", "QN95D", "QN90D", "QN85D", "QN80D", "QN70D", "LS03D", "U8000D", "U8090D", "2024"]):
                year_val = 2024

    # LG rules
    elif brand.upper() == "LG":
        # Check model code first
        if model_code != "Unknown":
            # 2026 Codes
            if any(x in model_code for x in ["C6", "G6", "B6", "W6", "M6", "QNED86B", "QNED80B", "QNED87B", "QNED81B", "QNED72B", "QNED71B", "QNED70B", "QNED7EB", "UA77", "MRGB87B", "MRGB87", "MRGB86B", "MRGB86", "LX7B", "LX6", "27LX6TDGA", "27LX6", "STANBYME 2", "STANBYME", "QLED7EB", "MRGB96B", "NU850", "NU800", "NU80", "NU8", "LB700"]):
                year_val = 2026
            # 2025 Codes
            elif any(x in model_code for x in ["C5", "G5", "B5", "W5", "M5", "QNED86A", "QNED80A", "QNED87A", "QNED87", "QNED72A", "QNED7EA", "UA75", "MRGB87A", "LX7A", "LX5", "QNED70A", "NANO81A", "NANO81", "NANO80A", "QNED93A"]):
                year_val = 2025
            # 2024 Codes
            elif any(x in model_code for x in ["C4", "G4", "B4", "M4", "UA73", "LX4"]):
                year_val = 2024
        
        # Fallback to title keywords
        if year_val is None:
            if any(x in title_upper for x in ["C6", "G6", "B6", "W6", "M6", "QNED86B", "QNED81B", "QNED80B", "QNED87B", "QNED72B", "QNED71B", "QNED70B", "QNED7EB", "UA77", "MRGB87B", "MRGB86B", "MRGB96B", "27LX6TDGA", "27LX6", "STANBYME 2", "STANBYME", "NU800", "NU80", "2026"]):
                year_val = 2026
            elif any(x in title_upper for x in ["C5", "G5", "B5", "W5", "M5", "QNED86A", "QNED80A", "QNED87A", "QNED72A", "QNED70A", "QNED7EA", "UA75", "NANO81A", "NANO80A", "QNED93A", "2025"]):
                year_val = 2025
            elif any(x in title_upper for x in ["C4", "G4", "B4", "M4", "2024"]):
                year_val = 2024
                
        if model_code != "Unknown" and "UA73" in model_code and "2025" in title_upper:
            year_val = 2025
            
    return year_val

def parse_product_name(title, brand):
    title_upper = title.upper()
    model_code = "Unknown"
    
    # 1. Size (Zoll / " / Inch / cm)
    size_val = 55  # Default fallback
    size_match = re.search(r'(\d+)\s*"', title)
    if not size_match:
        size_match = re.search(r'(\d+)\s*Zoll', title, re.IGNORECASE)
    if not size_match:
        size_match = re.search(r'(\d+)\s*inch', title, re.IGNORECASE)
    if not size_match:
        size_match = re.search(r'(\d{2,3})\s*(?:cm|pouces)', title, re.IGNORECASE)
        if size_match:
            cm_val = int(size_match.group(1))
            if cm_val > 100:
                size_val = int(round(cm_val / 2.54))
            elif cm_val >= 22:
                size_val = cm_val
            size_match = True
            
    if isinstance(size_match, re.Match):
        size_val = int(size_match.group(1))
    
    # 2. Model Code
    if brand.upper() == "LG":
        oled_match = re.search(r'\b(OLED\d{2,3}[A-Z]{1,3}\d{0,2}[A-Z]*)\b', title_upper)
        if oled_match:
            model_code = oled_match.group(1)
        else:
            lg_code_match = re.search(r'\b(\d{2,3}(?:QNED|NANO|MRGB|LX|NU|UA|UT|UR|UQ|LB|QLED)\w*)\b', title_upper)
            if lg_code_match:
                model_code = lg_code_match.group(1)
            else:
                words = title_upper.split()
                for w in words:
                    clean_w = re.sub(r'[^\w-]', '', w)
                    if any(junk in clean_w for junk in ["LEDLG", "OLEDLG", "TV", "CINEMA"]):
                        continue
                    if any(char.isdigit() for char in clean_w) and len(clean_w) >= 5:
                        model_code = clean_w
                        break
        # Smart fallback for LG series when model_code is Unknown or generic
        if model_code == "Unknown" or not any(x in model_code for x in ["QNED", "OLED", "MRGB", "NANO", "LX", "NU", "UA", "UT", "UR", "UQ", "LB", "QLED"]):
            if "27LX6" in title_upper or "STANBYME 2" in title_upper or "STANBYME" in title_upper:
                model_code = "27LX6TDGA"
                size_val = 27
            elif "QNED86B" in title_upper or "QNED86" in title_upper:
                model_code = f"{size_val}QNED86B"
            elif "QNED81B" in title_upper or "QNED81" in title_upper:
                model_code = f"{size_val}QNED81B"
            elif "QNED80B" in title_upper or "QNED80" in title_upper:
                model_code = f"{size_val}QNED80B"
            elif "QNED72B" in title_upper or "QNED72" in title_upper:
                model_code = f"{size_val}QNED72B"
            elif "QNED71B" in title_upper or "QNED71" in title_upper:
                model_code = f"{size_val}QNED71B"
            elif "QNED70B" in title_upper or "QNED70" in title_upper:
                model_code = f"{size_val}QNED70B"
            elif "MRGB87B" in title_upper or "MRGB87" in title_upper:
                model_code = f"{size_val}MRGB87B"
            elif "MRGB96B" in title_upper or "MRGB96" in title_upper:
                model_code = f"{size_val}MRGB96B"
    else: # SAMSUNG or others
        words = title_upper.split()
        for w in words:
            clean_w = w.replace("(", "").replace(")", "").replace(",", "").replace('"', '').strip()
            if any(char.isdigit() for char in clean_w) and len(clean_w) >= 6:
                if any(clean_w.startswith(pre) for pre in ["GQ", "TQ", "QE", "UE", "GU", "MRE"]):
                    model_code = clean_w
                    break
        if model_code == "Unknown":
            patterns = [
                r'\b(QN\d{2,3}[FH])\b',
                r'\b(S\d{2,3}[FH])\b',
                r'\b(U\d{3,4}[FH])\b',
                r'\b(M\d{2,3}[FH])\b',
                r'\b(R\d{2,3}[FH])\b',
                r'\b(LS03[A-Z]{1,2})\b',
                r'\b(F\d{4})\b',
                r'\b(MR\d{2,3}[FH])\b'
            ]
            for pat in patterns:
                m = re.search(pat, title_upper)
                if m:
                    series = m.group(1)
                    if series.startswith('S') or series.startswith('QN') or series.startswith('LS'):
                        model_code = f"QE{size_val}{series}"
                    elif series.startswith('U') or series.startswith('M') or series.startswith('F'):
                        model_code = f"UE{size_val}{series}"
                    elif series.startswith('R') or series.startswith('MR'):
                        model_code = f"MRE{size_val}{series}"
                    break
        if model_code == "Unknown":
            for w in words:
                clean_w = w.replace("(", "").replace(")", "").replace(",", "").replace('"', '').strip()
                if any(char.isdigit() for char in clean_w) and len(clean_w) >= 6:
                    model_code = clean_w
                    break
        
    # Always prioritize model code size digits if valid 2-digit screen size exists (e.g. TQ65S92H -> 65)
    if model_code != "Unknown":
        code_nums = re.findall(r'\d+', model_code)
        if code_nums and len(code_nums[0]) in [2, 3]:
            parsed_code_size = int(code_nums[0])
            if parsed_code_size in [42, 43, 48, 50, 55, 65, 75, 77, 83, 85, 86, 98, 100]:
                size_val = parsed_code_size
            
    # 3. Year
    year_val = classify_year(brand, model_code, title_upper)
    
    # 4. Display Type
    display_type = "LED"
    code_upper = model_code.upper()
    if brand.upper() == "SAMSUNG":
        if "OLED" in code_upper or any(x in code_upper for x in ["S90", "S91", "S92", "S93", "S94", "S95", "S99", "S85", "S83", "S84", "S86", "S89"]) or bool(re.search(r'S(?:9[0-9]|8[0-9])', code_upper)):
            display_type = "OLED"
        elif "MRE" in code_upper or "MRGB" in code_upper:
            display_type = "Micro RGB"
        elif "QN" in code_upper:
            display_type = "Neo QLED"
        elif any(x in code_upper for x in ["M70", "M72", "M74", "M80", "M82", "M84"]) or "MINI LED" in title_upper:
            display_type = "Mini LED"
        elif "Q" in code_upper or "LS" in code_upper:
            display_type = "QLED"
        elif any(x in code_upper for x in ["U8", "UA", "UT", "UR", "UQ", "UX"]):
            display_type = "UHD 4K"
    elif brand.upper() == "LG":
        if "MRGB" in code_upper or "MRGB" in title_upper:
            display_type = "Micro RGB"
        elif "QNED" in code_upper or "QNED" in title_upper:
            display_type = "QNED"
        elif "OLED" in code_upper or "OLED" in title_upper or re.search(r'OLED\d{2}', code_upper):
            display_type = "OLED evo" if ("EVO" in title_upper or any(x in code_upper for x in ["G5", "G6", "C5", "C6", "M5", "M6"])) else "OLED"
        elif "NANO" in code_upper or "NANO" in title_upper:
            display_type = "NanoCell"
        elif any(x in code_upper for x in ["NU900", "NU90", "NU85", "NU80", "NU800", "NU850", "LB650", "LB", "UA75", "UA73", "UA77", "UT", "UR", "UQ", "LQ"]):
            display_type = "UHD 4K"
        elif "LX" in code_upper or "STANBYME" in code_upper:
            display_type = "Lifestyle"

    return model_code, size_val, year_val, display_type

def clean_promotion_text(promo, brand, year, model_code, currency):
    if not promo or promo == "None":
        return "None"
    
    promo_upper = promo.upper()
    
    # 1. Cashback und oder 26% Rabatt auf Soundbar QS700F
    if "CASHBACK UND ODER 26%" in promo_upper or "CASHBACK AND/OR 26% OFF" in promo_upper:
        return "Samsung Cashback and/or 26% Off Soundbar QS700F"
        
    # 2. Gratis Music Studio 5
    if "MUSIC STUDIO 5" in promo_upper:
        return "Free Music Studio 5 Soundbar (HW-LS50H/EN)"
        
    # 3. mit kostenlosem Zusatzprodukt
    if "ZUSATZPRODUKT" in promo_upper or "WITH FREE PROMOTIONAL PRODUCT" in promo_upper:
        return "Free Promotional Product (Sound Device)"
        
    # 4. Restposten / Outlet
    if "RESTPOSTEN" in promo_upper:
        return "Clearance (Restposten)"
    if "OUTLET" in promo_upper:
        return "Outlet (Clearance)"
        
    # 5. SoundSuite / Soundbar (LG 2026)
    if "SOUNDSUITE" in promo_upper or "SOUNDBAR DAZU" in promo_upper:
        if brand.upper() == "LG" and year == 2026:
            return "Free SoundSuite/Soundbar with LG OLED TV 2026 Purchase"
            
    # 6. Sound Device geschenkt
    if "SOUND DEVICE GESCHENKT" in promo_upper:
        if year == 2026:
            return "Free Sound Device with 2026 TV Purchase"
            
    # 7. Aus unserer Werbung
    if "WERBUNG" in promo_upper or "WEEKLY AD" in promo_upper:
        return "Featured in Weekly Ad (Aus unserer Werbung)"
        
    # 8. Online Only
    if "ONLINE ONLY" in promo_upper or "ONLINE-ONLY" in promo_upper:
        return "Online Only"
        
    # 9. Direct Cut & Voucher Codes (Currys / UK & EU)
    if "DIRECT CUT" in promo_upper:
        return promo
        
    # 10. Sale / statt (Discount badges)
    if any(x in promo_upper for x in ["STATT", "WAS", "SPARA", "SAVE", "RABATT", "DISCOUNT"]):
        pct_match = re.search(r'(\d+)\s*%', promo)
        was_match = re.search(r'(?:was|statt|spara|save)\s*(?:chf|eur|gbp|sek|€|£)?\s*([\d’\x27\x60,.]+)', promo, re.IGNORECASE)
        
        pct = pct_match.group(1) if pct_match else None
        was_val = was_match.group(1).replace("’", "").replace("'", "").replace("`", "").strip() if was_match else None
        if was_val and (was_val.endswith(".") or was_val.endswith(",")):
            was_val = was_val[:-1]
            
        if pct and was_val:
            return f"Sale {pct}%; was {currency} {was_val}"
        elif was_val:
            return f"was {currency} {was_val}"
        elif pct:
            return f"Sale {pct}%"
            
    return promo

def clean_numeric_price(price_str):
    if price_str is None or price_str == "":
        return 0.0
    if isinstance(price_str, (int, float)):
        val = float(price_str)
        if 0 < val < 10.0 and round(val, 3) != round(val, 2):
            val = val * 1000.0
        return val if val >= 50.0 else 0.0
        
    clean_p = str(price_str).replace("€", "").replace("£", "").replace("kr", "").replace("SEK", "").replace("EUR", "").replace("GBP", "").replace("CHF", "").strip()
    
    # European format check: "1.799", "2.499", "1.499"
    if re.match(r'^\d{1,3}\.\d{3}$', clean_p):
        clean_p = clean_p.replace(".", "")
    elif re.match(r'^\d{1,3},\d{3}$', clean_p):
        clean_p = clean_p.replace(",", "")
    elif "," in clean_p and "." in clean_p:
        if clean_p.find(".") < clean_p.find(","):
            clean_p = clean_p.replace(".", "").replace(",", ".")
        else:
            clean_p = clean_p.replace(",", "")
    elif "," in clean_p:
        parts = clean_p.split(",")
        if len(parts[-1]) == 2:
            clean_p = clean_p.replace(",", ".")
        else:
            clean_p = clean_p.replace(",", "")
            
    clean_p = re.sub(r'[^\d.]', '', clean_p)
    try:
        val = float(clean_p) if clean_p else 0.0
        if 0 < val < 10.0:
            val = val * 1000.0
        return val if val >= 50.0 else 0.0
    except ValueError:
        return 0.0

def sync_retailer_sheet(filename, sheetname, brand, products, currency):
    if not os.path.exists(filename):
        wb = openpyxl.Workbook()
        # Remove default sheet
        default_sheet = wb.active
        wb.remove(default_sheet)
    else:
        wb = openpyxl.load_workbook(filename)
        
    print(f"  ➔ Syncing to sheet: {sheetname} ({currency})...")
    
    # Get or create sheet
    if sheetname in wb.sheetnames:
        sheet = wb[sheetname]
    else:
        sheet = wb.create_sheet(sheetname)
        
    # Set headers if sheet is empty
    headers = [
        "Brand", "Model Year", "Display Type", "Size (Inch)", "Model Code",
        f"Price ({currency})", "Shipping", "Installment", "Cashback", "General Promotions"
    ]
    if sheet.max_row <= 1:
        sheet.append(headers)
        
    # Delete old rows for this brand
    removed_count = 0
    for r in range(sheet.max_row, 1, -1):
        brand_val = sheet.cell(row=r, column=1).value
        if str(brand_val).upper() == str(brand).upper():
            sheet.delete_rows(r)
            removed_count += 1
    if removed_count > 0:
        print(f"    ➔ Removed {removed_count} obsolete {brand} rows.")
        
    # Append new rows
    added_count = 0
    for p in products:
        # Determine cashback amount if explicitly provided or parse from promo
        cashback_val = 0
        if p.get("cashback") and isinstance(p["cashback"], (int, float)) and p["cashback"] > 0:
            cashback_val = int(round(p["cashback"]))
            
        promo_text = p.get("promo", "None") or "None"
        
        # Standardize promo text
        promo_text = clean_promotion_text(promo_text, brand, p["year"], p["model_code"], currency)
        
        if cashback_val == 0 and any(x in promo_text.upper() for x in ["CASHBACK", "캐시백", "REEMBOLSO", "RIMBORSO", "RÜCKVERGÜTUNG", "RUCKVERGUTUNG", "REMISE", "ODR"]):
            # Match amount before keyword: e.g. "€100,- cashback", "100€ cashback", "150 EUR cashback"
            cb_match_before = re.search(r'(?:chf|eur|gbp|sek|€|£)?\s*(\d+)\s*(?:chf|eur|gbp|sek|€|£|[,-]+)?\s*(?:cashback|캐시백|reembolso|rimborso|rückvergütung|ruckvergutung|remise|odr)', promo_text, re.IGNORECASE)
            if cb_match_before:
                cashback_val = int(cb_match_before.group(1))
            else:
                # Match amount after keyword: e.g. "Cashback 100 €", "Cashback CHF 200", "Cashback sichern 150€"
                cb_match_after = re.search(r'(?:cashback|캐시백|reembolso|rimborso|rückvergütung|ruckvergutung|remise|odr)\s*(?:von|bis\s*zu|sichern|immédiate|différée)?\s*(?:chf|eur|gbp|sek|€|£)?\s*(\d+)', promo_text, re.IGNORECASE)
                if cb_match_after:
                    cashback_val = int(cb_match_after.group(1))
            if cashback_val < 10: # Ignore single-digit false positives
                cashback_val = 0
                    
        final_price = clean_numeric_price(p["price"])
        
        sheet.append([
            p["brand"],
            p["year"],
            p["display_type"],
            p["size"],
            p["model_code"],
            final_price,
            "Free",  # Shipping
            None,    # Installment
            cashback_val,
            promo_text
        ])
        added_count += 1
        
    print(f"    ➔ Added {added_count} live {brand} rows.")
    final_filename = save_workbook_with_retry(wb, filename)
    wb.close()
    return final_filename

def main():
    data_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
    if not os.path.exists(data_dir):
        os.makedirs(data_dir)
        
    # Resolve MMDD date
    from datetime import datetime
    target_date = None
    if len(sys.argv) > 1 and not sys.argv[1].startswith("-"):
        target_date = sys.argv[1]
    elif "--date" in sys.argv:
        idx = sys.argv.index("--date")
        if idx + 1 < len(sys.argv):
            target_date = sys.argv[idx + 1]
            
    today_mmdd = target_date if target_date else datetime.now().strftime("%m%d")
    
    excel_filename = os.path.join(data_dir, f"price tracker_EU_2026 {today_mmdd}_v1.xlsx")
    
    # Check if there is an existing file for today to append/modify
    latest_files = glob.glob(os.path.join(data_dir, f"price tracker_EU_2026 {today_mmdd}*.xlsx"))
    latest_files = [f for f in latest_files if not os.path.basename(f).startswith("~$")]
    if latest_files:
        latest_files.sort()
        excel_filename = latest_files[-1]
        print(f"[SYNC] Found existing workbook: {excel_filename}")
    else:
        # Check if there is a previous date workbook to clone
        prev_files = glob.glob(os.path.join(data_dir, "price tracker_EU_2026 *.xlsx"))
        prev_files = [f for f in prev_files if not os.path.basename(f).startswith("~$")]
        if prev_files:
            prev_files.sort()
            prev_file = prev_files[-1]
            print(f"[SYNC] Cloning latest workbook: {prev_file} -> {excel_filename}")
            shutil.copyfile(prev_file, excel_filename)
    # Mandatory Pre-flight Live Assertion Gate
    print("=" * 65)
    print(f" 🛡️ MANDATORY PRE-FLIGHT LIVE ASSERTION GATE (Survey Date: {today_mmdd})")
    print("=" * 65)
    today_date_str = datetime.now().strftime("%Y-%m-%d")
    stale_files = []
    checked_count = 0
    for s_key in EU_SOURCES:
        if s_key == 'mm-ch':
            continue
        for b in ["samsung", "lg"]:
            r_path = os.path.join(data_dir, f"raw_{s_key}_{b}.json")
            if not os.path.exists(r_path):
                g_cands = glob.glob(os.path.join(data_dir, f"raw_{s_key}_{b}*.json"))
                if g_cands:
                    r_path = g_cands[0]
            if os.path.exists(r_path):
                mtime = datetime.fromtimestamp(os.path.getmtime(r_path)).strftime("%Y-%m-%d")
                fsize = os.path.getsize(r_path)
                checked_count += 1
                if target_date:
                    if fsize < 100:
                        stale_files.append((os.path.basename(r_path), "EMPTY FILE (<100B)"))
                elif mtime != today_date_str:
                    stale_files.append((os.path.basename(r_path), mtime))
            else:
                stale_files.append((f"raw_{s_key}_{b}.json", "MISSING"))
                
    if stale_files:
        print(f"❌ [ASSERTION FAILURE] Found {len(stale_files)} invalid or missing raw datasets:")
        for fname, f_date in stale_files:
            print(f"   • {fname}: Status {f_date}")
        raise RuntimeError(f"PRE-FLIGHT ASSERTION FAILED: {len(stale_files)} files require valid live datasets.")
    print(f"✅ All {checked_count} raw EU datasets verified fresh for today ({today_date_str})!\n")

    # Iterate over all EU sources and brands
    for src_key, cfg in EU_SOURCES.items():
        for brand in ["samsung", "lg"]:
            products_list = []
            
            if src_key == 'mm-ch':
                swiss_file = os.path.join(data_dir, f"price tracker_swiss_2026 {today_mmdd}_v1.xlsx")
                if not os.path.exists(swiss_file):
                    swiss_files = glob.glob(os.path.join(data_dir, "price tracker_swiss_2026 *.xlsx"))
                    swiss_files = [f for f in swiss_files if not os.path.basename(f).startswith("~$")]
                    if swiss_files:
                        swiss_files.sort()
                        swiss_file = swiss_files[-1]
                if os.path.exists(swiss_file):
                    try:
                        wb_s = openpyxl.load_workbook(swiss_file, data_only=True)
                        s_name = "MediaMarkt_Samsung_Full" if brand.lower() == "samsung" else "MediaMarkt_LG_Full"
                        if s_name in wb_s.sheetnames:
                            ws_s = wb_s[s_name]
                            for r in range(2, ws_s.max_row + 1):
                                b_v = ws_s.cell(row=r, column=1).value
                                yr_v = ws_s.cell(row=r, column=2).value
                                disp_v = ws_s.cell(row=r, column=3).value
                                sz_v = ws_s.cell(row=r, column=4).value
                                code_v = ws_s.cell(row=r, column=5).value
                                pr_v = ws_s.cell(row=r, column=6).value
                                cb_v = ws_s.cell(row=r, column=9).value or 0
                                promo_v = ws_s.cell(row=r, column=10).value or "None"
                                if code_v and pr_v and str(code_v).upper() != "UNKNOWN":
                                    products_list.append({
                                        "brand": str(b_v).strip() if b_v else brand.upper(),
                                        "year": int(yr_v) if yr_v else 2025,
                                        "display_type": str(disp_v).strip() if disp_v else "LED",
                                        "size": int(sz_v) if sz_v else 55,
                                        "model_code": str(code_v).strip(),
                                        "price": float(pr_v),
                                        "promo": str(promo_v),
                                        "cashback": int(cb_v) if cb_v else 0
                                    })
                        wb_s.close()
                    except Exception as e_s:
                        print(f"[WARN] Failed loading Swiss Excel for mm-ch: {e_s}")
            else:
                if src_key == 'mm-at':
                    pattern = os.path.join(data_dir, f"raw_mm-at_{brand}*.json")
                else:
                    pattern = os.path.join(data_dir, f"raw_{src_key}_{brand}*.json")
                json_files = glob.glob(pattern)
                
                scraped_items = []
                for jf in json_files:
                    try:
                        with open(jf, "r", encoding="utf-8") as f:
                            data = json.load(f)
                            if isinstance(data, dict) and "items" in data:
                                scraped_items.extend(data["items"])
                            elif isinstance(data, list):
                                scraped_items.extend(data)
                    except Exception as e:
                        print(f"[WARN] Failed to load JSON {jf}: {e}")
                        
                if not scraped_items:
                    continue
                    
                seen_codes = {}
                for item in scraped_items:
                    if not isinstance(item, dict):
                        continue
                    if item.get("model_code") and item.get("year") in [2025, 2026] and item.get("size", 0) >= 22:
                        model_code = item["model_code"]
                        year = item["year"]
                        size = item["size"]
                        display_type = item.get("display") or item.get("display_type") or "LED"
                        price = clean_numeric_price(item.get("price"))
                        promo = item.get("promo") or "None"
                    else:
                        name = item.get("name") or item.get("title") or ""
                        raw_price = item.get("price")
                        promo = item.get("promo") or "None"
                        if not name:
                            continue
                        name_upper = name.upper()
                        if any(x in name_upper for x in ["MONITOR", "ODYSSEY", "ULTRAGEAR", "MYVIEW", "SOUNDBAR", "HIFI", "BEAMER", "PROJEKTOR", "ZUBEHÖR", "HALTERUNG", "WALLMOUNT", "CORNICE", "GALAXY BUDS"]):
                            continue
                        if any(x in name_upper for x in ["RETURNED", "REFURBISHED", "USED", "GEBRAUCHT", "REACONDICIONADO", "RICONDIZIONATO"]):
                            continue
                        model_code, size, year, display_type = parse_product_name(name, brand)
                        if size < 22 or year not in [2025, 2026]:
                            continue
                        price = clean_numeric_price(raw_price)
                        
                    if price <= 0.0:
                        continue
                    cashback = clean_numeric_price(item.get("cashback", 0))
                    p_rec = {
                        "brand": brand.upper() if brand.lower() == "samsung" else "LG",
                        "year": year,
                        "display_type": display_type,
                        "size": size,
                        "model_code": model_code,
                        "price": price,
                        "promo": promo,
                        "cashback": cashback
                    }
                    if model_code in seen_codes:
                        if price < seen_codes[model_code]["price"]:
                            seen_codes[model_code] = p_rec
                        elif seen_codes[model_code].get("cashback", 0) == 0 and cashback > 0:
                            seen_codes[model_code]["cashback"] = cashback
                            if seen_codes[model_code].get("promo") in ["None", ""]:
                                seen_codes[model_code]["promo"] = promo
                    else:
                        seen_codes[model_code] = p_rec
                products_list = list(seen_codes.values())
                
            if not products_list:
                continue
                
            products_list.sort(key=lambda x: (x["size"], x["model_code"]))
            sheetname = f"{cfg['retailer']}_{cfg['cc']}_{brand.upper() if brand.lower() == 'samsung' else 'LG'}_Full"
            excel_filename = sync_retailer_sheet(
                excel_filename, sheetname,
                "Samsung" if brand.lower() == "samsung" else "LG",
                products_list, cfg["currency"]
            )
            
    print(f"\n[EU SYNC DONE] Saved successfully to {excel_filename}")
    
    # Auto-mirror updated Excel workbook to History_EU folder
    try:
        from datetime import datetime
        today_folder = datetime.now().strftime("%Y %m%d")
        hist_dir = os.path.abspath(os.path.join(data_dir, "..", "History_EU", today_folder))
        if not os.path.exists(hist_dir):
            os.makedirs(hist_dir)
        hist_file = os.path.join(hist_dir, os.path.basename(excel_filename))
        shutil.copy2(excel_filename, hist_file)
        print(f"  ➔ [MIRRORED TO HISTORY_EU] {hist_file}")
    except Exception as e_h:
        print(f"[WARN] Failed mirroring to History_EU: {e_h}")

if __name__ == "__main__":
    main()
