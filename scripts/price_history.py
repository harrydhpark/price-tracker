import os
import sys
import re
import glob
import shutil
import sqlite3
from datetime import datetime
from typing import List, Dict, Any, Optional
import openpyxl

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DATA_DIR = os.path.join(ROOT_DIR, "data")
DB_PATH = os.path.join(DATA_DIR, "price_history.db")
BACKUP_PATH = os.path.join(DATA_DIR, "price_history.db.bak")
HISTORY_DIR = os.path.join(ROOT_DIR, "History")
HISTORY_EU_DIR = os.path.join(ROOT_DIR, "History_EU")
HISTORY_AU_DIR = os.path.join(ROOT_DIR, "History_AU")

def get_connection(db_path: str = DB_PATH) -> sqlite3.Connection:
    """Returns a connection to the SQLite price history database."""
    os.makedirs(os.path.dirname(db_path), exist_ok=True)
    con = sqlite3.connect(db_path)
    con.row_factory = sqlite3.Row
    return con

def migrate_db(db_path: str = DB_PATH, backup: bool = True) -> bool:
    """Migrates daily_prices table to the multi-country, currency-neutral schema.
    Preserves all existing historical records and is idempotent.
    """
    if not os.path.exists(db_path):
        print(f"[DB MIGRATE] DB file does not exist at {db_path}. Initializing clean schema.")
        con = get_connection(db_path)
        _create_schema(con)
        con.close()
        return True

    if backup and not os.path.exists(BACKUP_PATH):
        try:
            shutil.copy2(db_path, BACKUP_PATH)
            print(f"[DB MIGRATE] Backed up original database to {BACKUP_PATH}")
        except Exception as e:
            print(f"[WARN] Could not create database backup: {e}")

    con = get_connection(db_path)
    cur = con.cursor()

    # Check existing table
    table_exists = cur.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='daily_prices'").fetchone()
    if not table_exists:
        _create_schema(con)
        con.close()
        print("[DB MIGRATE] Initialized daily_prices schema.")
        return True

    cols = [c[1] for c in cur.execute("PRAGMA table_info(daily_prices)").fetchall()]
    
    # Check if migration is needed (legacy schema had original_price_chf)
    if "original_price_chf" in cols or "country" not in cols:
        print(f"[DB MIGRATE] Legacy schema detected (Columns: {cols}). Migrating to multi-country schema...")
        cur.execute("ALTER TABLE daily_prices RENAME TO daily_prices_legacy")
        _create_schema(con)

        # Transfer existing Swiss rows to new schema
        # country = 'CH', currency = 'CHF'
        cur.execute("""
            INSERT OR REPLACE INTO daily_prices (
                id, timestamp, date, country, retailer, manufacturer, model_code,
                base_model, display_size_inch, panel_type, selling_price, original_price,
                shipping_cost, total_price, currency, promo_text, in_stock, delivery_eta
            )
            SELECT
                id,
                timestamp,
                date,
                'CH' AS country,
                retailer,
                manufacturer,
                model_code,
                base_model,
                display_size_inch,
                panel_type,
                selling_price_chf AS selling_price,
                original_price_chf AS original_price,
                shipping_cost_chf AS shipping_cost,
                total_price_chf AS total_price,
                'CHF' AS currency,
                promo_text,
                in_stock,
                delivery_eta
            FROM daily_prices_legacy
        """)
        cur.execute("DROP TABLE daily_prices_legacy")
        con.commit()
        migrated_count = cur.execute("SELECT count(*) FROM daily_prices").fetchone()[0]
        print(f"[DB MIGRATE] Successfully migrated {migrated_count} legacy rows to new schema.")
    else:
        # Schema already up-to-date, verify indexes
        _ensure_indexes(con)

    con.close()
    return True

def _create_schema(con: sqlite3.Connection):
    cur = con.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS daily_prices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            timestamp TEXT,
            date TEXT,
            country TEXT,
            retailer TEXT,
            manufacturer TEXT,
            model_code TEXT,
            base_model TEXT,
            display_size_inch INTEGER,
            panel_type TEXT,
            selling_price REAL,
            original_price REAL,
            shipping_cost REAL,
            total_price REAL,
            currency TEXT,
            promo_text TEXT,
            in_stock INTEGER,
            delivery_eta TEXT
        )
    """)
    _ensure_indexes(con)
    con.commit()

def _ensure_indexes(con: sqlite3.Connection):
    cur = con.cursor()
    cur.execute("CREATE INDEX IF NOT EXISTS idx_lookup ON daily_prices(country, retailer, manufacturer, model_code, date)")
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_unique_record ON daily_prices(country, retailer, model_code, date)")
    con.commit()

def upsert_records(records: List[Dict[str, Any]], db_path: str = DB_PATH) -> int:
    """Upserts a list of product records into daily_prices.
    Uses conflict resolution on (country, retailer, model_code, date).
    Returns count of affected records.
    """
    if not records:
        return 0

    con = get_connection(db_path)
    cur = con.cursor()
    
    query = """
        INSERT INTO daily_prices (
            timestamp, date, country, retailer, manufacturer, model_code,
            base_model, display_size_inch, panel_type, selling_price, original_price,
            shipping_cost, total_price, currency, promo_text, in_stock, delivery_eta
        ) VALUES (
            :timestamp, :date, :country, :retailer, :manufacturer, :model_code,
            :base_model, :display_size_inch, :panel_type, :selling_price, :original_price,
            :shipping_cost, :total_price, :currency, :promo_text, :in_stock, :delivery_eta
        )
        ON CONFLICT(country, retailer, model_code, date) DO UPDATE SET
            selling_price = excluded.selling_price,
            original_price = excluded.original_price,
            shipping_cost = excluded.shipping_cost,
            total_price = excluded.total_price,
            currency = excluded.currency,
            promo_text = excluded.promo_text,
            display_size_inch = excluded.display_size_inch,
            panel_type = excluded.panel_type,
            base_model = excluded.base_model,
            timestamp = excluded.timestamp,
            in_stock = excluded.in_stock,
            delivery_eta = excluded.delivery_eta
    """
    
    batch = []
    now_iso = datetime.now().astimezone().isoformat()
    for r in records:
        rec = {
            "timestamp": r.get("timestamp") or now_iso,
            "date": r.get("date"),
            "country": str(r.get("country", "CH")).upper(),
            "retailer": r.get("retailer", "Unknown"),
            "manufacturer": str(r.get("manufacturer") or r.get("brand", "Unknown")).upper(),
            "model_code": str(r.get("model_code", "")).strip(),
            "base_model": r.get("base_model") or "",
            "display_size_inch": int(r.get("display_size_inch") or r.get("size") or 0),
            "panel_type": r.get("panel_type") or r.get("display_type") or "",
            "selling_price": float(r.get("selling_price") or r.get("price") or 0.0),
            "original_price": float(r.get("original_price") or r.get("selling_price") or r.get("price") or 0.0),
            "shipping_cost": float(r.get("shipping_cost") or 0.0),
            "total_price": float(r.get("total_price") or r.get("selling_price") or r.get("price") or 0.0),
            "currency": str(r.get("currency", "EUR")).upper(),
            "promo_text": str(r.get("promo_text") or r.get("promotion") or r.get("promo") or ""),
            "in_stock": int(r.get("in_stock", 1)),
            "delivery_eta": str(r.get("delivery_eta") or "")
        }
        if rec["model_code"] and rec["selling_price"] > 0 and rec["date"]:
            batch.append(rec)
            
    cur.executemany(query, batch)
    con.commit()
    con.close()
    return len(batch)

def _extract_date_from_path(file_path: str) -> Optional[str]:
    """Resolves YYYY-MM-DD string from directory path or filename."""
    # Pattern 1: '2026 0704' or '2026 0907'
    m = re.search(r'(\d{4})\s+(\d{2})(\d{2})', file_path)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    # Pattern 2: '2026_0704' or '2026-07-04'
    m = re.search(r'(\d{4})[-_](\d{2})[-_](\d{2})', file_path)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    # Pattern 3: '20260704'
    m = re.search(r'202\d[01]\d[0-3]\d', file_path)
    if m:
        s = m.group(0)
        return f"{s[:4]}-{s[4:6]}-{s[6:8]}"
    return None

def parse_survey_workbook(excel_path: str, default_country: str = None) -> List[Dict[str, Any]]:
    """Parses product records from all '_Full' sheets of an exported survey workbook."""
    if not os.path.exists(excel_path):
        return []

    date_str = _extract_date_from_path(excel_path)
    if not date_str:
        date_str = datetime.now().strftime("%Y-%m-%d")

    records = []
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
        for sname in wb.sheetnames:
            if not sname.endswith("_Full"):
                continue

            # Infer country, retailer, brand, currency from sheet name
            country = default_country
            retailer = "Unknown"
            brand = "Unknown"
            currency = "EUR"

            # Brand
            if "samsung" in sname.lower():
                brand = "SAMSUNG"
            elif "lg" in sname.lower():
                brand = "LG"

            # Swiss sheets
            if sname.startswith("MediaMarkt_Samsung") or sname.startswith("MediaMarkt_LG"):
                country = "CH"
                retailer = "MediaMarkt"
                currency = "CHF"
            elif sname.startswith("Interdiscount"):
                country = "CH"
                retailer = "Interdiscount"
                currency = "CHF"
            elif sname.startswith("Digitec"):
                country = "CH"
                retailer = "Digitec"
                currency = "CHF"
            # EU sheets: Retailer_CC_Brand_Full
            elif "_DE_" in sname:
                country = "DE"; retailer = "MediaMarkt"; currency = "EUR"
            elif "_UK_" in sname or sname.startswith("Currys"):
                country = "UK"; retailer = "Currys"; currency = "GBP"
            elif "_ES_" in sname:
                country = "ES"; retailer = "MediaMarkt"; currency = "EUR"
            elif "_NL_" in sname:
                country = "NL"; retailer = "MediaMarkt"; currency = "EUR"
            elif "_IT_" in sname or sname.startswith("MediaWorld"):
                country = "IT"; retailer = "MediaWorld"; currency = "EUR"
            elif "_AT_" in sname:
                country = "AT"; retailer = "MediaMarkt"; currency = "EUR"
            elif "_CZ_" in sname or sname.startswith("Alza"):
                country = "CZ"; retailer = "Alza"; currency = "CZK"
            elif "_GR_" in sname or sname.startswith("Public"):
                country = "GR"; retailer = "Public"; currency = "EUR"
            elif "_HU_" in sname:
                country = "HU"; retailer = "MediaMarkt"; currency = "HUF"
            elif "_FR_" in sname:
                country = "FR"; retailer = "Fnac Darty"; currency = "EUR"
            # AU sheets
            elif sname.startswith("JBHIFI"):
                country = "AU"; retailer = "JB Hi-Fi"; currency = "AUD"
            elif sname.startswith("TheGoodGuys"):
                country = "AU"; retailer = "The Good Guys"; currency = "AUD"
            elif sname.startswith("HarveyNorman"):
                country = "AU"; retailer = "Harvey Norman"; currency = "AUD"
            elif not country:
                country = "EU"

            ws = wb[sname]
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                continue

            if not header_row:
                continue

            header_map = {}
            for idx, h in enumerate(header_row):
                if h:
                    header_map[str(h).strip().lower()] = idx

            def get_val(row, *aliases):
                for a in aliases:
                    for k, idx in header_map.items():
                        if a in k:
                            if idx < len(row):
                                return row[idx]
                return None

            for row in rows_iter:
                if not row or all(v is None for v in row):
                    continue

                model_code = get_val(row, "model code", "model_code", "model")
                if not model_code or str(model_code).strip() == "" or str(model_code).strip().lower() in ["none", "model code"]:
                    continue
                model_code = str(model_code).strip()

                raw_price = get_val(row, "price", "preis", "selling price")
                if raw_price is None:
                    continue
                try:
                    price_val = float(str(raw_price).replace("CHF", "").replace("EUR", "").replace("GBP", "").replace("AUD", "").replace("Ft", "").replace(",", ".").replace("'", "").replace(" ", "").strip())
                except (ValueError, TypeError):
                    continue

                if price_val <= 0:
                    continue

                raw_orig = get_val(row, "original price", "was price", "uvp", "original")
                orig_val = price_val
                if raw_orig is not None:
                    try:
                        parsed_orig = float(str(raw_orig).replace("CHF", "").replace("EUR", "").replace("GBP", "").replace("AUD", "").replace("Ft", "").replace(",", ".").replace("'", "").replace(" ", "").strip())
                        if parsed_orig >= price_val:
                            orig_val = parsed_orig
                    except (ValueError, TypeError):
                        pass

                raw_size = get_val(row, "size", "inch", "screen")
                size_val = 0
                if raw_size is not None:
                    try:
                        m_size = re.search(r'\d+', str(raw_size))
                        if m_size:
                            size_val = int(m_size.group(0))
                    except Exception:
                        pass

                panel_type = get_val(row, "display type", "display", "category", "panel") or ""
                promo = get_val(row, "promo", "promotion", "general promotions") or ""
                shipping = get_val(row, "shipping") or ""
                shipping_val = 0.0
                if shipping and str(shipping).lower() not in ["free", "kostenlos", "none", "0"]:
                    try:
                        shipping_val = float(re.sub(r'[^\d.]', '', str(shipping).replace(",", ".")))
                    except Exception:
                        shipping_val = 0.0

                records.append({
                    "date": date_str,
                    "country": country,
                    "retailer": retailer,
                    "manufacturer": brand,
                    "model_code": model_code,
                    "base_model": "",
                    "display_size_inch": size_val,
                    "panel_type": str(panel_type),
                    "selling_price": price_val,
                    "original_price": orig_val,
                    "shipping_cost": shipping_val,
                    "total_price": price_val + shipping_val,
                    "currency": currency,
                    "promo_text": str(promo),
                    "in_stock": 1,
                    "delivery_eta": ""
                })
        wb.close()
    except Exception as e:
        print(f"[WARN] Error reading {excel_path}: {e}")

    return records

def backfill_from_history(db_path: str = DB_PATH) -> Dict[str, Any]:
    """Reconstructs historical time-series database from History/, History_EU/, and History_AU/ workbooks.
    Complies strictly with AGENTS.md §1 Zero Historical Price Injection Policy:
    rebuilding the historical baseline repository only.
    """
    print("=" * 65)
    print(" 📚 REBUILDING PRICE HISTORY TIME-SERIES DB (BACKFILL)")
    print("=" * 65)
    migrate_db(db_path)

    all_records = []
    
    # 1. Swiss History workbooks
    swiss_files = glob.glob(os.path.join(HISTORY_DIR, "*", "price tracker_swiss_*.xlsx"))
    swiss_files = [f for f in swiss_files if not os.path.basename(f).startswith("~$")]
    print(f"➔ Found {len(swiss_files)} Swiss history workbooks.")
    for f in sorted(swiss_files):
        recs = parse_survey_workbook(f, default_country="CH")
        all_records.extend(recs)

    # 2. Pan-European History workbooks
    eu_files = glob.glob(os.path.join(HISTORY_EU_DIR, "*", "price tracker_EU_*.xlsx"))
    eu_files = [f for f in eu_files if not os.path.basename(f).startswith("~$")]
    print(f"➔ Found {len(eu_files)} Pan-European history workbooks.")
    for f in sorted(eu_files):
        recs = parse_survey_workbook(f)
        all_records.extend(recs)

    # 3. Australia History workbooks
    au_files = glob.glob(os.path.join(HISTORY_AU_DIR, "*", "price tracker_AU_*.xlsx"))
    au_files = [f for f in au_files if not os.path.basename(f).startswith("~$")]
    print(f"➔ Found {len(au_files)} Australia history workbooks.")
    for f in sorted(au_files):
        recs = parse_survey_workbook(f, default_country="AU")
        all_records.extend(recs)

    print(f"\n[BACKFILL] Extracted {len(all_records)} raw product records across all history archives.")
    inserted_count = upsert_records(all_records, db_path)
    print(f"[BACKFILL] Upserted {inserted_count} records into daily_prices.")

    # Summary verification
    con = get_connection(db_path)
    cur = con.cursor()
    total_rows = cur.execute("SELECT count(*) FROM daily_prices").fetchone()[0]
    distinct_dates = cur.execute("SELECT count(distinct date) FROM daily_prices").fetchone()[0]
    dates = cur.execute("SELECT min(date), max(date) FROM daily_prices").fetchone()
    breakdown = cur.execute("SELECT country, count(*), count(distinct retailer), count(distinct model_code) FROM daily_prices GROUP BY country ORDER BY count(*) DESC").fetchall()
    con.close()

    summary = {
        "total_rows": total_rows,
        "distinct_dates": distinct_dates,
        "date_range": [dates[0], dates[1]],
        "countries": {row[0]: {"records": row[1], "retailers": row[2], "models": row[3]} for row in breakdown}
    }

    print("\n" + "=" * 50)
    print(" PRICE HISTORY DATABASE STATUS SUMMARY")
    print("=" * 50)
    print(f"Total Records   : {total_rows:,}")
    print(f"Distinct Dates  : {distinct_dates} (Range: {dates[0]} ~ {dates[1]})")
    print("Country Breakdown:")
    for c, stats in summary["countries"].items():
        print(f"  • {c:4s}: {stats['records']:,} rows | {stats['retailers']} retailers | {stats['models']} unique models")
    print("=" * 50)

    return summary

def record_survey_from_excel(excel_path: str, country: str = None, db_path: str = DB_PATH) -> int:
    """Convenience function called by sync scripts to record today's survey into DB."""
    try:
        migrate_db(db_path)
        records = parse_survey_workbook(excel_path, default_country=country)
        if records:
            inserted = upsert_records(records, db_path)
            print(f"📊 [PRICE HISTORY] Recorded {inserted} survey records from {os.path.basename(excel_path)} into price_history.db")
            return inserted
    except Exception as e:
        print(f"[WARN] Failed to record survey to price_history.db: {e}")
    return 0

if __name__ == "__main__":
    backfill_from_history()
