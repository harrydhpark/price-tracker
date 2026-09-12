# -*- coding: utf-8 -*-
import sys, os, json, re
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from scrape_and_sync_hungary import clean_and_standardize, write_standard_sheet

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DATA_DIR = os.path.join(ROOT_DIR, "data")
HISTORY_DIR = os.path.join(ROOT_DIR, "History_EU")

def process_hungary_dump(dump_path):
    print(f"[HU SYNC] Reading dump from {dump_path}")
    with open(dump_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    m = re.search(r'### Result\s*\n(.*?)\n### Ran Playwright code', content, re.DOTALL)
    if m:
        raw_json_str = m.group(1).strip()
    else:
        s_idx = content.find('{')
        e_idx = content.rfind('}')
        raw_json_str = content[s_idx:e_idx+1]
        
    data = json.loads(raw_json_str)
    items = data.get("items", [])
    print(f"[HU SYNC] Loaded {len(items)} raw items from dump.")
    
    raw_sec = [it for it in items if it.get("brand", "").upper() == "SAMSUNG"]
    raw_lg = [it for it in items if it.get("brand", "").upper() == "LG"]
    
    raw_cache_file = os.path.join(DATA_DIR, "raw_mediamarkt_hu_deep.json")
    with open(raw_cache_file, "w", encoding="utf-8") as f:
        json.dump({"samsung": raw_sec, "lg": raw_lg}, f, ensure_ascii=False, indent=2)
    print(f"✅ Saved {raw_cache_file} ({len(raw_sec)} Samsung, {len(raw_lg)} LG)")
    
    clean_sec = clean_and_standardize(raw_sec, "Samsung")
    clean_lg = clean_and_standardize(raw_lg, "LG")
    
    print(f"📊 [CLEANED] Samsung: {len(clean_sec)}, LG: {len(clean_lg)} (Total: {len(clean_sec) + len(clean_lg)})")
    
    all_json = clean_sec + clean_lg
    with open(os.path.join(DATA_DIR, "mediamarkt_hu_full.json"), "w", encoding="utf-8") as f:
        json.dump(all_json, f, ensure_ascii=False, indent=2)
        
    today_folder = datetime.now().strftime("%Y %m%d")
    today_mmdd = datetime.now().strftime("%m%d")
    hist_dir = os.path.join(HISTORY_DIR, today_folder)
    os.makedirs(hist_dir, exist_ok=True)
    
    import openpyxl
    hu_excel_path = os.path.join(hist_dir, f"price tracker_HU_{today_folder}.xlsx")
    wb_hu = openpyxl.Workbook()
    ws_sec = wb_hu.active
    ws_sec.title = "MediaMarkt_HU_Samsung"
    write_standard_sheet(ws_sec, clean_sec, "MediaMarkt_HU_Samsung")
    ws_lg = wb_hu.create_sheet(title="MediaMarkt_HU_LG")
    write_standard_sheet(ws_lg, clean_lg, "MediaMarkt_HU_LG")
    wb_hu.save(hu_excel_path)
    print(f"💾 [SAVED] Standalone Hungary Excel: {hu_excel_path}")
    
    eu_wb_path = os.path.join(DATA_DIR, f"price tracker_EU_2026 {today_mmdd}_v1.xlsx")
    if not os.path.exists(eu_wb_path):
        import glob, shutil
        existing = sorted(glob.glob(os.path.join(DATA_DIR, "price tracker_EU_2026 *.xlsx")))
        existing = [f for f in existing if not os.path.basename(f).startswith("~$")]
        if existing:
            shutil.copy2(existing[-1], eu_wb_path)
            print(f"[CLONE] Created today's EU workbook from {existing[-1]} -> {eu_wb_path}")
            
    if os.path.exists(eu_wb_path):
        try:
            wb_eu = openpyxl.load_workbook(eu_wb_path)
            for sname in ["MediaMarkt_HU_Samsung", "MediaMarkt_HU_LG"]:
                if sname in wb_eu.sheetnames:
                    del wb_eu[sname]
            ws_eu_sec = wb_eu.create_sheet(title="MediaMarkt_HU_Samsung")
            write_standard_sheet(ws_eu_sec, clean_sec, "MediaMarkt_HU_Samsung")
            ws_eu_lg = wb_eu.create_sheet(title="MediaMarkt_HU_LG")
            write_standard_sheet(ws_eu_lg, clean_lg, "MediaMarkt_HU_LG")
            wb_eu.save(eu_wb_path)
            print(f"✅ [MERGED] Successfully updated Pan-European Excel with Hungary sheets: {eu_wb_path}")
        except Exception as e:
            print(f"⚠️ Error merging Hungary into EU Excel: {e}")

if __name__ == '__main__':
    dump_p = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\harry.park\.gemini\antigravity\brain\1c8f1fcd-c6b2-4806-abbc-780188e26593\.system_generated\steps\920\output.txt"
    process_hungary_dump(dump_p)
