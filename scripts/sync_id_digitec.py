import sys
import os
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
import re
import json
import glob
import parse_digitec

def find_latest_file(pattern_prefix):
    data_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
    files = glob.glob(os.path.join(data_dir, pattern_prefix + "*.xlsx"))
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    if not files:
        return None
    # Sort files by name so v2 is selected over v1, etc.
    files.sort()
    return files[-1]


def load_digitec_live():
    # steps/729/output.txt (Samsung) & steps/731/output.txt (LG) 에서 데이터 로드
    try:
        samsungs = parse_digitec.parse_digitec_file("C:/Users/harry.park/.gemini/antigravity/brain/f7b3e514-69f3-40c5-9d95-806bc168c4c7/.system_generated/steps/729/output.txt", "Samsung")
        lgs = parse_digitec.parse_digitec_file("C:/Users/harry.park/.gemini/antigravity/brain/f7b3e514-69f3-40c5-9d95-806bc168c4c7/.system_generated/steps/731/output.txt", "LG")
    except Exception as e:
        print(f"[WARN] Failed to load digitec raw files: {e}")
        samsungs, lgs = [], []
    
    # 27인치 등 모니터 및 라이프스타일 27인치 제외 (TV 조사 대상 아님)
    s_filtered = [p for p in samsungs if p["size"] >= 32 and not p["is_partner"]]
    l_filtered = [p for p in lgs if p["size"] >= 32 and not p["is_partner"]]
    
    return s_filtered, l_filtered

def load_interdiscount_live():
    data_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
    try:
        with open(os.path.join(data_dir, "raw_interdiscount_samsung.json"), "r", encoding="utf-8") as f:
            samsungs = json.load(f)
        with open(os.path.join(data_dir, "raw_interdiscount_lg.json"), "r", encoding="utf-8") as f:
            lgs = json.load(f)
    except Exception as e:
        print(f"[WARN] Failed to load interdiscount raw files: {e}")
        samsungs, lgs = [], []
        
    s_filtered = [p for p in samsungs if p["size"] >= 32]
    l_filtered = [p for p in lgs if p["size"] >= 32]
    
    return s_filtered, l_filtered

def save_workbook_with_retry(wb, filename):
    base, ext = os.path.splitext(filename)
    version = 1
    current_filename = filename
    saved = False
    
    while not saved:
        try:
            wb.save(current_filename)
            saved = True
            print(f"  ➔ [SAVED] Successfully saved to: {current_filename}")
        except PermissionError:
            v_match = re.search(r'(_v|v)(\d+)$', base)
            if v_match:
                prefix = v_match.group(1)
                num = int(v_match.group(2)) + 1
                base_clean = base[:v_match.start()]
                current_filename = f"{base_clean}{prefix}{num}{ext}"
                base = f"{base_clean}{prefix}{num}"
            else:
                current_filename = f"{base}_v{version}{ext}"
                version += 1
            print(f"  ➔ [LOCKED] Permission denied for {filename}. Retrying with: {current_filename}...")
            
    return current_filename

def sync_retailer_sheet(filename, sheetname, brand, products):
    if not os.path.exists(filename):
        print(f"[ERROR] {filename} not found.")
        return filename
        
    print(f"\n[SYNC] Loading {filename} -> {sheetname}...")
    wb = openpyxl.load_workbook(filename)
    sheet = wb[sheetname]
    
    # 해당 Brand의 기존 행 삭제
    removed_count = 0
    for r in range(sheet.max_row, 1, -1):
        brand_val = sheet.cell(row=r, column=1).value
        if brand_val == brand:
            sheet.delete_rows(r)
            removed_count += 1
            
    print(f"  ➔ Removed {removed_count} obsolete {brand} rows.")
    
    # 새로운 라이브 모델 추가
    added_count = 0
    for p in products:
        # (Brand, Year, Display Type, Size, Model Code, Price, Shipping, Installment, Cashback, Promo)
        # 디스플레이 유형 매핑
        display_type = "OLED"
        code_upper = p["model_code"].upper()
        if "OLED" in code_upper or "S90" in code_upper or "S95" in code_upper or "S85" in code_upper or "S99" in code_upper:
            display_type = "OLED"
        elif "QNED" in code_upper:
            display_type = "QNED"
        elif "QN" in code_upper or "Q8" in code_upper or "Q7" in code_upper:
            display_type = "QLED"
        elif "U8" in code_upper or "UA" in code_upper or "UT" in code_upper or "UR" in code_upper:
            display_type = "UHD 4K"
            
        sheet.append([
            p["brand"],
            p["year"],
            display_type,
            p["size"],
            p["model_code"],
            p["price"],
            "Free",
            None,
            0,
            p.get("promo", "None")
        ])
        added_count += 1
        
    print(f"  ➔ Added {added_count} live {brand} rows.")
    
    final_filename = save_workbook_with_retry(wb, filename)
    wb.close()
    return final_filename

def find_matching_product(series_key, year, brand, products):
    # series_key 예: "77G5", "65S90H", "75QNED86B", "65QN70h", "75U8000H"
    # size와 series_key의 family 이름을 기준으로 매칭
    
    # 1. size 파싱 (첫 몇글자 숫자)
    size_match = re.match(r'^(\d+)', series_key)
    if not size_match:
        return None
    size = int(size_match.group(1))
    
    family = series_key[len(size_match.group(0)):].strip().upper()
    # "QN70H" -> "QN70"
    if family.endswith("H") or family.endswith("F"):
        family = family[:-1]
        
    for p in products:
        if p["size"] != size or p["brand"].upper() != brand.upper():
            continue
            
        if p.get("year") != year:
            continue
            
        code = p["model_code"].upper()
        # 가족명이 코드 내부에 포함되어 있는지 검사 (예: QN70F -> QN70)
        # 단, OLED vs QNED 교차 오매핑을 방지하기 위해 QNED는 따로 걸러야 함
        if "QNED" in family and "QNED" not in code:
            continue
        if "QNED" not in family and "QNED" in code:
            continue
            
        # 매칭 조건 정의
        is_match = False
        if family == "G5":
            is_match = "G5" in code
        elif family == "G6":
            is_match = "G6" in code
        elif family == "C5":
            is_match = "C5" in code
        elif family == "C6":
            is_match = "C6" in code
        elif family == "B5":
            is_match = "B5" in code
        elif family == "B6":
            is_match = "B6" in code
        elif family == "QNED86A":
            is_match = "QNED86" in code or "QNED85" in code
        elif family == "QNED86B":
            is_match = "QNED86" in code
        elif family == "QNED80A":
            is_match = "QNED80" in code
        elif family == "QNED80B":
            is_match = "QNED80" in code or "QNED7E" in code or "QNED72" in code
        elif family == "UA75":
            is_match = "UA75" in code or "UA73" in code
        elif family == "UA77":
            is_match = "UA77" in code
        elif family == "S99H":
            is_match = "S99H" in code or "S99" in code
        elif family == "S95F" or family == "S95":
            is_match = "S95F" in code or "S95" in code
        elif family == "S95H":
            is_match = "S95H" in code
        elif family == "S90F" or family == "S90":
            is_match = "S90F" in code or "S90" in code
        elif family == "S90H":
            is_match = "S90H" in code
        elif family == "S85F" or family == "S85":
            is_match = "S85F" in code or "S85" in code
        elif family == "S85H":
            is_match = "S85H" in code
        elif family == "QN70":
            is_match = "QN70" in code
        elif family == "Q8":
            is_match = "Q8" in code
        elif family == "Q7":
            is_match = "Q7" in code
        elif family == "U8000":
            is_match = "U8000" in code or "U8090" in code or "U80" in code
            
        if is_match:
            return p
            
    return None

def sync_ata_sheet(filename, sheetname, year, id_products, digi_products):
    if not os.path.exists(filename):
        print(f"[ERROR] {filename} not found.")
        return filename
        
    print(f"\n[SYNC ATA] Loading {filename} -> {sheetname}...")
    wb = openpyxl.load_workbook(filename)
    sheet = wb[sheetname]
    
    # 4행부터 85행까지 순회
    updated_id = 0
    updated_digi = 0
    
    for r in range(4, 86):
        series_val = sheet.cell(row=r, column=2).value
        if not series_val:
            continue
            
        series_key = str(series_val).strip()
        # brand 판별
        # LG 시리즈인 경우: G5, C5, B5, QNED86A, QNED80A, UA75, G6, C6, B6, QNED86B, QNED80B, UA77
        is_lg = any(s in series_key for s in ["G5", "C5", "B5", "QNED86A", "QNED80A", "UA75", "G6", "C6", "B6", "QNED86B", "QNED80B", "UA77"])
        brand = "LG" if is_lg else "Samsung"
        
        # 1. Interdiscount (H: Model, J: Price, K: Cashback)
        id_match = find_matching_product(series_key, year, brand, id_products)
        if id_match:
            sheet.cell(row=r, column=8).value = id_match["model_code"]
            sheet.cell(row=r, column=10).value = id_match["price"]
            sheet.cell(row=r, column=11).value = 0
            updated_id += 1
        else:
            sheet.cell(row=r, column=8).value = ""
            sheet.cell(row=r, column=10).value = 0
            sheet.cell(row=r, column=11).value = 0
            
        # 2. Digitec (M: Model, O: Price, P: Cashback)
        digi_match = find_matching_product(series_key, year, brand, digi_products)
        if digi_match:
            sheet.cell(row=r, column=13).value = digi_match["model_code"]
            sheet.cell(row=r, column=15).value = digi_match["price"]
            sheet.cell(row=r, column=16).value = 0
            updated_digi += 1
        else:
            sheet.cell(row=r, column=13).value = ""
            sheet.cell(row=r, column=15).value = 0
            sheet.cell(row=r, column=16).value = 0
            
    print(f"  ➔ Updated Interdiscount: {updated_id} rows | Digitec: {updated_digi} rows.")
    final_filename = save_workbook_with_retry(wb, filename)
    wb.close()
    return final_filename

if __name__ == "__main__":
    digi_s, digi_l = load_digitec_live()
    id_s, id_l = load_interdiscount_live()
    
    print(f"[INFO] Loaded Digitec: Samsung={len(digi_s)}, LG={len(digi_l)}")
    print(f"[INFO] Loaded Interdiscount: Samsung={len(id_s)}, LG={len(id_l)}")
    
    # 1. price tracker 파일 업데이트 (Digitec, Interdiscount)
    data_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
    pt_file = find_latest_file("price tracker_swiss_2026 0703") or os.path.join(data_dir, "price tracker_swiss_2026 0703_v1.xlsx")
    pt_file = sync_retailer_sheet(pt_file, "Digitec_Samsung_Full", "Samsung", digi_s)
    pt_file = sync_retailer_sheet(pt_file, "Digitec_LG_Full", "LG", digi_l)
    pt_file = sync_retailer_sheet(pt_file, "Interdiscount_Samsung_Full", "Samsung", id_s)
    pt_file = sync_retailer_sheet(pt_file, "Interdiscount_LG_Full", "LG", id_l)
    
    # 2. 통합 파일 업데이트
    comp_file = find_latest_file("TV_Price_Comparison") or os.path.join(data_dir, "TV_Price_Comparison_v5.xlsx")
    comp_file = sync_retailer_sheet(comp_file, "Digitec_Samsung_Full", "Samsung", digi_s)
    comp_file = sync_retailer_sheet(comp_file, "Digitec_LG_Full", "LG", digi_l)
    comp_file = sync_retailer_sheet(comp_file, "Interdiscount_Samsung_Full", "Samsung", id_s)
    comp_file = sync_retailer_sheet(comp_file, "Interdiscount_LG_Full", "LG", id_l)
    
    # 3. ATA 비교 시트 업데이트 (Swiss_2025, Swiss_2026)
    ata_file = find_latest_file("Swiss_ATA_Comparison_2026_0703") or os.path.join(data_dir, "Swiss_ATA_Comparison_2026_0703.xlsx")
    ata_file = sync_ata_sheet(ata_file, "Swiss_2025", 2025, id_s + id_l, digi_s + digi_l)
    ata_file = sync_ata_sheet(ata_file, "Swiss_2026", 2026, id_s + id_l, digi_s + digi_l)
