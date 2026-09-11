import openpyxl
import os

swiss_file = r"data\price tracker_swiss_2026 0814_v1.xlsx"
eu_file = r"data\price tracker_EU_2026 0814_v1.xlsx"

print("==================================================")
print("1. Inspecting Swiss Excel:", swiss_file)
print("==================================================")

if os.path.exists(swiss_file):
    wb = openpyxl.load_workbook(swiss_file, data_only=True)
    print("Sheets:", wb.sheetnames)
    for s in wb.sheetnames:
        ws = wb[s]
        rows = list(ws.iter_rows(values_only=True))
        if not rows: continue
        headers = [str(c).strip() if c is not None else "" for c in rows[0]]
        
        cb_col = -1
        promo_col = -1
        price_col = -1
        model_col = -1
        for idx, h in enumerate(headers):
            if "cashback" in h.lower() or "캐시백" in h.lower(): cb_col = idx
            if "promo" in h.lower() or "혜택" in h.lower() or "promotion" in h.lower(): promo_col = idx
            if "price" in h.lower() or "가격" in h.lower(): price_col = idx
            if "model" in h.lower() or "모델" in h.lower(): model_col = idx
            
        cb_count = 0
        promo_count = 0
        all_cb_entries = []
        all_promo_entries = []
        
        for r in rows[1:]:
            cb_val = r[cb_col] if cb_col >= 0 and cb_col < len(r) else None
            pr_val = r[promo_col] if promo_col >= 0 and promo_col < len(r) else None
            m_val = r[model_col] if model_col >= 0 and model_col < len(r) else None
            p_val = r[price_col] if price_col >= 0 and price_col < len(r) else None
            
            if cb_val is not None and str(cb_val).strip() not in ["", "0", "0.0", "None"]:
                cb_count += 1
                all_cb_entries.append((m_val, p_val, cb_val, pr_val))
            if pr_val is not None and str(pr_val).strip() not in ["", "None"]:
                promo_count += 1
                all_promo_entries.append((m_val, p_val, cb_val, pr_val))
                
        print(f"\n[Sheet: {s}] Total Rows: {len(rows)-1} | Cashback > 0: {cb_count} | Promo Text: {promo_count}")
        if all_cb_entries:
            print("  Cashback items:")
            for m, p, cb, pr in all_cb_entries[:10]:
                print(f"    - {m}: Price={p}, Cashback={cb}, Promo={pr}")
        if all_promo_entries:
            print("  Promo text items:")
            for m, p, cb, pr in all_promo_entries[:10]:
                print(f"    - {m}: Price={p}, Cashback={cb}, Promo={pr}")

print("\n==================================================")
print("2. Inspecting EU Excel (Swiss Sheets):", eu_file)
print("==================================================")
if os.path.exists(eu_file):
    wb_eu = openpyxl.load_workbook(eu_file, data_only=True)
    swiss_sheets = [s for s in wb_eu.sheetnames if "CH" in s or "Swiss" in s]
    print("Swiss sheets in EU workbook:", swiss_sheets)
    for s in swiss_sheets:
        ws = wb_eu[s]
        rows = list(ws.iter_rows(values_only=True))
        if not rows: continue
        headers = [str(c).strip() if c is not None else "" for c in rows[0]]
        print(f"\n[EU Sheet: {s}] Rows: {len(rows)-1}")
        for idx, h in enumerate(headers):
            print(f"  Col {idx+1}: {h}")
