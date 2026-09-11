import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb_eu_path = 'data/price tracker_EU_2026 0901_v1.xlsx'
wb_ch_path = 'data/price tracker_swiss_2026 0901_v1.xlsx'

def audit_workbook(path, name):
    if not os.path.exists(path):
        return
    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"==================================================")
    print(f" 🔍 AUDITING WORKBOOK: {name}")
    print(f"==================================================")
    
    anomalies = []
    
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        currency = "EUR"
        if "UK" in sheet or "Currys" in sheet: currency = "GBP"
        elif "CH" in sheet or "Swiss" in sheet or "Digitec" in sheet or "Interdiscount" in sheet: currency = "CHF"
        elif "CZ" in sheet or "Alza" in sheet: currency = "CZK"
        elif "HU" in sheet or "Hungary" in sheet: currency = "HUF"
        
        for r in range(2, ws.max_row + 1):
            year = ws.cell(r, 2).value
            disp = str(ws.cell(r, 3).value or "LED")
            size = ws.cell(r, 4).value
            mc = str(ws.cell(r, 5).value or "")
            price = ws.cell(r, 6).value
            promo = str(ws.cell(r, 10).value or "")
            
            if not mc or not price:
                continue
                
            try:
                pval = float(price)
                sval = int(size) if size and str(size).isdigit() else 55
                d_up = disp.upper()
                
                # Check for anomalies in EUR / CHF / GBP
                is_anom = False
                reason = ""
                
                if currency in ["EUR", "CHF", "GBP"]:
                    # 1. OLED Anomaly
                    if "OLED" in d_up or "OLED" in mc.upper():
                        if sval >= 83 and pval < 1800: is_anom = True; reason = f"83\" OLED price too low ({currency} {pval})"
                        elif sval >= 77 and pval < 1300: is_anom = True; reason = f"77\" OLED price too low ({currency} {pval})"
                        elif sval >= 65 and pval < 850: is_anom = True; reason = f"65\" OLED price too low ({currency} {pval})"
                        elif sval >= 55 and pval < 650: is_anom = True; reason = f"55\" OLED price too low ({currency} {pval})"
                        elif sval in [42, 48] and pval < 500: is_anom = True; reason = f"42/48\" OLED price too low ({currency} {pval})"
                        
                    # 2. Micro RGB Anomaly
                    elif "MRGB" in d_up or "MICRO RGB" in d_up or "MRGB" in mc.upper() or "R85" in mc.upper() or "R95" in mc.upper():
                        if sval >= 75 and pval < 1500: is_anom = True; reason = f"75\"+ MRGB price too low ({currency} {pval})"
                        elif sval >= 50 and pval < 750: is_anom = True; reason = f"55/65\" MRGB price too low ({currency} {pval})"
                        
                    # 3. QNED / QLED / Neo QLED Anomaly
                    elif any(x in d_up for x in ["QNED", "QLED", "NEO QLED"]) or any(x in mc.upper() for x in ["QNED", "QN8", "QN9", "QN7", "LS03"]):
                        if sval >= 75 and pval < 600: is_anom = True; reason = f"75\"+ QNED/QLED price too low ({currency} {pval})"
                        elif sval >= 65 and pval < 450: is_anom = True; reason = f"65\" QNED/QLED price too low ({currency} {pval})"
                        elif sval >= 50 and pval < 280: is_anom = True; reason = f"50/55\" QNED/QLED price too low ({currency} {pval})"
                        elif sval >= 43 and pval < 200: is_anom = True; reason = f"43\" QNED/QLED price too low ({currency} {pval})"
                        
                    # 4. UHD 4K General Anomaly
                    elif sval >= 55 and pval < 150: is_anom = True; reason = f"55\"+ TV price too low ({currency} {pval})"
                    elif sval >= 43 and pval < 100: is_anom = True; reason = f"43\"+ TV price too low ({currency} {pval})"
                    
                elif currency == "CZK": # CZK ~ 25 per EUR
                    if ("OLED" in d_up or "OLED" in mc.upper()) and sval >= 55 and pval < 15000:
                        is_anom = True; reason = f"55\"+ OLED price too low (CZK {pval})"
                        
                elif currency == "HUF": # HUF ~ 400 per EUR
                    if ("OLED" in d_up or "OLED" in mc.upper()) and sval >= 55 and pval < 250000:
                        is_anom = True; reason = f"55\"+ OLED price too low (HUF {pval})"
                        
                if is_anom:
                    anomalies.append({
                        "sheet": sheet,
                        "year": year,
                        "disp": disp,
                        "size": sval,
                        "mc": mc,
                        "price": pval,
                        "currency": currency,
                        "reason": reason,
                        "promo": promo
                    })
            except Exception:
                pass
                
    print(f"Total Anomalies Found in {name}: {len(anomalies)}")
    for a in anomalies:
        print(f"  🔻 [{a['sheet']:<28}] {a['size']:<2}\" {a['disp']:<10} {a['mc']:<18} : {a['currency']} {a['price']:>8.2f} | Reason: {a['reason']}")
    print()

audit_workbook(wb_eu_path, "Pan-European Workbook")
audit_workbook(wb_ch_path, "Swiss Workbook")
