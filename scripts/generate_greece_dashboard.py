# -*- coding: utf-8 -*-
import openpyxl
import json
import os
import re
import sys
import glob
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

def clean_price(val):
    if val is None or val == "" or val == "None":
        return 0
    if isinstance(val, (int, float)):
        fval = float(val)
        return int(round(fval)) if fval >= 50.0 else 0
        
    s = str(val).replace("€", "").replace("EUR", "").strip()
    if re.match(r'^\d{1,3}\.\d{3}$', s):
        s = s.replace(".", "")
    elif re.match(r'^\d{1,3},\d{3}$', s):
        s = s.replace(",", "")
    elif "," in s and "." in s:
        if s.find(".") < s.find(","):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        parts = s.split(",")
        if len(parts[-1]) == 2:
            s = s.replace(",", ".")
        else:
            s = s.replace(",", "")
            
    cleaned = re.sub(r'[^\d.]', '', s)
    try:
        fval = float(cleaned) if cleaned else 0.0
        return int(round(fval)) if fval >= 50.0 else 0
    except ValueError:
        return 0

def load_excel_data(excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    all_products = []
    
    for sheet_name in wb.sheetnames:
        if sheet_name in ["Greece_2025", "Greece_2026"]:
            continue
            
        sheet = wb[sheet_name]
        parts = sheet_name.split("_")
        retailer = parts[0] if len(parts) > 0 else "Greece"
        
        for r in range(2, sheet.max_row + 1):
            brand = sheet.cell(row=r, column=1).value
            if not brand:
                continue
                
            year = sheet.cell(row=r, column=2).value
            display_type = sheet.cell(row=r, column=3).value
            size = sheet.cell(row=r, column=4).value
            model_code = sheet.cell(row=r, column=5).value
            price = clean_price(sheet.cell(row=r, column=6).value)
            cashback = clean_price(sheet.cell(row=r, column=9).value)
            promo = str(sheet.cell(row=r, column=10).value or "None")
            
            if price < 50:
                continue
                
            all_products.append({
                "retailer": retailer,
                "brand": str(brand).upper(),
                "year": int(year) if year else 2026,
                "display_type": str(display_type or "LED"),
                "size": int(size) if size else 55,
                "model_code": str(model_code or "Unknown"),
                "price": price,
                "cashback": cashback,
                "net_price": price - cashback,
                "promo": promo
            })
            
    wb.close()
    return all_products

def generate_dashboard_html(products, output_html_path):
    json_payload = json.dumps(products, ensure_ascii=False, indent=2)
    
    html_content = f"""<!DOCTYPE html>
<html lang="el">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Greece TV Price Tracker & Promo Dashboard (Kotsovolos vs Public)</title>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&display=swap" rel="stylesheet">
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <style>
        :root {{
            --bg-main: #0f172a;
            --bg-card: #1e293b;
            --bg-card-hover: #334155;
            --accent-blue: #3b82f6;
            --accent-purple: #8b5cf6;
            --text-main: #f8fafc;
            --text-muted: #94a3b8;
            --border-color: #334155;
            --samsung-blue: #2563eb;
            --lg-red: #ef4444;
            --kotsovolos-red: #dc2626;
            --public-yellow: #f59e0b;
        }}

        * {{
            box-sizing: border-box;
            margin: 0;
            padding: 0;
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif;
        }}

        body {{
            background-color: var(--bg-main);
            color: var(--text-main);
            padding: 24px;
            min-height: 100vh;
        }}

        header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            padding-bottom: 24px;
            border-bottom: 1px solid var(--border-color);
            margin-bottom: 24px;
        }}

        .logo-area h1 {{
            font-size: 26px;
            font-weight: 800;
            background: linear-gradient(135deg, #60a5fa 0%, #a78bfa 100%);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
        }}

        .logo-area p {{
            font-size: 14px;
            color: var(--text-muted);
            margin-top: 4px;
        }}

        .stats-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 16px;
            margin-bottom: 24px;
        }}

        .stat-card {{
            background: var(--bg-card);
            border: 1px solid var(--border-color);
            border-radius: 12px;
            padding: 20px;
            display: flex;
            flex-direction: column;
        }}

        .stat-card .label {{
            font-size: 13px;
            color: var(--text-muted);
            font-weight: 500;
        }}

        .stat-card .value {{
            font-size: 28px;
            font-weight: 700;
            margin-top: 8px;
            color: #ffffff;
        }}

        .filter-bar {{
            background: var(--bg-card);
            border: 1px solid var(--border-color);
            border-radius: 12px;
            padding: 16px 20px;
            display: flex;
            flex-wrap: wrap;
            gap: 16px;
            align-items: center;
            margin-bottom: 24px;
        }}

        .filter-group {{
            display: flex;
            align-items: center;
            gap: 8px;
        }}

        .filter-group label {{
            font-size: 13px;
            font-weight: 600;
            color: var(--text-muted);
        }}

        .filter-select, .filter-input {{
            background: #0f172a;
            border: 1px solid var(--border-color);
            color: var(--text-main);
            padding: 8px 14px;
            border-radius: 8px;
            font-size: 14px;
            outline: none;
            transition: all 0.2s ease;
        }}

        .filter-select:focus, .filter-input:focus {{
            border-color: var(--accent-blue);
        }}

        .chart-section {{
            background: var(--bg-card);
            border: 1px solid var(--border-color);
            border-radius: 12px;
            padding: 24px;
            margin-bottom: 24px;
        }}

        .chart-header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 20px;
        }}

        .chart-title {{
            font-size: 18px;
            font-weight: 700;
        }}

        .chart-container {{
            position: relative;
            height: 380px;
            width: 100%;
        }}

        .table-section {{
            background: var(--bg-card);
            border: 1px solid var(--border-color);
            border-radius: 12px;
            overflow: hidden;
        }}

        .table-header {{
            padding: 20px 24px;
            border-bottom: 1px solid var(--border-color);
            display: flex;
            justify-content: space-between;
            align-items: center;
        }}

        table {{
            width: 100%;
            border-collapse: collapse;
            text-align: left;
        }}

        th {{
            background: #111827;
            color: var(--text-muted);
            font-weight: 600;
            font-size: 12px;
            text-transform: uppercase;
            letter-spacing: 0.05em;
            padding: 14px 20px;
            border-bottom: 1px solid var(--border-color);
        }}

        td {{
            padding: 16px 20px;
            border-bottom: 1px solid var(--border-color);
            font-size: 14px;
            color: var(--text-main);
        }}

        tr:hover {{
            background: var(--bg-card-hover);
        }}

        .badge-brand {{
            padding: 4px 10px;
            border-radius: 6px;
            font-size: 12px;
            font-weight: 700;
            display: inline-block;
        }}

        .badge-samsung {{ background: rgba(37, 99, 235, 0.2); color: #60a5fa; border: 1px solid rgba(37, 99, 235, 0.4); }}
        .badge-lg {{ background: rgba(239, 68, 68, 0.2); color: #f87171; border: 1px solid rgba(239, 68, 68, 0.4); }}
        
        .badge-retailer {{
            padding: 4px 10px;
            border-radius: 6px;
            font-size: 12px;
            font-weight: 600;
        }}
        .badge-kotsovolos {{ background: rgba(220, 38, 38, 0.2); color: #fca5a5; }}
        .badge-public {{ background: rgba(245, 158, 11, 0.2); color: #fcd34d; }}

        .price-text {{
            font-weight: 700;
            color: #38bdf8;
        }}
    </style>
</head>
<body>

    <header>
        <div class="logo-area">
            <h1>🇬🇷 Greece TV Price Tracker & Promo Dashboard</h1>
            <p>Kotsovolos vs Public Retailer Pricing & Promotional Analysis (2025/2026 Models)</p>
        </div>
        <div style="font-size: 13px; color: var(--text-muted);">
            Last Updated: <strong>{datetime.now().strftime("%Y-%m-%d %H:%M")}</strong>
        </div>
    </header>

    <div class="stats-grid">
        <div class="stat-card">
            <span class="label">Total Active Models</span>
            <span class="value" id="stat-total">0</span>
        </div>
        <div class="stat-card">
            <span class="label">Kotsovolos Offerings</span>
            <span class="value" id="stat-kotsovolos">0</span>
        </div>
        <div class="stat-card">
            <span class="label">Public Offerings</span>
            <span class="value" id="stat-public">0</span>
        </div>
        <div class="stat-card">
            <span class="label">Average Selling Price</span>
            <span class="value" id="stat-avg-price">€0</span>
        </div>
    </div>

    <div class="filter-bar">
        <div class="filter-group">
            <label>Model Year:</label>
            <select id="filter-year" class="filter-select" onchange="renderDashboard()">
                <option value="ALL">All Years (2025 & 2026)</option>
                <option value="2026">2026 Models</option>
                <option value="2025">2025 Models</option>
            </select>
        </div>
        <div class="filter-group">
            <label>Retailer:</label>
            <select id="filter-retailer" class="filter-select" onchange="renderDashboard()">
                <option value="ALL">All Retailers (Kotsovolos & Public)</option>
                <option value="Kotsovolos">Kotsovolos</option>
                <option value="Public">Public</option>
            </select>
        </div>
        <div class="filter-group">
            <label>Brand:</label>
            <select id="filter-brand" class="filter-select" onchange="renderDashboard()">
                <option value="ALL">All Brands (Samsung & LG)</option>
                <option value="SAMSUNG">Samsung</option>
                <option value="LG">LG</option>
            </select>
        </div>
        <div class="filter-group">
            <label>Display Type:</label>
            <select id="filter-display" class="filter-select" onchange="renderDashboard()">
                <option value="ALL">All Display Types</option>
                <option value="OLED">OLED / OLED evo</option>
                <option value="QNED">QNED / Neo QLED / QLED</option>
                <option value="UHD 4K">UHD 4K</option>
            </select>
        </div>
        <div class="filter-group" style="flex: 1;">
            <label>Search:</label>
            <input type="text" id="filter-search" class="filter-input" placeholder="Search model code or title..." oninput="renderDashboard()" style="width: 100%;">
        </div>
    </div>

    <div class="chart-section">
        <div class="chart-header">
            <span class="chart-title">Retailer Price Comparison (Kotsovolos vs Public)</span>
        </div>
        <div class="chart-container">
            <canvas id="priceChart"></canvas>
        </div>
    </div>

    <div class="table-section">
        <div class="table-header">
            <span class="chart-title">Active Model Inventory & Pricing</span>
        </div>
        <table>
            <thead>
                <tr>
                    <th>Retailer</th>
                    <th>Brand</th>
                    <th>Year</th>
                    <th>Display</th>
                    <th>Size</th>
                    <th>Model Code</th>
                    <th>Selling Price</th>
                    <th>Cashback</th>
                    <th>Promotions</th>
                </tr>
            </thead>
            <tbody id="table-body">
            </tbody>
        </table>
    </div>

    <script>
        const RAW_DATA = {json_payload};
        let chartInstance = null;

        function renderDashboard() {{
            const selectedYear = document.getElementById('filter-year').value;
            const selectedRetailer = document.getElementById('filter-retailer').value;
            const selectedBrand = document.getElementById('filter-brand').value;
            const selectedDisplay = document.getElementById('filter-display').value;
            const searchText = document.getElementById('filter-search').value.toLowerCase().strip();

            let filtered = RAW_DATA.filter(item => {{
                if (selectedYear !== 'ALL' && item.year.toString() !== selectedYear) return false;
                if (selectedRetailer !== 'ALL' && item.retailer !== selectedRetailer) return false;
                if (selectedBrand !== 'ALL' && item.brand !== selectedBrand) return false;
                if (selectedDisplay !== 'ALL') {{
                    if (selectedDisplay === 'OLED' && !item.display_type.includes('OLED')) return false;
                    if (selectedDisplay === 'QNED' && !item.display_type.includes('QNED') && !item.display_type.includes('QLED')) return false;
                    if (selectedDisplay === 'UHD 4K' && !item.display_type.includes('UHD')) return false;
                }}
                if (searchText) {{
                    const matchTitle = item.model_code.toLowerCase().includes(searchText);
                    if (!matchTitle) return false;
                }}
                return true;
            }});

            // Stats
            document.getElementById('stat-total').innerText = filtered.length;
            const kotsoCount = filtered.filter(i => i.retailer === 'Kotsovolos').length;
            const publicCount = filtered.filter(i => i.retailer === 'Public').length;
            document.getElementById('stat-kotsovolos').innerText = kotsoCount;
            document.getElementById('stat-public').innerText = publicCount;
            
            const avgPrice = filtered.length > 0 ? Math.round(filtered.reduce((a, b) => a + b.price, 0) / filtered.length) : 0;
            document.getElementById('stat-avg-price').innerText = '€' + avgPrice.toLocaleString();

            // Table
            const tbody = document.getElementById('table-body');
            tbody.innerHTML = '';
            filtered.forEach(item => {{
                const tr = document.createElement('tr');
                tr.innerHTML = `
                    <td><span class="badge-retailer badge-${{item.retailer.toLowerCase()}}">${{item.retailer}}</span></td>
                    <td><span class="badge-brand badge-${{item.brand.toLowerCase()}}">${{item.brand}}</span></td>
                    <td>${{item.year}}</td>
                    <td>${{item.display_type}}</td>
                    <td>${{item.size}}"</td>
                    <td><strong>${{item.model_code}}</strong></td>
                    <td class="price-text">€${{item.price.toLocaleString()}}</td>
                    <td>${{item.cashback > 0 ? '€' + item.cashback : '-'}}</td>
                    <td>${{item.promo || '-'}}</td>
                `;
                tbody.appendChild(tr);
            }});

            // Chart
            updateChart(filtered);
        }}

        function updateChart(data) {{
            const ctx = document.getElementById('priceChart').getContext('2d');
            
            // Group by model code
            const models = [...new Set(data.map(d => d.model_code))].slice(0, 15);
            const kotsoPrices = models.map(m => {{
                const item = data.find(d => d.model_code === m && d.retailer === 'Kotsovolos');
                return item ? item.price : 0;
            }});
            const publicPrices = models.map(m => {{
                const item = data.find(d => d.model_code === m && d.retailer === 'Public');
                return item ? item.price : 0;
            }});

            if (chartInstance) {{
                chartInstance.destroy();
            }}

            chartInstance = new Chart(ctx, {{
                type: 'bar',
                data: {{
                    labels: models,
                    datasets: [
                        {{
                            label: 'Kotsovolos (€)',
                            data: kotsoPrices,
                            backgroundColor: '#ef4444'
                        }},
                        {{
                            label: 'Public (€)',
                            data: publicPrices,
                            backgroundColor: '#f59e0b'
                        }}
                    ]
                }},
                options: {{
                    responsive: true,
                    maintainAspectRatio: false,
                    scales: {{
                        y: {{
                            beginAtZero: true,
                            grid: {{ color: '#334155' }},
                            ticks: {{ color: '#94a3b8' }}
                        }},
                        x: {{
                            grid: {{ color: '#334155' }},
                            ticks: {{ color: '#94a3b8' }}
                        }}
                    }},
                    plugins: {{
                        legend: {{ labels: {{ color: '#f8fafc' }} }}
                    }}
                }}
            }});
        }}

        window.onload = function() {{
            renderDashboard();
        }};
    </script>
</body>
</html>
"""
    with open(output_html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"[SUCCESS] Dashboard generated: {output_html_path}")

def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    root_dir = os.path.abspath(os.path.join(script_dir, ".."))
    data_dir = os.path.join(root_dir, "data")
    
    excel_files = glob.glob(os.path.join(data_dir, "price tracker_greece_2026 *.xlsx"))
    excel_files = [f for f in excel_files if not os.path.basename(f).startswith("~$")]
    if not excel_files:
        print("[WARN] No Greece excel file found in data directory.")
        return
        
    excel_files.sort()
    latest_excel = excel_files[-1]
    print(f"[DASHBOARD] Reading data from: {latest_excel}")
    
    products = load_excel_data(latest_excel)
    output_html = os.path.join(data_dir, "greece_price_dashboard.html")
    generate_dashboard_html(products, output_html)

if __name__ == '__main__':
    main()
