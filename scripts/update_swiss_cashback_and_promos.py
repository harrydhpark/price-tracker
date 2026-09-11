import os
import sys
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
import re
from swiss_promo_parser import parse_swiss_promo_and_cashback, clean_swiss_promo_text

def update_swiss_excel(file_path):
    if not os.path.exists(file_path):
        print(f"[ERROR] {file_path} does not exist.")
        return
        
    print(f"\n[UPDATING SWISS EXCEL] Processing: {file_path}")
    wb = openpyxl.load_workbook(file_path)
    
    total_updated = 0
    total_cb_found = 0
    total_promo_found = 0
    
    for s_name in wb.sheetnames:
        ws = wb[s_name]
        rows = list(ws.iter_rows(values_only=False))
        if not rows or len(rows) <= 1:
            continue
            
        headers = [str(c.value).strip() if c.value is not None else "" for c in rows[0]]
        
        brand_col = 0
        year_col = 1
        disp_col = 2
        size_col = 3
        code_col = 4
        price_col = 5
        cb_col = 8
        promo_col = 9
        
        for idx, h in enumerate(headers):
            h_low = h.lower()
            if "brand" in h_low or "브랜드" in h_low: brand_col = idx
            elif "year" in h_low or "연도" in h_low: year_col = idx
            elif "display" in h_low or "디스플레이" in h_low: disp_col = idx
            elif "size" in h_low or "인치" in h_low: size_col = idx
            elif "code" in h_low or "모델" in h_low: code_col = idx
            elif "price" in h_low or "가격" in h_low: price_col = idx
            elif "cashback" in h_low or "캐시백" in h_low: cb_col = idx
            elif "promo" in h_low or "혜택" in h_low: promo_col = idx
            
        sheet_cb = 0
        sheet_promo = 0
        
        for r_idx in range(1, len(rows)):
            row = rows[r_idx]
            brand_val = str(row[brand_col].value or "").strip()
            year_val = row[year_col].value
            try:
                year_int = int(year_val) if year_val else 2025
            except Exception:
                year_int = 2025
                
            size_val = row[size_col].value
            try:
                size_int = int(size_val) if size_val else 55
            except Exception:
                size_int = 55
                
            code_val = str(row[code_col].value or "").strip()
            price_val = row[price_col].value
            try:
                price_float = float(price_val) if price_val else 0.0
            except Exception:
                price_float = 0.0
                
            old_cb = row[cb_col].value if cb_col < len(row) else 0
            old_promo = row[promo_col].value if promo_col < len(row) else "None"
            
            cb_num, clean_promo = parse_swiss_promo_and_cashback(
                old_promo, "", brand_val, year_int, code_val, size_int, price_float
            )
            
            # Write updated values
            row[cb_col].value = cb_num if cb_num > 0 else 0
            row[promo_col].value = clean_promo if clean_promo != "None" else None
            
            if cb_num > 0:
                sheet_cb += 1
            if clean_promo != "None":
                sheet_promo += 1
            total_updated += 1
            
        print(f"  ➔ Sheet '{s_name}': {len(rows)-1} rows | Cashback > 0: {sheet_cb} | Promos: {sheet_promo}")
        total_cb_found += sheet_cb
        total_promo_found += sheet_promo
        
    wb.save(file_path)
    wb.close()
    print(f"[SUCCESS] Updated {file_path} (Total Rows: {total_updated}, Cashback > 0: {total_cb_found}, Promos: {total_promo_found})")

if __name__ == "__main__":
    swiss_file = r"data\price tracker_swiss_2026 0814_v1.xlsx"
    update_swiss_excel(swiss_file)
