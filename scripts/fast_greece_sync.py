# -*- coding: utf-8 -*-
import os
import sys
import re
import json
import glob
import shutil
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

DATA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
HISTORY_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "History_EU"))

def parse_price(price_str):
    if not price_str:
        return 0.0
    cleaned = str(price_str).replace('€', '').replace(' ', '').replace('\xa0', '').strip()
    if not cleaned:
        return 0.0
    if ',' in cleaned and '.' in cleaned:
        cleaned = cleaned.replace('.', '').replace(',', '.')
    elif ',' in cleaned:
        cleaned = cleaned.replace(',', '.')
    try:
        val = float(cleaned)
        return val if val >= 50.0 else 0.0
    except ValueError:
        return 0.0

def extract_model_code_and_info(title_upper, brand, href):
    year_val = None
    series_code = "Unknown"
    model_code = "Unknown"
    size_val = 0

    m_size = re.search(r'(\d{2,3})[\"\”\s]*(?:4K|Smart|OLED|QLED|LED|QNED|Nano|Mini|Micro|\b)', title_upper)
    if m_size:
        try:
            size_val = int(m_size.group(1))
        except ValueError:
            size_val = 0
            
    if size_val < 22:
        return None  # Exclude small screens, monitors and accessories

    if brand.lower() == "samsung":
        patterns_2026 = [
            (r'S99H', 'S99H', 2026), (r'S95H', 'S95H', 2026), (r'S90H', 'S90H', 2026), (r'S85H', 'S85H', 2026),
            (r'QN95H', 'QN95H', 2026), (r'QN90H', 'QN90H', 2026), (r'QN85H', 'QN85H', 2026),
            (r'QN80H', 'QN80H', 2026), (r'QN70H', 'QN70H', 2026), (r'M80H', 'M80H', 2026),
            (r'M70H', 'M70H', 2026), (r'R95H', 'R95H', 2026), (r'R86H', 'R86H', 2026), (r'R85H', 'R85H', 2026),
            (r'U8072H', 'U8000H', 2026), (r'U8070H', 'U8000H', 2026), (r'U8092H', 'U8000H', 2026),
            (r'U8072', 'U8000H', 2026), (r'U8092', 'U8000H', 2026), (r'U8000H', 'U8000H', 2026),
            (r'LS03HW', 'The Frame Pro', 2026), (r'LS03H', 'The Frame', 2026), (r'F6000H', 'F6000H', 2026)
        ]
        patterns_2025 = [
            (r'S95F', 'S95F', 2025), (r'S90F', 'S90F', 2025), (r'S85F', 'S85F', 2025),
            (r'QN95F', 'QN95F', 2025), (r'QN90F', 'QN90F', 2025), (r'QN85F', 'QN85F', 2025),
            (r'QN80F', 'QN80F', 2025), (r'QN70F', 'QN70F', 2025), (r'Q80F', 'Q8F', 2025),
            (r'Q8F', 'Q8F', 2025), (r'Q70F', 'Q7F', 2025), (r'Q7F', 'Q7F', 2025),
            (r'Q60F', 'Q6F', 2025), (r'Q6F', 'Q6F', 2025), (r'U8072F', 'U8000F', 2025), (r'U8000F', 'U8000F', 2025),
            (r'LS03F', 'The Frame', 2025)
        ]

        m_code = re.search(r'([A-Z]{2}\d{2}[A-Z0-9]+|\b\d{2}[A-Z]{1,3}\d{2,3}[A-Z0-9]*\b)', title_upper)
        if m_code:
            model_code = m_code.group(1)
        else:
            slug_m = re.search(r'/([a-z0-9\-]+)/?\d*$', href.lower())
            if slug_m:
                model_code = slug_m.group(1).upper()

        for pat, ser, yr in patterns_2026:
            if re.search(pat, title_upper):
                series_code = ser
                year_val = yr
                break

        if not year_val:
            for pat, ser, yr in patterns_2025:
                if re.search(pat, title_upper):
                    series_code = ser
                    year_val = yr
                    break

        if not year_val:
            if "2026" in title_upper or ("AI TV" in title_upper and any(x in title_upper for x in ["H ", "H\b", "95H", "90H", "85H", "80H", "70H"])):
                year_val = 2026
            elif "2025" in title_upper:
                year_val = 2025

    elif brand.lower() == "lg":
        lg_2026 = [
            (r'QNED93B', 'QNED93B', 2026), (r'QNED87B', 'QNED87B', 2026), (r'QNED86B', 'QNED86B', 2026),
            (r'QNED85B', 'QNED85B', 2026), (r'QNED83B', 'QNED81B', 2026), (r'QNED82B', 'QNED82B', 2026),
            (r'QNED81B', 'QNED81B', 2026), (r'QNED80B', 'QNED80B', 2026),
            (r'QNED72B', 'QNED72B', 2026), (r'QNED71B', 'QNED70B', 2026), (r'QNED70B', 'QNED70B', 2026),
            (r'QNED7EB', 'QNED7EB', 2026),
            (r'MRGB96B', 'MRGB96B', 2026), (r'MRGB88B', 'MRGB88B', 2026),
            (r'MRGB87B', 'MRGB87B', 2026), (r'MRGB86B', 'MRGB86B', 2026), (r'MRGB85B', 'MRGB85B', 2026),
            (r'LX7B', 'LX7B', 2026), (r'LX6', 'LX6', 2026), (r'27LX6', '27LX6', 2026), (r'STANBYME', 'STANBYME', 2026),
            (r'OLED\d{2}G6', 'OLED G6', 2026), (r'G6[0-9]L', 'OLED G6', 2026), (r'G6[0-9]', 'OLED G6', 2026),
            (r'OLED\d{2}C6', 'OLED C6', 2026), (r'C6[0-9]L', 'OLED C6', 2026), (r'C6[0-9]', 'OLED C6', 2026),
            (r'OLED\d{2}B6', 'OLED B6', 2026), (r'B6[0-9]L', 'OLED B6', 2026), (r'B6[0-9]', 'OLED B6', 2026),
            (r'OLED\d{2}W6', 'OLED W6', 2026), (r'W6[0-9]L', 'OLED W6', 2026),
            (r'NU90', 'NU90', 2026), (r'NU85', 'NU85', 2026), (r'NU8E', 'NU85', 2026),
            (r'UA77', 'UA77', 2026), (r'LB650B', 'LB650B', 2026)
        ]
        lg_2025 = [
            (r'QNED93A', 'QNED93A', 2025), (r'QNED87A', 'QNED87A', 2025), (r'QNED86A', 'QNED86A', 2025),
            (r'QNED85A', 'QNED85A', 2025), (r'QNED82A', 'QNED80A', 2025), (r'QNED81A', 'QNED81A', 2025),
            (r'QNED80A', 'QNED80A', 2025), (r'QNED72A', 'QNED72A', 2025), (r'QNED70A', 'QNED70A', 2025),
            (r'QNED7EA', 'QNED7EA', 2025),
            (r'MRGB87A', 'MRGB87A', 2025), (r'LX7A', 'LX7A', 2025), (r'LX5', 'LX5', 2025),
            (r'OLED\d{2}G5', 'OLED G5', 2025), (r'G5[0-9]L', 'OLED G5', 2025), (r'G5[0-9]', 'OLED G5', 2025),
            (r'OLED\d{2}C5', 'OLED C5', 2025), (r'C5[0-9]L', 'OLED C5', 2025), (r'C5[0-9]', 'OLED C5', 2025),
            (r'OLED\d{2}B5', 'OLED B5', 2025), (r'B5[0-9]L', 'OLED B5', 2025), (r'B5[0-9]', 'OLED B5', 2025),
            (r'NANO82A', 'NANO80A', 2025), (r'NANO81A', 'NANO80A', 2025), (r'NANO80A', 'NANO80A', 2025),
            (r'UA75', 'UA75', 2025)
        ]

        m_code = re.search(r'([A-Z]{2}\d{2}[A-Z0-9]+|\b\d{2}[A-Z]{1,3}\d{2,3}[A-Z0-9]*\b)', title_upper)
        if m_code:
            model_code = m_code.group(1)
        else:
            slug_m = re.search(r'/([a-z0-9\-]+)/?\d*$', href.lower())
            if slug_m:
                model_code = slug_m.group(1).upper()

        for pat, ser, yr in lg_2026:
            if re.search(pat, title_upper):
                series_code = ser
                year_val = yr
                break

        if not year_val:
            for pat, ser, yr in lg_2025:
                if re.search(pat, title_upper):
                    series_code = ser
                    year_val = yr
                    break

        if not year_val:
            if "2026" in title_upper:
                year_val = 2026
            elif "2025" in title_upper:
                year_val = 2025

    if not year_val:
        return None  # Purge older legacy models

    disp_type = "UHD 4K"
    if "OLED" in series_code.upper() or "OLED" in title_upper:
        disp_type = "OLED"
    elif "MRGB" in series_code.upper() or "MICRO RGB" in title_upper:
        disp_type = "MRGB"
    elif "QNED" in series_code.upper() or "QNED" in title_upper:
        disp_type = "QNED/QLED"
    elif any(x in series_code.upper() or x in title_upper for x in ["QLED", "NEO QLED", "MINI LED", "Q8", "Q7", "QN"]):
        disp_type = "QNED/QLED"

    return {
        "brand": brand.upper(),
        "title": title_upper,
        "series": series_code,
        "model_code": model_code,
        "size": size_val,
        "year": year_val,
        "display_type": disp_type
    }

def sync_greece_dump(dump_path=None):
    if not dump_path:
        dump_path = r"C:\Users\harry.park\.gemini\antigravity\brain\1c8f1fcd-c6b2-4806-abbc-780188e26593\.system_generated\steps\971\output.txt"

    print(f"[FAST GREECE SYNC] Reading dump from: {dump_path}")
    with open(dump_path, 'r', encoding='utf-8', errors='ignore') as f:
        text = f.read()

    pos = text.find('### Ran Playwright code')
    start = text.find('{')
    if start == -1:
        raise ValueError("Could not find JSON start in dump file.")
    
    json_str = text[start:pos].strip() if pos != -1 else text[start:].strip()
    data = json.loads(json_str)

    results = {}
    for brand in ['Samsung', 'LG']:
        raw_list = data.get('data', {}).get(brand, [])
        valid = []
        seen_urls = set()
        
        for item in raw_list:
            title = item.get('title', '')
            href = item.get('href', '')
            card_text = item.get('cardText', '')

            info = extract_model_code_and_info(title.upper(), brand, href)
            if not info:
                continue

            prices = re.findall(r'(\d[\d\.\s]*[,\.]\d{2})\s*€', card_text)
            if not prices:
                continue

            size_val = info['size']
            parsed_prices = []
            for p in prices:
                val = parse_price(p)
                if size_val >= 70 and val < 700.0:
                    continue
                if size_val >= 55 and val < 350.0:
                    continue
                if size_val >= 48 and val < 250.0:
                    continue
                if val >= 100.0:
                    parsed_prices.append(val)

            if not parsed_prices:
                continue

            selling_price = min(parsed_prices)
            orig_price = max(parsed_prices) if len(parsed_prices) > 1 else selling_price

            url = f"https://www.public.gr{href}" if href.startswith('/') else href
            if url in seen_urls:
                continue
            seen_urls.add(url)

            # Promotion detection
            promo = "None"
            if "Δώρο" in card_text or "Soundbar" in card_text or "επιστροφή" in card_text:
                m_pr = re.search(r'(\d+%\s*Public\s*επιστροφή|Δώρο[^\n\.\,]+|Soundbar[^\n\.\,]+)', card_text)
                if m_pr:
                    promo = m_pr.group(1).strip()

            clean_item = {
                "brand": brand,
                "title": title,
                "series": info["series"],
                "model_code": info["model_code"],
                "size": info["size"],
                "year": info["year"],
                "display_type": info["display_type"],
                "original_price": orig_price,
                "price": selling_price,
                "cashback": 0.0,
                "net_price": selling_price,
                "promotion": promo,
                "url": url,
                "image_url": ""
            }
            valid.append(clean_item)

        results[brand] = valid
        out_json = os.path.join(DATA_DIR, f"raw_public_gr_{brand.lower()}.json")
        with open(out_json, "w", encoding="utf-8") as f:
            json.dump(valid, f, ensure_ascii=False, indent=2)
        print(f"  ➔ Saved {len(valid)} items to {out_json}")

    # Now update Excel workbook
    candidates = glob.glob(os.path.join(DATA_DIR, "price tracker_EU_2026 0912*.xlsx"))
    if not candidates:
        candidates = glob.glob(os.path.join(DATA_DIR, "price tracker_EU_2026 *.xlsx"))
    target_wb_path = sorted(candidates)[-1]
    print(f"  ➔ Updating Excel workbook: {target_wb_path}")

    wb = openpyxl.load_workbook(target_wb_path)
    sheets_config = [
        ("Public_GR_Samsung", results["Samsung"]),
        ("Public_GR_LG", results["LG"])
    ]

    for sheet_name, items in sheets_config:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(title=sheet_name)
        headers = [
            "Brand", "Model Year", "Display Type", "Size (Inch)", "Model Code",
            "Price (EUR)", "Original Price (EUR)", "Series", "Cashback", "General Promotions",
            "URL", "Title"
        ]
        ws.append(headers)
        for it in items:
            ws.append([
                it.get("brand"),
                it.get("year"),
                it.get("display_type"),
                it.get("size"),
                it.get("model_code"),
                it.get("price"),
                it.get("original_price"),
                it.get("series"),
                it.get("cashback", 0),
                it.get("promotion", "None"),
                it.get("url"),
                it.get("title")
            ])

    wb.save(target_wb_path)
    wb.close()
    print(f"  ➔ Successfully updated sheets Public_GR_Samsung and Public_GR_LG in {target_wb_path}")

    # Mirror to History_EU
    hist_dir = os.path.join(HISTORY_DIR, "2026 0912")
    os.makedirs(hist_dir, exist_ok=True)
    hist_dest = os.path.join(hist_dir, os.path.basename(target_wb_path))
    shutil.copy2(target_wb_path, hist_dest)
    print(f"  ➔ Mirrored workbook to: {hist_dest}")

if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else None
    sync_greece_dump(path)
